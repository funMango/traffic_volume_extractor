from __future__ import annotations

import importlib.util
import sys
import tempfile
from datetime import date
from pathlib import Path

from openpyxl import Workbook


MODULE_PATH = (
    Path(__file__).resolve().parents[1]
    / "01_Program"
    / "09_else"
    / "samjeong_smart_edge_compare.py"
)
SPEC = importlib.util.spec_from_file_location("samjeong_smart_edge_compare", MODULE_PATH)
comparison = importlib.util.module_from_spec(SPEC)
assert SPEC and SPEC.loader
sys.modules[SPEC.name] = comparison
SPEC.loader.exec_module(comparison)


def test_target_direction_mapping_includes_samjeong_bridge_west_substitution() -> None:
    targets = {target.label: target for target in comparison.TARGETS}
    assert targets["삼정교사거리"].smart_approaches == ("삼정교사거리-서 (동향)",)
    assert targets["박촌교 삼거리"].smart_approaches == (
        "박촌교삼거리-북 (남향)",
        "박촌교삼거리-남 (북향)",
    )


def test_weekday_excludes_weekends_and_configured_holidays() -> None:
    assert not comparison.is_weekday(date(2026, 5, 1))
    assert not comparison.is_weekday(date(2026, 5, 5))
    assert not comparison.is_weekday(date(2026, 5, 9))
    assert comparison.is_weekday(date(2026, 5, 6))


def test_edge_missing_day_remains_in_denominator_but_smart_missing_day_is_excluded() -> None:
    first = date(2026, 5, 1)
    second = date(2026, 5, 2)
    values = {
        first: comparison.DailyTraffic(100, True),
        second: comparison.DailyTraffic(0, False),
    }
    assert comparison.average_daily(values, (first, second), False) == (50, 2)
    assert comparison.average_daily(values, (first, second), True) == (100, 1)


def test_peak_hours_are_exactly_the_two_half_open_windows() -> None:
    assert comparison.PEAK_HOURS == {7, 8, 17, 18}


def test_rounding_is_half_up() -> None:
    assert comparison.round_half_up(comparison.Decimal("10.5")) == 11


def test_v2_smart_peak_uses_fixed_divisor_without_hourly_quality_filter() -> None:
    class FakeReader:
        def load(self):
            values = {}
            for target in comparison.TARGETS:
                daily = {}
                for day in comparison.calendar_days(date(2026, 5, 1), date(2026, 6, 30)):
                    daily[day] = {}
                daily[date(2026, 5, 1)] = {hour: 10 for hour in range(24)}
                daily[date(2026, 5, 6)] = {hour: 10 for hour in range(24)}
                daily[date(2026, 5, 7)] = {hour: 20 for hour in range(24)}
                values[target.label] = daily
            return values

    payload = comparison.smart_v2_payload(FakeReader())
    result = payload[(5, "박촌교 삼거리")]
    assert result == {
        "monthly": 320,
        "weekday": 360,
        "peak": 6,
        "monthly_days": 3,
        "weekday_days": 2,
        "peak_days": 18,
    }


def test_hourly_quality_allows_three_invalid_hours_and_rejects_four() -> None:
    usable = {hour: 10 for hour in range(24)}
    for hour in (1, 2, 3):
        usable[hour] = 0
    rejected = {hour: 10 for hour in range(24)}
    for hour in (1, 2, 3, 4):
        rejected[hour] = 0
    values = {date(2026, 5, 6): usable, date(2026, 5, 7): rejected}

    assert comparison.average_hourly_days(values, values) == (210, 1)


def test_weekday_peak_sums_four_hours_and_truncates_with_fixed_month_divisor() -> None:
    days = [date(2026, 5, 6), date(2026, 5, 7)]
    values = {
        days[0]: {7: 10, 8: 0, 17: 5},
        days[1]: {7: 9, 8: 8, 17: 4, 18: 1},
    }

    assert comparison.average_weekday_peak_hours(values, days, 18) == (2, 18)
    assert comparison.average_weekday_peak_hours(values, days, 20) == (1, 20)


def test_weekday_peak_keeps_missing_and_zero_hours_in_fixed_denominator() -> None:
    days = [date(2026, 6, 4), date(2026, 6, 5)]

    assert comparison.average_weekday_peak_hours({days[0]: {7: 20}}, days, 20) == (1, 20)


def test_write_v2_workbook_changes_only_smart_peak_column() -> None:
    with tempfile.TemporaryDirectory(dir=Path.cwd()) as temporary_directory:
        workbook_path = Path(temporary_directory) / "comparison.xlsx"
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = comparison.V2_SHEET_NAME
        payload = {}
        for row in comparison.V2_TARGET_ROWS:
            month = 5 if row <= 9 else 6
            label = f"target-{row}"
            worksheet.cell(row, 2).value = label
            worksheet.cell(row, 4).value = f"=1+{row}"
            worksheet.cell(row, 7).value = f"=2+{row}"
            worksheet.cell(row, 10).value = -1
            worksheet.cell(row, 11).value = f"=J{row}/100"
            payload[(month, label)] = {"peak": row * 10}
        workbook.create_sheet("Preserved")["A1"] = "unchanged"
        workbook.save(workbook_path)
        workbook.close()

        before = comparison.v2_non_target_snapshot(workbook_path)
        comparison.write_v2_workbook(workbook_path, payload)
        comparison.verify_v2_workbook(workbook_path, payload, before)
