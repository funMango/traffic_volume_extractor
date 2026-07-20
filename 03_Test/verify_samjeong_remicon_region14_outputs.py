#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Verify the Samjeong-dong expanded-remicon output workbooks against SQLite."""

from __future__ import annotations

import re
import sqlite3
from collections import defaultdict
from datetime import date
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parents[1]
DB_PATH = BASE_DIR / "00_Data" / "삼정동_레미콘" / "삼정동_레미콘_데이터.db"
BASE_WORKBOOK = BASE_DIR / "02_Result" / "09_기타" / "삼정동_레미콘_교통량.xlsx"
OUTPUT_WORKBOOK = BASE_DIR / "02_Result" / "09_기타" / "삼정동_레미콘_교통량_지역명14포함.xlsx"
DETAIL_WORKBOOK = BASE_DIR / "02_Result" / "09_기타" / "삼정동_레미콘_지역명14차량_교차로별.xlsx"

LOCATION_MAPPING = {
    "박촌교 삼거리": [
        "박촌교 삼거리[남향]",
        "박촌교 삼거리[북향]",
        "박촌교 삼거리[서향(순방향)]",
        "박촌교 삼거리[서향(역방향)]",
    ],
    "봉오고가교사거리": ["봉오고가교 사거리[남향]", "봉오고가교 사거리[서향]"],
    "삼정고가삼거리": ["삼정고가 삼거리[동향]", "삼정고가 삼거리[서향]"],
    "삼정교사거리": ["삼정교 사거리[북동향]"],
    "산업길사거리": ["산업길 사거리[남향]", "산업길 사거리[북향]"],
    "봉오대로사거리": ["봉오대로 사거리[동향]", "봉오대로 사거리[서향]"],
    "자동차검사소": ["자동차검사소"],
    "삼정동 320-1": ["삼정동320-1 부천IC"],
}
MONTHS = {
    "5월": ("2026-05-01 00:00:00", "2026-06-01 00:00:00", range(4, 12)),
    "6월": ("2026-06-01 00:00:00", "2026-07-01 00:00:00", range(12, 20)),
}
HOLIDAYS = {date(2026, 5, 1), date(2026, 5, 5), date(2026, 5, 25), date(2026, 6, 3)}
OLD_PATTERN = re.compile(r"^014[가-힣]")
AREA14_PATTERN = re.compile(r"^[가-힣]{2}14[가-힣][0-9]\*{3}$")
DETAIL_HEADERS = [
    "id",
    "collected_at",
    "device_id",
    "location_name",
    "lane",
    "vehicle_number",
    "owner_registered_area",
    "source_file",
    "source_sheet",
]


def is_remicon(value: object) -> int:
    vehicle_number = str(value or "")
    return int(bool(OLD_PATTERN.match(vehicle_number) or AREA14_PATTERN.fullmatch(vehicle_number)))


def ratio(remicon: int, total: int) -> float:
    return (
        float(
            (Decimal(remicon) * 100 / Decimal(total)).quantize(
                Decimal("0.01"), rounding=ROUND_HALF_UP
            )
        )
        if total
        else 0.0
    )


def fetch_counts(
    conn: sqlite3.Connection,
    locations: list[str],
    start: str,
    end: str,
    where: str = "",
    params: tuple[object, ...] = (),
) -> tuple[int, int]:
    marks = ", ".join("?" for _ in locations)
    total, remicon = conn.execute(
        f"SELECT COUNT(*), SUM(IS_REMICON(vehicle_number)) FROM vehicle_detection WHERE collected_at >= ? AND collected_at < ? AND location_name IN ({marks}) {where}",
        (start, end, *locations, *params),
    ).fetchone()
    return int(total or 0), int(remicon or 0)


def verify_traffic_workbook() -> None:
    conn = sqlite3.connect(f"file:{DB_PATH.resolve().as_posix()}?mode=ro", uri=True)
    conn.create_function("IS_REMICON", 1, is_remicon)
    base = load_workbook(BASE_WORKBOOK, data_only=True)
    result = load_workbook(OUTPUT_WORKBOOK, data_only=True)
    try:
        base_ws, result_ws = base["작성서식"], result["작성서식"]
        for month, (start, end, rows) in MONTHS.items():
            for row in rows:
                label = str(result_ws[f"B{row}"].value or "").rstrip("*")
                locations = LOCATION_MAPPING[label]
                weekday = "AND CAST(strftime('%w', collected_at) AS INTEGER) NOT IN (0, 6) AND date(collected_at) NOT IN (?, ?, ?, ?)"
                weekday_params = tuple(day.isoformat() for day in sorted(HOLIDAYS))
                peak = (
                    weekday
                    + " AND ((time(collected_at) >= '07:00:00' AND time(collected_at) < '09:00:00') OR (time(collected_at) >= '17:00:00' AND time(collected_at) < '19:00:00'))"
                )
                expected = [
                    *fetch_counts(conn, locations, start, end),
                    *fetch_counts(conn, locations, start, end, weekday, weekday_params),
                    *fetch_counts(conn, locations, start, end, peak, weekday_params),
                ]
                actual = [result_ws.cell(row, col).value for col in (3, 4, 6, 7, 9, 10)]
                if actual != expected:
                    raise AssertionError(
                        f"DB 재집계 불일치: {month} {label}: {actual} != {expected}"
                    )
                for count_col, total, remicon in (
                    (5, expected[0], expected[1]),
                    (8, expected[2], expected[3]),
                    (11, expected[4], expected[5]),
                ):
                    if result_ws.cell(row, count_col).value != ratio(remicon, total):
                        raise AssertionError(f"비율 불일치: {month} {label} col {count_col}")
                if [base_ws.cell(row, col).value for col in (3, 6, 9)] != [
                    result_ws.cell(row, col).value for col in (3, 6, 9)
                ]:
                    raise AssertionError(f"총 통행량 변경: {month} {label}")
    finally:
        result.close()
        base.close()
        conn.close()


def verify_detail_workbook() -> None:
    workbook = load_workbook(DETAIL_WORKBOOK, read_only=False, data_only=True)
    try:
        if len(workbook.sheetnames) != 15:
            raise AssertionError(f"상세 시트 수 불일치: {len(workbook.sheetnames)}")
        counts: dict[str, int] = defaultdict(int)
        for sheet_name in workbook.sheetnames:
            ws = workbook[sheet_name]
            if (
                [cell.value for cell in ws[1]] != DETAIL_HEADERS
                or ws.freeze_panes != "A2"
                or ws.auto_filter.ref != ws.dimensions
            ):
                raise AssertionError(f"상세 시트 형식 불일치: {sheet_name}")
            previous: tuple[str, int] | None = None
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not AREA14_PATTERN.fullmatch(str(row[5] or "")):
                    raise AssertionError(f"상세 번호판 형식 불일치: {sheet_name}")
                key = (str(row[1] or ""), int(row[0]))
                if previous is not None and key < previous:
                    raise AssertionError(f"상세 정렬 불일치: {sheet_name}")
                previous = key
                counts[sheet_name] += 1
        if sum(counts.values()) != 47_935:
            raise AssertionError(f"상세 총 건수 불일치: {sum(counts.values())}")
    finally:
        workbook.close()


if __name__ == "__main__":
    verify_traffic_workbook()
    verify_detail_workbook()
    print("traffic_db_reaggregation=matched")
    print("totals_preserved=matched")
    print("detail_workbook=matched")
