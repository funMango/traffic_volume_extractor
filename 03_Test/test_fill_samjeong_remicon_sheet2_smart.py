from __future__ import annotations

import importlib.util
import sys
import tempfile
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill
import pytest


MODULE_PATH = (
    Path(__file__).resolve().parents[1]
    / "01_Program"
    / "09_else"
    / "fill_samjeong_remicon_sheet2_smart.py"
)
SPEC = importlib.util.spec_from_file_location("fill_sheet2_smart", MODULE_PATH)
sheet2 = importlib.util.module_from_spec(SPEC)
assert SPEC and SPEC.loader
sys.modules[SPEC.name] = sheet2
SPEC.loader.exec_module(sheet2)


def test_weekday_dates_use_fixed_calendar_divisors() -> None:
    assert len(sheet2.weekday_dates(sheet2.PERIOD_BY_MONTH["5월"])) == 18
    assert len(sheet2.weekday_dates(sheet2.PERIOD_BY_MONTH["6월"])) == 21


def test_metrics_keep_missing_dates_and_hours_as_zero_and_round_half_up() -> None:
    period = sheet2.PERIOD_BY_MONTH["5월"]
    values = {
        date(2026, 5, 6): {7: 31, 8: 1, 17: 1, 18: 1},
        date(2026, 5, 7): {7: 1},
    }
    assert sheet2.calculate_metrics(values, period) == (1, 2, 2)


def test_monthly_metrics_are_limited_to_each_month() -> None:
    values = {
        date(2026, 5, 6): {7: 31},
        date(2026, 6, 4): {7: 300},
    }

    assert sheet2.calculate_metrics(values, sheet2.PERIOD_BY_MONTH["5월"]) == (1, 2, 2)
    assert sheet2.calculate_metrics(values, sheet2.PERIOD_BY_MONTH["6월"]) == (10, 14, 14)


def test_monthly_average_validation_reports_mapping_and_daily_values() -> None:
    with pytest.raises(AssertionError, match="박촌교삼거리-북") as error:
        sheet2.validate_monthly_average(
            (100, 99, 10),
            sheet2.PERIOD_BY_MONTH["5월"],
            "박촌교 삼거리",
            "남향",
            {date(2026, 5, 6): {7: 100}},
        )

    assert "2026-05-06" in str(error.value)


def test_apply_only_changes_smart_cells_and_writes_total_values() -> None:
    with tempfile.TemporaryDirectory(dir=Path.cwd()) as temporary_directory:
        workbook_path = Path(temporary_directory) / "sheet2.xlsx"
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = sheet2.SHEET_NAME
        row = 4
        for month in ("5월", "6월"):
            for intersection in (
                "박촌교 삼거리",
                "봉오고가교 사거리",
                "삼정고가 삼거리",
                "삼정교 사거리",
                "산업길 사거리",
                "봉오대로 사거리",
            ):
                directions = [key[1] for key in sheet2.APPROACH_NAMES if key[0] == intersection]
                worksheet.cell(row, 1).value = month
                worksheet.cell(row, 2).value = intersection
                start_row = row
                for direction in directions:
                    worksheet.cell(row, 3).value = direction
                    row += 1
                if len(directions) > 1:
                    worksheet.cell(row, 3).value = "합계"
                    row += 1
                worksheet.cell(start_row, 4).value = "preserved"
        workbook.create_sheet("Sheet1")["A1"] = "unchanged"
        workbook["Sheet1"]["A1"].fill = PatternFill("solid", fgColor="FFFF00")
        worksheet["D4"] = "preserved style"
        worksheet["D4"].fill = PatternFill("solid", fgColor="00FF00")
        worksheet.merge_cells("D4:E4")
        workbook.save(workbook_path)
        workbook.close()

        workbook = sheet2.load_workbook(workbook_path, data_only=False)
        try:
            before = {
                (sheet.title, cell.coordinate): (cell.value, cell.style_id)
                for sheet in workbook.worksheets
                for row_cells in sheet.iter_rows()
                for cell in row_cells
                if cell.value is not None
            }
            groups = sheet2.find_direction_groups(workbook[sheet2.SHEET_NAME])
            targets = sheet2.target_coordinates(groups)
            before_merges = {
                sheet.title: tuple(str(merged) for merged in sheet.merged_cells.ranges)
                for sheet in workbook.worksheets
            }
        finally:
            workbook.close()

        workbook = sheet2.load_workbook(workbook_path)
        try:
            groups = sheet2.find_direction_groups(workbook[sheet2.SHEET_NAME])
            values = {
                direction_row: (11, 22, 33)
                for group in groups
                for direction_row in group.direction_rows
            }
        finally:
            workbook.close()
        sheet2.apply_or_verify(workbook_path, values, False)

        workbook = sheet2.load_workbook(workbook_path, data_only=False)
        try:
            assert workbook["Sheet1"]["A1"].value == "unchanged"
            assert workbook["Sheet1"]["A1"].fill.fgColor.rgb == "00FFFF00"
            assert workbook[sheet2.SHEET_NAME]["D4"].fill.fgColor.rgb == "0000FF00"
            assert {
                sheet.title: tuple(str(merged) for merged in sheet.merged_cells.ranges)
                for sheet in workbook.worksheets
            } == before_merges
            for group in sheet2.find_direction_groups(workbook[sheet2.SHEET_NAME]):
                for direction_row in group.direction_rows:
                    assert tuple(
                        workbook[sheet2.SHEET_NAME].cell(direction_row, col).value
                        for col in (7, 9, 11)
                    ) == (11, 22, 33)
                if group.total_row is not None:
                    assert workbook[sheet2.SHEET_NAME].cell(group.total_row, 7).value == (
                        11 * len(group.direction_rows)
                    )
            after = {
                (sheet.title, cell.coordinate): (cell.value, cell.style_id)
                for sheet in workbook.worksheets
                for row_cells in sheet.iter_rows()
                for cell in row_cells
                if cell.value is not None
            }
            assert {
                coordinate: state
                for coordinate, state in after.items()
                if coordinate not in targets
            } == {
                coordinate: state
                for coordinate, state in before.items()
                if coordinate not in targets
            }
        finally:
            workbook.close()
