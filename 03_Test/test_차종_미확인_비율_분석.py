from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook


MODULE_PATH = Path(__file__).parents[1] / "01_Program" / "00_main" / "차종_미확인_비율_분석.py"
SPEC = importlib.util.spec_from_file_location("unidentified_vehicle_analysis", MODULE_PATH)
assert SPEC is not None and SPEC.loader is not None
analysis = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = analysis
SPEC.loader.exec_module(analysis)


def test_period_target_and_hour_parsing_preserves_direction_spaces():
    periods = analysis.parse_periods("260708~260709, 2608, 260901")
    assert [(period.start.isoformat(), period.end.isoformat()) for period in periods] == [
        ("2026-07-08", "2026-07-09"),
        ("2026-08-01", "2026-08-31"),
        ("2026-09-01", "2026-09-01"),
    ]
    assert analysis.parse_target("A교차로-북 (동향)") == analysis.Target("A교차로", "북 (동향)")
    assert analysis.parse_hours("첨두시") == [7, 8]
    assert analysis.parse_hours("일반시간") == [7, 8, 9]
    assert analysis.parse_hours("07~09, 17~19") == [7, 8, 17, 18]
    assert analysis.parse_hours("첨두시, 17~19") == [7, 8, 17, 18]


@pytest.mark.parametrize(
    "sql", ["DELETE FROM x", "SELECT 1; SELECT 2", "SELECT 1;", "SET TRANSACTION READ ONLY"]
)
def test_select_guard_blocks_everything_except_one_select(sql):
    with pytest.raises(RuntimeError):
        analysis.ensure_select_sql(sql)


def test_sql_aggregates_codes_zero_and_one_without_quantity_threshold():
    sql, params = analysis.build_day_sql(analysis.Target("A", "북 (동향)"), [7, 8])
    compact = " ".join(sql.split())
    assert "S_CRSRD_VKND_TRF_1HH" in compact
    assert "IN ('0', '1')" in compact
    assert "TRF_QNTY >" not in compact
    assert "M_CRSRD_INF" in compact and "M_CRSRD_ACSR_INF" in compact and "M_CD_INF" in compact
    assert params["direction_name"] == "북 (동향)"


def test_combine_rows_and_zero_denominator_xlsx_layout(tmp_path):
    rows = analysis.combine_rows(
        [
            analysis.AnalysisRow(
                "2026.07.08 00:00~2026.07.09 23:59", "A", "북 (동향)", "좌", 2, 10
            ),
            analysis.AnalysisRow("2026.07.08 00:00~2026.07.09 23:59", "A", "북 (동향)", "좌", 3, 5),
            analysis.AnalysisRow("2026.07.08 00:00~2026.07.09 23:59", "B", "남", "직", 0, 0),
        ]
    )
    output = tmp_path / "result.xlsx"
    analysis.save_xlsx(
        rows,
        [analysis.Period(analysis.date(2026, 7, 8), analysis.date(2026, 7, 9))],
        analysis.Target(),
        [7, 8],
        output,
    )
    workbook = load_workbook(output, data_only=False)
    sheet = workbook["미확인 분석"]
    assert [sheet.cell(10, column).value for column in range(1, 8)] == [
        "기간",
        "교차로",
        "방향",
        "접근로",
        "미확인 교통량",
        "전체 교통량",
        "미확인 교통량 비율(%)",
    ]
    assert (
        sheet["E11"].value == 5
        and sheet["F11"].value == 15
        and sheet["G11"].value == "=IF(F11=0,0,E11/F11)"
    )
    assert sheet["G11"].number_format == "0.00%" and sheet["G12"].value == "=IF(F12=0,0,E12/F12)"
    assert sheet.auto_filter.ref == "A10:G12" and sheet.freeze_panes == "A11"


def test_worker_choice_prefers_lower_count_within_five_percent():
    assert analysis.choose_worker_count({1: 10.4, 2: 10.0, 3: 9.0}) == 3
    assert analysis.choose_worker_count({1: 10.6, 2: 10.0, 3: 10.1}) == 2


def test_menu_hides_tests_and_keeps_analysis_last():
    init_spec = importlib.util.spec_from_file_location(
        "project_init", Path(__file__).parents[1] / "init.py"
    )
    assert init_spec is not None and init_spec.loader is not None
    project_init = importlib.util.module_from_spec(init_spec)
    init_spec.loader.exec_module(project_init)
    programs = project_init._programs()
    assert all(not path.name.startswith("test_") for path in programs)
    assert programs[-1].name == "차종_미확인_비율_분석.py"
