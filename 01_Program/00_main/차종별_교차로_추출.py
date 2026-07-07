#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""차종별 교차로 교통량 추출 프로그램.

S_CRSRD_VKND_TRF_1HH를 읽기 전용 SELECT로 조회해
시간, 교차로, 교차로 방향, 방향, 차종, 교통량 컬럼의 결과를 생성한다.
"""

from __future__ import annotations

import argparse
from concurrent.futures import FIRST_COMPLETED, ProcessPoolExecutor, wait
import csv
import difflib
import json
import os
import re
import sqlite3
import sys
import time
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from multiprocessing import Manager
from pathlib import Path
from queue import Empty
from tempfile import TemporaryDirectory
from threading import Event, Lock, Thread
from typing import Any, Callable, Iterable, Iterator

try:
    import oracledb
except ImportError:  # pragma: no cover - 실행 환경 안내용
    oracledb = None

try:
    from dotenv import load_dotenv
except ImportError:  # pragma: no cover - 실행 환경 안내용
    load_dotenv = None


BASE_DIR = Path(__file__).resolve().parent.parent
ENV_PATH = BASE_DIR.parent / "00_Data" / ".env"
VKND_KIND_PATH = BASE_DIR.parent / "00_Data" / "VKND_KIND.json"
RESULT_DIR = BASE_DIR / "02_Result" / "차종별_교통량_추출"
FIXED_SQLITE_OUTPUT_PATH = (
    BASE_DIR
    / "02_Result"
    / "교통량_추출"
    / "교통량"
    / "전체교차로_차종별_2505~2605.db"
)
TRAFFIC_TABLE = "S_CRSRD_VKND_TRF_1HH"
INTERSECTION_TABLE = "M_CRSRD_INF"
APPROACH_TABLE = "M_CRSRD_ACSR_INF"
CODE_TABLE = "M_CD_INF"
FIXED_START_SQL = "DATE '2025-05-01'"
FIXED_END_SQL = "DATE '2026-06-01'"
FIXED_START_TEXT = "2025-05-01 00:00:00"
FIXED_END_TEXT = "2026-06-01 00:00:00"
SQLITE_INSERT_BATCH_SIZE = 10000
EXCLUDED_VEHICLE_CODES = {"0", "1"}
RESULT_HEADER = ["시간", "교차로", "교차로 방향", "방향", "차종", "교통량"]
CSV_HEADER = RESULT_HEADER
SQLITE_TABLE = "traffic_by_vehicle"
SQLITE_COLUMNS = [
    "observed_at",
    "intersection",
    "approach_direction",
    "direction",
    "vehicle_type",
    "traffic_volume",
    "vehicle_code",
]
SQLITE_INDEX_DEFINITIONS = [
    (
        "idx_traffic_intersection_time",
        f"""
        CREATE INDEX IF NOT EXISTS idx_traffic_intersection_time
            ON {SQLITE_TABLE} (observed_at, intersection)
        """,
    ),
    (
        "idx_traffic_time",
        f"""
        CREATE INDEX IF NOT EXISTS idx_traffic_time
            ON {SQLITE_TABLE} (observed_at)
        """,
    ),
    (
        "idx_traffic_vehicle_type",
        f"""
        CREATE INDEX IF NOT EXISTS idx_traffic_vehicle_type
            ON {SQLITE_TABLE} (vehicle_type)
        """,
    ),
    (
        "idx_traffic_direction",
        f"""
        CREATE INDEX IF NOT EXISTS idx_traffic_direction
            ON {SQLITE_TABLE} (direction)
        """,
    ),
    (
        "idx_traffic_sort",
        f"""
        CREATE INDEX IF NOT EXISTS idx_traffic_sort
            ON {SQLITE_TABLE} (
                observed_at,
                intersection,
                approach_direction,
                direction,
                vehicle_type
            )
        """,
    ),
]
ALL_INTERSECTION_TOKENS = {"*", "전체", "all"}
PEAK_HOURS = [7, 8, 16, 17, 18]
OUTPUT_FORMAT_EXTENSIONS = {
    "csv": ".csv",
    "xlsx": ".xlsx",
    "db": ".db",
}
OUTPUT_FORMAT_LABELS = {
    "csv": "CSV 저장",
    "xlsx": "XLSX 저장",
    "db": "SQLite DB 저장",
}
DEFAULT_PARALLEL_WORKERS = 5
BLOCKED_SQL_PATTERN = re.compile(
    r"\b(INSERT|UPDATE|DELETE|MERGE|DROP|TRUNCATE|ALTER|CREATE|COMMIT|ROLLBACK|GRANT|REVOKE)\b",
    re.IGNORECASE,
)


@dataclass(frozen=True)
class Period:
    start: date
    end: date
    token: str


@dataclass(frozen=True)
class Intersection:
    node_id: int | str
    name: str


@dataclass(frozen=True)
class WorkerChunkResult:
    worker_index: int
    chunk_path: str
    row_count: int
    chunk_count: int = 1


@dataclass
class WorkerProgressState:
    worker_index: int
    assigned_days: int
    completed_days: int = 0
    row_count: int = 0
    status: str = "pending"
    current_label: str = ""
    last_completed_label: str = ""
    merge_total_rows: int = 0
    merge_completed_rows: int = 0
    error_message: str = ""


class StepProgress:
    """스피너와 퍼센트를 함께 출력하는 간단한 진행 표시기."""

    def __init__(self, label: str, total: int = 1) -> None:
        self.label = label
        self.total = max(total, 1)
        self.done = 0
        self._frames = ("|", "/", "-", "\\")
        self._frame_index = 0
        self._last_len = 0
        self._lock = Lock()
        self._stop = Event()
        self._thread: Thread | None = None

    def start(self) -> None:
        if self._thread is not None:
            return
        self._thread = Thread(target=self._run, daemon=True)
        self._thread.start()

    def advance(self, step: int = 1) -> None:
        with self._lock:
            self.done = min(self.total, self.done + step)
        self._render()

    def finish(self) -> None:
        with self._lock:
            self.done = self.total
        self._stop.set()
        if self._thread is not None:
            self._thread.join(timeout=1.0)
        self._render()
        print()

    def _run(self) -> None:
        while not self._stop.wait(0.12):
            self._frame_index = (self._frame_index + 1) % len(self._frames)
            self._render()

    def _render(self) -> None:
        with self._lock:
            done = self.done
            total = self.total
        pct = int((done / total) * 100)
        frame = self._frames[self._frame_index]
        line = f"\r{self.label} {frame} {pct:3d}% ({done}/{total})"
        pad = " " * max(0, self._last_len - len(line))
        print(line + pad, end="", flush=True)
        self._last_len = len(line)


def run_instant_step(label: str, work) -> object:
    progress = StepProgress(label)
    progress.start()
    try:
        result = work()
    finally:
        progress.finish()
    return result


def print_step_done(label: str) -> None:
    print(f"{label} | 100% (1/1)")


def connect_db():
    """00_Data/.env를 로드한 뒤 DB에 연결한다."""
    if oracledb is None:
        raise RuntimeError("oracledb 패키지가 설치되어 있지 않습니다.")
    if load_dotenv is None:
        raise RuntimeError("python-dotenv 패키지가 설치되어 있지 않습니다.")

    load_dotenv(ENV_PATH)

    required = ["DB_USER", "DB_PASSWORD", "DB_HOST", "DB_SERVICE"]
    missing = [name for name in required if not os.getenv(name)]
    if missing:
        raise RuntimeError(
            "DB 접속 환경변수가 없습니다: "
            + ", ".join(missing)
            + f"\n.env 경로를 확인하세요: {ENV_PATH}"
        )

    conn = oracledb.connect(
        user=os.getenv("DB_USER", "").strip(),
        password=os.getenv("DB_PASSWORD", "").strip(),
        host=os.getenv("DB_HOST", "").strip(),
        port=int(os.getenv("DB_PORT", "1521").strip()),
        service_name=os.getenv("DB_SERVICE", "").strip(),
    )
    with conn.cursor() as cur:
        cur.execute("SET TRANSACTION READ ONLY")
    return conn


def strip_sql_comments(sql: str) -> str:
    sql = re.sub(r"/\*.*?\*/", " ", sql, flags=re.DOTALL)
    return re.sub(r"--[^\n\r]*", " ", sql)


def ensure_select_sql(sql: str) -> None:
    normalized = strip_sql_comments(sql).strip()
    if not normalized:
        raise RuntimeError("빈 SQL은 실행할 수 없습니다.")
    if not normalized.upper().startswith("SELECT"):
        raise RuntimeError("읽기 전용 보호: SELECT 문만 실행할 수 있습니다.")
    if ";" in normalized.rstrip(";"):
        raise RuntimeError("읽기 전용 보호: 복수 SQL 문은 실행할 수 없습니다.")
    if BLOCKED_SQL_PATTERN.search(normalized):
        raise RuntimeError("읽기 전용 보호: 쓰기/DDL/트랜잭션 SQL은 실행할 수 없습니다.")


def execute_select(cursor: Any, sql: str, params: dict[str, Any] | None = None) -> list[tuple]:
    ensure_select_sql(sql)
    cursor.execute(sql, params or {})
    return list(cursor.fetchall())


def normalize_code(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def load_vehicle_kind_names(path: Path = VKND_KIND_PATH) -> dict[str, str]:
    with path.open("r", encoding="utf-8-sig") as f:
        raw_data = json.load(f)

    if not isinstance(raw_data, dict):
        raise RuntimeError(f"차종 JSON 형식이 올바르지 않습니다: {path}")

    vehicle_kind_names: dict[str, str] = {}
    for raw_code, raw_name in raw_data.items():
        code = normalize_code(raw_code)
        if not code or code in EXCLUDED_VEHICLE_CODES:
            continue
        name = normalize_code(raw_name) or code
        vehicle_kind_names[code] = name

    if not vehicle_kind_names:
        raise RuntimeError(f"사용 가능한 차종 코드가 없습니다: {path}")
    return vehicle_kind_names


def parse_yymmdd(value: str) -> date:
    if not re.fullmatch(r"\d{6}", value):
        raise ValueError(f"날짜 형식이 올바르지 않습니다: {value}")
    return date(2000 + int(value[:2]), int(value[2:4]), int(value[4:6]))


def month_end(year: int, month: int) -> date:
    if month == 12:
        return date(year, 12, 31)
    return date(year, month + 1, 1) - timedelta(days=1)


def parse_period_segment(segment: str) -> Period:
    token = re.sub(r"\s+", "", segment)
    if re.fullmatch(r"\d{6}~\d{6}", token):
        start_raw, end_raw = token.split("~", 1)
        start = parse_yymmdd(start_raw)
        end = parse_yymmdd(end_raw)
    elif re.fullmatch(r"\d{4}", token):
        year = 2000 + int(token[:2])
        month = int(token[2:4])
        start = date(year, month, 1)
        end = month_end(year, month)
    elif re.fullmatch(r"\d{6}", token):
        start = parse_yymmdd(token)
        end = start
    else:
        raise ValueError(f"기간 형식이 올바르지 않습니다: {segment}")
    if start > end:
        raise ValueError(f"시작일이 종료일보다 늦습니다: {segment}")
    return Period(start=start, end=end, token=token)


def parse_periods(raw: str) -> list[Period]:
    parts = [part.strip() for part in raw.split(",") if part.strip()]
    if not parts:
        raise ValueError("기간을 입력해 주세요.")
    return [parse_period_segment(part) for part in parts]


def parse_time_range(raw: str) -> list[int]:
    value = raw.strip().lower()
    if value in {"24시간", "1"}:
        return list(range(24))
    if value in {"첨두시", "2"}:
        return PEAK_HOURS.copy()

    if value in {"직접입력", "직접", "3"}:
        value = input("시간대 직접 입력 (예: 07~09, 18~20): ").strip()

    hours: set[int] = set()
    for segment in [part.strip() for part in value.split(",") if part.strip()]:
        match = re.fullmatch(r"(\d{1,2})(?::00)?\s*~\s*(\d{1,2})(?::00)?", segment)
        if not match:
            raise ValueError(f"시간대 형식이 올바르지 않습니다: {segment}")
        start_hour = int(match.group(1))
        end_hour = int(match.group(2))
        if not (0 <= start_hour <= 23 and 1 <= end_hour <= 24 and start_hour < end_hour):
            raise ValueError(f"시간대 범위가 올바르지 않습니다: {segment}")
        hours.update(range(start_hour, end_hour))
    if not hours:
        raise ValueError("시간대를 입력해 주세요.")
    return sorted(hours)


def input_periods() -> list[Period]:
    print("\n[기간 입력]")
    print("  예: 260410~260412 / 2604 / 260410 / 2604, 2605")
    while True:
        try:
            periods = parse_periods(input("기간: "))
        except ValueError as exc:
            print(f"  오류: {exc}")
            continue
        for period in periods:
            print(f"  - {period.token}: {period.start} ~ {period.end}")
        return periods


def input_hours() -> list[int]:
    print("\n[시간대 입력]")
    print("  1 또는 24시간: 0~23시")
    print("  2 또는 첨두시: 07~09, 16~19 (7, 8, 16, 17, 18시)")
    print("  3 또는 직접입력: 예) 07~09, 18~20")
    while True:
        try:
            hours = parse_time_range(input("시간대: "))
        except ValueError as exc:
            print(f"  오류: {exc}")
            continue
        print(f"  선택 시간: {', '.join(str(h) for h in hours)}시")
        return hours


def input_output_format() -> str:
    print("\n[출력 방식 입력]")
    print("  1 또는 csv  : CSV 파일")
    print("  2 또는 xlsx : Excel 파일")
    print("  3 또는 db   : SQLite DB 파일")
    print("  기본값: csv")
    choices = {
        "": "csv",
        "1": "csv",
        "csv": "csv",
        "2": "xlsx",
        "xlsx": "xlsx",
        "3": "db",
        "db": "db",
    }
    while True:
        selected = choices.get(input("출력 방식: ").strip().lower())
        if selected is None:
            print("  오류: csv, xlsx, db 또는 1, 2, 3 중 하나를 입력해 주세요.")
            continue
        print(f"  선택 출력 방식: {selected}")
        return selected


def load_reference_data(conn) -> tuple[list[Intersection], dict[str, str], dict[str, str]]:
    with conn.cursor() as cur:
        intersections = [
            Intersection(node_id=row[0], name=row[1])
            for row in execute_select(
                cur,
                f"SELECT NODE_ID, CRSRD_NM FROM {INTERSECTION_TABLE} ORDER BY CRSRD_NM",
            )
        ]

        drct_rows = execute_select(
            cur,
            """
            SELECT TRIM(TO_CHAR(CD)) AS CD, CD_NM
            FROM M_CD_INF
            WHERE GRP_CD = 'DRCT_CD'
            """,
        )
        drct_names = {normalize_code(code): name for code, name in drct_rows}

    vknd_names = load_vehicle_kind_names()
    return intersections, drct_names, vknd_names


def get_table_columns(cursor: Any, table_name: str) -> set[str]:
    rows = execute_select(
        cursor,
        """
        SELECT COLUMN_NAME
        FROM ALL_TAB_COLUMNS
        WHERE TABLE_NAME = :table_name
        """,
        {"table_name": table_name.upper()},
    )
    return {str(row[0]).upper() for row in rows}


def validate_vehicle_extract_tables(cursor: Any) -> None:
    required_columns = {
        TRAFFIC_TABLE: {"NODE_ID", "TOT_DT", "ACSR_ID", "DRCT_CD", "VKND_CD", "TRF_QNTY"},
        INTERSECTION_TABLE: {"NODE_ID", "CRSRD_NM"},
        APPROACH_TABLE: {"NODE_ID", "ACSR_ID", "ACSR_NM"},
        CODE_TABLE: {"GRP_CD", "CD", "CD_NM"},
    }

    missing_messages = []
    for table_name, required in required_columns.items():
        columns = get_table_columns(cursor, table_name)
        missing = sorted(required - columns)
        if missing:
            missing_messages.append(f"{table_name}: {', '.join(missing)}")

    if missing_messages:
        raise RuntimeError("필수 컬럼이 없습니다: " + " / ".join(missing_messages))


def resolve_intersections(raw: str, all_intersections: list[Intersection]) -> tuple[list[Intersection], bool]:
    raw = raw.strip()
    if raw.lower() in ALL_INTERSECTION_TOKENS:
        return all_intersections.copy(), True

    by_name = {item.name: item for item in all_intersections}
    names = [item.name for item in all_intersections]
    selected: list[Intersection] = []
    seen: set[int | str] = set()

    for part in [part.strip() for part in raw.split(",") if part.strip()]:
        item = by_name.get(part)
        if item is None:
            candidates = difflib.get_close_matches(part, names, n=5, cutoff=0.3)
            candidate_text = ", ".join(candidates) if candidates else "후보 없음"
            raise LookupError(f"'{part}'을(를) 찾을 수 없습니다. 유사 교차로: {candidate_text}")
        if item.node_id not in seen:
            selected.append(item)
            seen.add(item.node_id)

    if not selected:
        raise ValueError("교차로를 입력해 주세요.")
    return selected, False


def input_intersections(all_intersections: list[Intersection]) -> tuple[list[Intersection], bool]:
    print("\n[교차로 입력]")
    print("  전체: *, 전체, all")
    print("  복수: 교차로이름, 교차로이름")
    print(f"  전체 교차로 수: {len(all_intersections)}개")
    while True:
        try:
            selected, is_all = resolve_intersections(input("교차로: "), all_intersections)
        except (LookupError, ValueError) as exc:
            print(f"  오류: {exc}")
            continue
        if is_all:
            print(f"  전체 교차로 {len(selected)}개 선택")
        else:
            print(f"  선택 교차로: {', '.join(item.name for item in selected)}")
        return selected, is_all


def make_safe_filename_part(value: str) -> str:
    cleaned = re.sub(r'[\\/*?:"<>|]', "_", value).strip()
    return cleaned or "_"


def make_filename(intersections: list[Intersection], is_all: bool, period_tokens: Iterable[str]) -> str:
    period_part = "_".join(period_tokens)
    if is_all:
        name_part = "전체교차로"
    elif len(intersections) == 1:
        name_part = intersections[0].name
    else:
        name_part = f"{intersections[0].name}외{len(intersections) - 1}개"
    return f"{make_safe_filename_part(name_part)}_{period_part}"


def build_in_clause(prefix: str, values: Iterable[object], bind_params: dict) -> str:
    names = []
    for index, value in enumerate(values):
        bind_name = f"{prefix}{index}"
        bind_params[bind_name] = value
        names.append(f":{bind_name}")
    return ", ".join(names)


def build_vehicle_name_case_expression(
    vehicle_kind_names: dict[str, str],
    code_expression: str,
    bind_params: dict[str, object],
) -> str:
    parts = ["CASE"]
    for index, (code, name) in enumerate(vehicle_kind_names.items()):
        code_bind = f"vknd_case_code{index}"
        name_bind = f"vknd_case_name{index}"
        bind_params[code_bind] = code
        bind_params[name_bind] = name
        parts.append(f"WHEN {code_expression} = :{code_bind} THEN :{name_bind}")
    parts.append(f"ELSE {code_expression} END")
    return "\n                ".join(parts)


def build_vehicle_traffic_sql(
    vehicle_kind_names: dict[str, str],
    bind_params: dict[str, object],
    *,
    start_expression: str,
    end_expression: str,
    node_filter: bool = False,
    hour_clause: str = "",
    order_by_expressions: tuple[str, ...] | None = None,
) -> str:
    if not vehicle_kind_names:
        raise RuntimeError("조회 가능한 차종 코드가 없습니다.")

    vknd_code_expression = "TRIM(TO_CHAR(v.VKND_CD))"
    drct_code_expression = "TRIM(TO_CHAR(v.DRCT_CD))"
    vknd_in_clause = build_in_clause("vknd", vehicle_kind_names.keys(), bind_params)
    vknd_name_expression = build_vehicle_name_case_expression(
        vehicle_kind_names,
        vknd_code_expression,
        bind_params,
    )

    node_condition = "AND v.NODE_ID = :node_id" if node_filter else ""
    if order_by_expressions is None:
        order_by_expressions = (
            "v.TOT_DT",
            "c.CRSRD_NM",
            "NVL(a.ACSR_NM, '')",
            f"NVL(d.CD_NM, {drct_code_expression})",
            "VKND_NM",
        )
    order_by_clause = ",\n            ".join(order_by_expressions)
    return f"""
        SELECT
            v.TOT_DT AS TOT_DT,
            c.CRSRD_NM AS CRSRD_NM,
            a.ACSR_NM AS ACSR_NM,
            {drct_code_expression} AS DRCT_CD,
            NVL(d.CD_NM, {drct_code_expression}) AS DRCT_NM,
            {vknd_code_expression} AS VKND_CD,
            {vknd_name_expression} AS VKND_NM,
            SUM(NVL(v.TRF_QNTY, 0)) AS TRF_QNTY
        FROM {TRAFFIC_TABLE} v
        JOIN {INTERSECTION_TABLE} c
          ON c.NODE_ID = v.NODE_ID
        LEFT JOIN {APPROACH_TABLE} a
          ON a.NODE_ID = v.NODE_ID
         AND a.ACSR_ID = v.ACSR_ID
        LEFT JOIN {CODE_TABLE} d
          ON d.GRP_CD = 'DRCT_CD'
         AND TRIM(TO_CHAR(d.CD)) = {drct_code_expression}
        WHERE v.TOT_DT >= {start_expression}
          AND v.TOT_DT < {end_expression}
          {node_condition}
          {hour_clause}
          AND {vknd_code_expression} IN ({vknd_in_clause})
        GROUP BY
            v.TOT_DT,
            c.CRSRD_NM,
            a.ACSR_NM,
            {drct_code_expression},
            d.CD_NM,
            {vknd_code_expression},
            {vknd_name_expression}
        HAVING SUM(NVL(v.TRF_QNTY, 0)) > 0
        ORDER BY
            {order_by_clause}
    """


def build_all_intersections_vehicle_sql(
    vehicle_kind_names: dict[str, str],
) -> tuple[str, dict[str, object]]:
    bind_params: dict[str, object] = {}
    sql = build_vehicle_traffic_sql(
        vehicle_kind_names,
        bind_params,
        start_expression=FIXED_START_SQL,
        end_expression=FIXED_END_SQL,
    )
    return sql, bind_params


def build_hour_filter_clause(hours: list[int], bind_params: dict[str, object]) -> str:
    if set(hours) == set(range(24)):
        return ""
    return (
        "AND TO_NUMBER(TO_CHAR(v.TOT_DT, 'HH24')) IN "
        f"({build_in_clause('hour', sorted(set(hours)), bind_params)})"
    )


def build_requested_period_clause(
    periods: list[Period],
    bind_params: dict[str, object],
) -> str:
    if len(periods) <= 1:
        return ""

    conditions = []
    for index, period in enumerate(periods):
        start_bind = f"period_start{index}"
        end_bind = f"period_end{index}"
        bind_params[start_bind] = datetime.combine(period.start, datetime.min.time())
        bind_params[end_bind] = datetime.combine(period.end + timedelta(days=1), datetime.min.time())
        conditions.append(f"(v.TOT_DT >= :{start_bind} AND v.TOT_DT < :{end_bind})")

    return "AND (\n              " + "\n           OR ".join(conditions) + "\n          )"


def build_fetch_all_intersections_periods_sql_params(
    periods: list[Period],
    hours: list[int],
    vehicle_kind_names: dict[str, str],
) -> tuple[str, dict[str, object]]:
    if not periods:
        raise RuntimeError("조회 기간이 없습니다.")

    start = min(period.start for period in periods)
    end = max(period.end for period in periods)
    bind_params: dict[str, object] = {
        "start_dt": datetime.combine(start, datetime.min.time()),
        "end_next": datetime.combine(end + timedelta(days=1), datetime.min.time()),
    }
    extra_clauses = [
        clause
        for clause in (
            build_requested_period_clause(periods, bind_params),
            build_hour_filter_clause(hours, bind_params),
        )
        if clause
    ]

    sql = build_vehicle_traffic_sql(
        vehicle_kind_names,
        bind_params,
        start_expression=":start_dt",
        end_expression=":end_next",
        node_filter=False,
        hour_clause="\n          ".join(extra_clauses),
        order_by_expressions=("v.TOT_DT",),
    )
    ensure_select_sql(sql)
    return sql, bind_params


def build_fetch_rows_sql_params(
    intersection: Intersection,
    period: Period,
    hours: list[int],
    vehicle_kind_names: dict[str, str],
) -> tuple[str, dict[str, object]]:
    start_dt = datetime.combine(period.start, datetime.min.time())
    end_next = datetime.combine(period.end + timedelta(days=1), datetime.min.time())
    bind_params: dict[str, object] = {
        "node_id": intersection.node_id,
        "start_dt": start_dt,
        "end_next": end_next,
    }
    hour_clause = build_hour_filter_clause(hours, bind_params)

    sql = build_vehicle_traffic_sql(
        vehicle_kind_names,
        bind_params,
        start_expression=":start_dt",
        end_expression=":end_next",
        node_filter=True,
        hour_clause=hour_clause,
    )
    ensure_select_sql(sql)
    return sql, bind_params


def fetch_rows_for_step(
    conn,
    intersection: Intersection,
    period: Period,
    hours: list[int],
    vehicle_kind_names: dict[str, str],
) -> list[tuple]:
    sql, bind_params = build_fetch_rows_sql_params(
        intersection,
        period,
        hours,
        vehicle_kind_names,
    )
    with conn.cursor() as cur:
        cur.arraysize = 10000
        cur.execute(sql, bind_params)
        return cur.fetchall()


def format_time(value) -> str:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    return str(value)


def normalize_result_row(row: tuple) -> list[object]:
    tot_dt, crsrd_nm, acsr_nm, drct_cd, drct_nm, vknd_cd, vknd_nm, trf_qnty = row
    return [
        format_time(tot_dt),
        crsrd_nm,
        acsr_nm or "",
        drct_nm or normalize_code(drct_cd),
        vknd_nm or normalize_code(vknd_cd),
        int(trf_qnty or 0),
    ]


def result_sort_key(row: list[object]) -> tuple[object, object, object, object, object]:
    return (row[0], row[1], row[2], row[3], row[4])


def count_period_days(periods: Iterable[Period]) -> int:
    return sum((period.end - period.start).days + 1 for period in periods)


def iter_period_dates(periods: Iterable[Period]) -> Iterator[date]:
    for period in periods:
        current = period.start
        while current <= period.end:
            yield current
            current += timedelta(days=1)


def format_period_token(start: date, end: date) -> str:
    start_text = start.strftime("%y%m%d")
    end_text = end.strftime("%y%m%d")
    if start == end:
        return start_text
    return f"{start_text}~{end_text}"


def compact_dates_to_periods(dates: Iterable[date]) -> list[Period]:
    ordered_dates = list(dates)
    if not ordered_dates:
        return []

    periods: list[Period] = []
    start = ordered_dates[0]
    previous = ordered_dates[0]
    for current in ordered_dates[1:]:
        if current == previous + timedelta(days=1):
            previous = current
            continue
        periods.append(Period(start=start, end=previous, token=format_period_token(start, previous)))
        start = current
        previous = current

    periods.append(Period(start=start, end=previous, token=format_period_token(start, previous)))
    return periods


def split_periods_for_parallel_workers(periods: list[Period]) -> list[list[Period]]:
    ordered_dates = list(iter_period_dates(periods))
    if not ordered_dates:
        return []

    worker_count = min(DEFAULT_PARALLEL_WORKERS, len(ordered_dates))
    base = len(ordered_dates) // worker_count
    remainder = len(ordered_dates) % worker_count
    chunk_sizes = [
        base + (1 if index >= worker_count - remainder else 0)
        for index in range(worker_count)
    ]
    chunks: list[list[Period]] = []
    offset = 0
    for size in chunk_sizes:
        if size <= 0:
            continue
        chunk_dates = ordered_dates[offset:offset + size]
        chunks.append(compact_dates_to_periods(chunk_dates))
        offset += size
    return chunks


def split_periods_by_month(periods: list[Period]) -> list[list[Period]]:
    chunks: list[list[Period]] = []
    current_month: tuple[int, int] | None = None
    current_dates: list[date] = []

    for current_date in iter_period_dates(periods):
        month_key = (current_date.year, current_date.month)
        if current_month is not None and month_key != current_month:
            chunks.append(compact_dates_to_periods(current_dates))
            current_dates = []
        current_month = month_key
        current_dates.append(current_date)

    if current_dates:
        chunks.append(compact_dates_to_periods(current_dates))
    return chunks


def format_month_chunk_label(periods: list[Period]) -> str:
    if not periods:
        return "-"
    first = periods[0].start
    last = periods[-1].end
    if (first.year, first.month) == (last.year, last.month):
        return first.strftime("%y%m")
    return f"{first.strftime('%y%m')}~{last.strftime('%y%m')}"


def fetch_rows_for_periods(
    conn,
    intersections: list[Intersection],
    periods: list[Period],
    hours: list[int],
    vehicle_kind_names: dict[str, str],
) -> list[list]:
    rows: list[list] = []
    for intersection in intersections:
        for period in periods:
            fetched = fetch_rows_for_step(conn, intersection, period, hours, vehicle_kind_names)
            rows.extend(normalize_result_row(row) for row in fetched)
    rows.sort(key=result_sort_key)
    return rows


def fetch_all_rows(
    conn,
    intersections: list[Intersection],
    periods: list[Period],
    hours: list[int],
    vehicle_kind_names: dict[str, str],
) -> list[list]:
    progress = StepProgress("[3/4] 교통량 조회", total=len(intersections) * len(periods))
    progress.start()
    rows: list[list] = []
    try:
        for intersection in intersections:
            for period in periods:
                fetched = fetch_rows_for_step(conn, intersection, period, hours, vehicle_kind_names)
                rows.extend(normalize_result_row(row) for row in fetched)
                progress.advance()
    finally:
        progress.finish()
    rows.sort(key=result_sort_key)
    return rows


def save_csv(rows: list[list], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(CSV_HEADER)
        writer.writerows(rows)


def normalize_chunk_csv_row(row: list[str], chunk_path: Path) -> list[object]:
    if len(row) != len(CSV_HEADER):
        raise RuntimeError(
            f"청크 CSV 컬럼 수가 올바르지 않습니다: {chunk_path} "
            f"({len(row)}개, 기대 {len(CSV_HEADER)}개)"
        )
    normalized: list[object] = row.copy()
    normalized[5] = int(normalized[5] or 0)
    return normalized


def read_chunk_rows(chunk_path: Path) -> list[list]:
    with chunk_path.open("r", newline="", encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        try:
            header = next(reader)
        except StopIteration:
            return []
        if header != CSV_HEADER:
            raise RuntimeError(f"청크 CSV 헤더가 올바르지 않습니다: {chunk_path}")
        return [normalize_chunk_csv_row(row, chunk_path) for row in reader]


def merge_chunk_files(chunk_paths: Iterable[Path]) -> list[list]:
    rows: list[list] = []
    for chunk_path in chunk_paths:
        if chunk_path.suffix.lower() == ".db":
            rows.extend(read_output_rows_from_sqlite(chunk_path))
        else:
            rows.extend(read_chunk_rows(chunk_path))
    rows.sort(key=result_sort_key)
    return rows


def build_worker_progress_states(
    period_chunks: list[list[Period]],
) -> dict[int, WorkerProgressState]:
    return {
        index: WorkerProgressState(
            worker_index=index,
            assigned_days=count_period_days(chunk),
        )
        for index, chunk in enumerate(period_chunks, start=1)
    }


def apply_worker_progress_event(
    states: dict[int, WorkerProgressState],
    event: dict[str, object],
) -> None:
    worker_index = int(event["worker_index"])
    state = states[worker_index]
    event_type = str(event.get("type", ""))

    if event_type == "worker_started":
        state.status = "pending"
        state.assigned_days = int(event.get("assigned_days", state.assigned_days))
    elif event_type == "chunk_started":
        state.status = "running"
        state.current_label = str(event.get("label", ""))
    elif event_type == "day_progress":
        state.status = "running"
        state.completed_days = min(
            state.assigned_days,
            max(
                state.completed_days,
                int(event.get("completed_days", state.completed_days)),
            ),
        )
    elif event_type == "chunk_completed":
        state.status = "merge_pending"
        state.current_label = ""
        state.last_completed_label = str(event.get("label", ""))
        state.completed_days = min(
            state.assigned_days,
            int(event.get("completed_days", state.completed_days)),
        )
        state.merge_total_rows = int(event.get("row_count", 0))
        state.merge_completed_rows = 0
    elif event_type == "chunk_merge_started":
        state.status = "merging"
        state.current_label = str(event.get("label", state.last_completed_label))
        state.merge_total_rows = int(event.get("expected_rows", state.merge_total_rows))
        state.merge_completed_rows = 0
    elif event_type == "chunk_merge_progress":
        state.status = "merging"
        state.merge_total_rows = int(event.get("expected_rows", state.merge_total_rows))
        state.merge_completed_rows = int(
            event.get("merged_rows", state.merge_completed_rows)
        )
        state.row_count = int(event.get("row_count", state.row_count))
    elif event_type == "chunk_merge_completed":
        state.status = "completed_chunk"
        state.current_label = ""
        state.last_completed_label = str(event.get("label", state.last_completed_label))
        state.merge_total_rows = int(event.get("merged_rows", state.merge_total_rows))
        state.merge_completed_rows = state.merge_total_rows
        state.row_count = int(event.get("row_count", state.row_count))
    elif event_type == "worker_completed":
        state.status = "completed"
        state.current_label = ""
        state.completed_days = state.assigned_days
    elif event_type == "worker_failed":
        state.status = "failed"
        state.current_label = ""
        state.error_message = str(event.get("error", ""))


def format_progress_bar(done: int, total: int, width: int = 20) -> tuple[str, int]:
    if total <= 0:
        ratio = 1.0
    else:
        ratio = max(0.0, min(1.0, done / total))
    filled = min(width, int(round(width * ratio)))
    percent = int(round(ratio * 100))
    return f"[{'#' * filled}{'-' * (width - filled)}]", percent


def format_worker_progress_line(state: WorkerProgressState, frame: str) -> str:
    bar, percent = format_progress_bar(state.completed_days, state.assigned_days)
    if state.status == "running":
        marker = frame
        status_text = f"진행중: {state.current_label or '-'}"
    elif state.status == "merge_pending":
        marker = "-"
        status_text = (
            f"병합 대기: {state.last_completed_label or '-'} "
            f"({state.merge_total_rows:,} rows)"
        )
    elif state.status == "merging":
        marker = frame
        label = state.current_label or state.last_completed_label or "-"
        total = state.merge_total_rows
        status_text = f"병합중: {label} {state.merge_completed_rows:,}/{total:,} rows"
    elif state.status == "completed":
        marker = "-"
        status_text = "완료"
    elif state.status == "failed":
        marker = "!"
        message = state.error_message.splitlines()[0] if state.error_message else "-"
        status_text = f"실패: {message[:48]}"
    elif state.last_completed_label:
        marker = "-"
        status_text = f"완료: {state.last_completed_label}"
    else:
        marker = "-"
        status_text = "대기"

    return (
        f"W{state.worker_index} {marker} {bar} {percent:3d}% "
        f"{state.completed_days}/{state.assigned_days}일 "
        f"{status_text}  rows={state.row_count:,}"
    )


def format_elapsed(seconds: float) -> str:
    total_seconds = max(0, int(seconds))
    minutes, second = divmod(total_seconds, 60)
    hour, minute = divmod(minutes, 60)
    if hour:
        return f"{hour:02d}:{minute:02d}:{second:02d}"
    return f"{minute:02d}:{second:02d}"


class ParallelProgressBoard:
    def __init__(
        self,
        label: str,
        worker_states: dict[int, WorkerProgressState],
        total_chunks: int,
    ) -> None:
        self.label = label
        self.worker_states = worker_states
        self.total_chunks = total_chunks
        self.started = time.perf_counter()
        self._frames = ("/", "-", "\\", "|")
        self._frame_index = 0
        self._last_line_count = 0
        self._last_rendered = 0.0
        self._last_printed_chunk_count = -1
        self._interactive = sys.stdout.isatty()

    def render(self, completed_chunks: int, merged_rows: int, force: bool = False) -> None:
        now = time.perf_counter()
        if not force and now - self._last_rendered < 0.12:
            return
        if not self._interactive and not force and completed_chunks == self._last_printed_chunk_count:
            return

        frame = self._frames[self._frame_index % len(self._frames)]
        self._frame_index += 1
        lines = [
            self.label,
            "",
            *[
                format_worker_progress_line(self.worker_states[index], frame)
                for index in sorted(self.worker_states)
            ],
            "",
            (
                f"완료 chunk: {completed_chunks}/{self.total_chunks}개 | "
                f"병합 rows: {merged_rows:,} | "
                f"경과 {format_elapsed(now - self.started)}"
            ),
        ]

        if self._interactive:
            if self._last_line_count:
                print(f"\x1b[{self._last_line_count}F", end="")
            for line in lines:
                print(f"\x1b[2K{line}")
        else:
            print("\n".join(lines), flush=True)

        self._last_line_count = len(lines)
        self._last_rendered = now
        self._last_printed_chunk_count = completed_chunks

    def finish(self, completed_chunks: int, merged_rows: int) -> None:
        self.render(completed_chunks, merged_rows, force=True)
        if self._interactive:
            print()


def emit_worker_event(event_queue: Any, event: dict[str, object]) -> None:
    if event_queue is not None:
        event_queue.put(event)


def fetch_worker_chunk(
    worker_index: int,
    intersections: list[Intersection],
    periods: list[Period],
    hours: list[int],
    vehicle_kind_names: dict[str, str],
    is_all_intersections: bool,
    chunk_dir: str,
    event_queue: Any,
) -> WorkerChunkResult:
    month_chunks = split_periods_by_month(periods)
    assigned_days = count_period_days(periods)
    emit_worker_event(
        event_queue,
        {
            "type": "worker_started",
            "worker_index": worker_index,
            "assigned_days": assigned_days,
            "chunk_count": len(month_chunks),
        },
    )

    total_rows = 0
    completed_days = 0
    completed_chunks = 0
    try:
        for month_index, month_periods in enumerate(month_chunks, start=1):
            month_label = format_month_chunk_label(month_periods)
            chunk_path = (
                Path(chunk_dir)
                / f"worker_{worker_index:02d}_month_{month_index:02d}_"
                f"{make_safe_filename_part(month_label)}.db"
            )
            emit_worker_event(
                event_queue,
                {
                    "type": "chunk_started",
                    "worker_index": worker_index,
                    "label": month_label,
                },
            )

            conn = connect_db()
            try:
                def emit_day_progress(chunk_completed_days: int) -> None:
                    emit_worker_event(
                        event_queue,
                        {
                            "type": "day_progress",
                            "worker_index": worker_index,
                            "completed_days": completed_days + chunk_completed_days,
                        },
                    )

                row_count = write_oracle_periods_to_sqlite_chunk(
                    conn,
                    chunk_path,
                    intersections,
                    month_periods,
                    hours,
                    vehicle_kind_names,
                    is_all_intersections=is_all_intersections,
                    day_progress_callback=emit_day_progress if is_all_intersections else None,
                )
            finally:
                conn.close()

            completed_chunks += 1
            chunk_days = count_period_days(month_periods)
            completed_days += chunk_days
            total_rows += row_count
            emit_worker_event(
                event_queue,
                {
                    "type": "chunk_completed",
                    "worker_index": worker_index,
                    "label": month_label,
                    "chunk_path": str(chunk_path),
                    "chunk_days": chunk_days,
                    "completed_days": completed_days,
                    "row_count": row_count,
                },
            )

        emit_worker_event(
            event_queue,
            {"type": "worker_completed", "worker_index": worker_index},
        )
        return WorkerChunkResult(
            worker_index=worker_index,
            chunk_path="",
            row_count=total_rows,
            chunk_count=completed_chunks,
        )
    except Exception as exc:
        emit_worker_event(
            event_queue,
            {
                "type": "worker_failed",
                "worker_index": worker_index,
                "error": str(exc),
            },
        )
        raise


def drain_worker_events(
    event_queue: Any,
    worker_states: dict[int, WorkerProgressState],
    accumulator_conn: sqlite3.Connection,
    batch_size: int,
    progress_callback: Callable[[int, int], None] | None = None,
) -> tuple[int, int, int]:
    event_count = 0
    completed_chunks = 0
    merged_rows = 0

    def notify_progress(
        local_completed_chunks: int | None = None,
        local_merged_rows: int | None = None,
    ) -> None:
        if progress_callback is not None:
            progress_callback(
                completed_chunks if local_completed_chunks is None else local_completed_chunks,
                merged_rows if local_merged_rows is None else local_merged_rows,
            )

    while True:
        try:
            event = event_queue.get_nowait()
        except Empty:
            break

        event_count += 1
        if not isinstance(event, dict):
            continue

        if event.get("type") == "chunk_completed":
            apply_worker_progress_event(worker_states, event)
            notify_progress()

            worker_index = int(event["worker_index"])
            chunk_path = Path(str(event["chunk_path"]))
            expected_count = int(event.get("row_count", 0))
            base_worker_rows = worker_states[worker_index].row_count
            merge_started_event = {
                "type": "chunk_merge_started",
                "worker_index": worker_index,
                "label": str(event.get("label", "")),
                "expected_rows": expected_count,
            }
            apply_worker_progress_event(worker_states, merge_started_event)
            notify_progress()

            def emit_merge_progress(current_chunk_rows: int) -> None:
                apply_worker_progress_event(
                    worker_states,
                    {
                        "type": "chunk_merge_progress",
                        "worker_index": worker_index,
                        "label": str(event.get("label", "")),
                        "expected_rows": expected_count,
                        "merged_rows": current_chunk_rows,
                        "row_count": base_worker_rows + current_chunk_rows,
                    },
                )
                notify_progress(
                    completed_chunks,
                    merged_rows + current_chunk_rows,
                )

            merged_count = merge_sqlite_chunk_file(
                accumulator_conn,
                chunk_path,
                batch_size=batch_size,
                progress_callback=emit_merge_progress,
            )
            if merged_count != expected_count:
                raise RuntimeError(
                    f"SQLite chunk 병합 행 수가 맞지 않습니다: {chunk_path} "
                    f"(chunk={expected_count:,}, merged={merged_count:,})"
                )
            apply_worker_progress_event(
                worker_states,
                {
                    "type": "chunk_merge_completed",
                    "worker_index": worker_index,
                    "label": str(event.get("label", "")),
                    "merged_rows": merged_count,
                    "row_count": base_worker_rows + merged_count,
                },
            )
            try:
                chunk_path.unlink()
            except FileNotFoundError:
                pass
            completed_chunks += 1
            merged_rows += merged_count
            notify_progress()
            continue

        apply_worker_progress_event(worker_states, event)
        notify_progress()

    return event_count, completed_chunks, merged_rows


def terminate_process_pool_workers(executor: ProcessPoolExecutor) -> None:
    processes = getattr(executor, "_processes", None)
    if not processes:
        return
    for process in processes.values():
        if process.is_alive():
            process.terminate()


def fetch_all_rows_parallel_to_sqlite(
    intersections: list[Intersection],
    periods: list[Period],
    hours: list[int],
    vehicle_kind_names: dict[str, str],
    accumulator_path: Path,
    batch_size: int = SQLITE_INSERT_BATCH_SIZE,
    is_all_intersections: bool = False,
) -> int:
    period_chunks = split_periods_for_parallel_workers(periods)
    accumulator_conn = recreate_sqlite_file(accumulator_path)
    if not period_chunks:
        create_sqlite_indexes(
            accumulator_conn,
            progress_callback=print_sqlite_index_progress,
        )
        accumulator_conn.commit()
        accumulator_conn.close()
        return 0

    total_days = count_period_days(periods)
    chunk_day_counts = [count_period_days(chunk) for chunk in period_chunks]
    total_month_chunks = sum(len(split_periods_by_month(chunk)) for chunk in period_chunks)
    print(
        "  병렬 조회 작업: "
        f"총 {total_days}일 -> "
        + " / ".join(f"{day_count}일" for day_count in chunk_day_counts)
    )

    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    worker_states = build_worker_progress_states(period_chunks)
    worker_intersections = [] if is_all_intersections else intersections
    completed_chunk_count = 0
    merged_row_count = 0

    try:
        with TemporaryDirectory(prefix="chunks_", dir=RESULT_DIR) as temp_dir:
            manager = Manager()
            try:
                event_queue = manager.Queue()
                board = ParallelProgressBoard(
                    "[3/4] 교통량 병렬 조회",
                    worker_states,
                    total_month_chunks,
                )
                executor = ProcessPoolExecutor(max_workers=len(period_chunks))
                shutdown_cancel = False
                futures = {
                    executor.submit(
                        fetch_worker_chunk,
                        index,
                        worker_intersections,
                        chunk,
                        hours,
                        vehicle_kind_names,
                        is_all_intersections,
                        temp_dir,
                        event_queue,
                    ): index
                    for index, chunk in enumerate(period_chunks, start=1)
                }
                pending = set(futures)
                board.render(completed_chunk_count, merged_row_count, force=True)
                try:
                    def render_drain_progress(
                        drained_chunk_count: int,
                        drained_row_count: int,
                    ) -> None:
                        board.render(
                            completed_chunk_count + drained_chunk_count,
                            merged_row_count + drained_row_count,
                            force=True,
                        )

                    while pending:
                        event_count, chunk_count, row_count = drain_worker_events(
                            event_queue,
                            worker_states,
                            accumulator_conn,
                            batch_size,
                            progress_callback=render_drain_progress,
                        )
                        completed_chunk_count += chunk_count
                        merged_row_count += row_count
                        if event_count:
                            board.render(completed_chunk_count, merged_row_count, force=True)
                        else:
                            board.render(completed_chunk_count, merged_row_count)

                        done, pending = wait(
                            pending,
                            timeout=0.12,
                            return_when=FIRST_COMPLETED,
                        )
                        for future in done:
                            worker_index = futures[future]
                            try:
                                future.result()
                            except Exception as exc:
                                worker_states[worker_index].status = "failed"
                                worker_states[worker_index].error_message = str(exc)
                                shutdown_cancel = True
                                board.render(completed_chunk_count, merged_row_count, force=True)
                                raise RuntimeError(f"W{worker_index} 조회 실패: {exc}") from exc

                    while True:
                        event_count, chunk_count, row_count = drain_worker_events(
                            event_queue,
                            worker_states,
                            accumulator_conn,
                            batch_size,
                            progress_callback=render_drain_progress,
                        )
                        if not event_count:
                            break
                        completed_chunk_count += chunk_count
                        merged_row_count += row_count
                        board.render(completed_chunk_count, merged_row_count, force=True)
                except Exception:
                    terminate_process_pool_workers(executor)
                    raise
                finally:
                    executor.shutdown(wait=True, cancel_futures=shutdown_cancel)
                    board.finish(completed_chunk_count, merged_row_count)
            finally:
                manager.shutdown()

        create_sqlite_indexes(
            accumulator_conn,
            progress_callback=print_sqlite_index_progress,
        )
        accumulator_conn.commit()
        return merged_row_count
    except Exception:
        accumulator_conn.rollback()
        raise
    finally:
        accumulator_conn.close()


def fetch_all_rows_parallel(
    intersections: list[Intersection],
    periods: list[Period],
    hours: list[int],
    vehicle_kind_names: dict[str, str],
    is_all_intersections: bool = False,
) -> list[list]:
    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    with TemporaryDirectory(prefix="accumulator_", dir=RESULT_DIR) as temp_dir:
        accumulator_path = Path(temp_dir) / "traffic_accumulator.db"
        fetch_all_rows_parallel_to_sqlite(
            intersections,
            periods,
            hours,
            vehicle_kind_names,
            accumulator_path,
            is_all_intersections=is_all_intersections,
        )
        return read_output_rows_from_sqlite(accumulator_path)


def save_xlsx(rows: list[list], output_path: Path) -> None:
    try:
        from openpyxl import Workbook
    except ImportError as exc:  # pragma: no cover - 실행 환경 안내용
        raise RuntimeError(
            "XLSX 저장에는 openpyxl 패키지가 필요합니다. "
            "설치 후 다시 실행해 주세요: python -m pip install openpyxl"
        ) from exc

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "교통량"
    ws.append(CSV_HEADER)
    for row in rows:
        ws.append(row)
    wb.save(output_path)


def recreate_sqlite_table(conn: sqlite3.Connection) -> None:
    conn.executescript(
        f"""
        DROP TABLE IF EXISTS {SQLITE_TABLE};

        CREATE TABLE {SQLITE_TABLE} (
            observed_at TEXT NOT NULL,
            intersection TEXT NOT NULL,
            approach_direction TEXT NOT NULL,
            direction TEXT NOT NULL,
            vehicle_type TEXT NOT NULL,
            traffic_volume INTEGER NOT NULL,
            vehicle_code TEXT NOT NULL
        );
        """
    )


def create_sqlite_indexes(
    conn: sqlite3.Connection,
    progress_callback: Callable[[int, int, str], None] | None = None,
) -> None:
    total = len(SQLITE_INDEX_DEFINITIONS)
    for index, (index_name, sql) in enumerate(SQLITE_INDEX_DEFINITIONS, start=1):
        conn.execute(sql)
        if progress_callback is not None:
            progress_callback(index, total, index_name)


def print_sqlite_index_progress(current: int, total: int, index_name: str) -> None:
    print(f"  인덱스 생성 {current}/{total}: {index_name}", flush=True)


def to_sqlite_record(row: list[object]) -> tuple[str, str, str, str, str, int, str]:
    return (
        str(row[0]),
        str(row[1]),
        str(row[2]),
        str(row[3]),
        str(row[4]),
        int(row[5] or 0),
        str(row[6]) if len(row) > 6 else "",
    )


def normalize_sqlite_record(row: tuple) -> tuple[str, str, str, str, str, int, str]:
    tot_dt, crsrd_nm, acsr_nm, drct_cd, drct_nm, vknd_cd, vknd_nm, trf_qnty = row
    vehicle_code = normalize_code(vknd_cd)
    return (
        format_time(tot_dt),
        str(crsrd_nm),
        str(acsr_nm or ""),
        str(drct_nm or normalize_code(drct_cd)),
        str(vknd_nm or vehicle_code),
        int(trf_qnty or 0),
        vehicle_code,
    )


def observed_date_from_value(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    match = re.match(r"(\d{4})-(\d{2})-(\d{2})", str(value).strip())
    if not match:
        raise RuntimeError(f"조회 row의 날짜를 해석할 수 없습니다: {value}")
    return date(int(match.group(1)), int(match.group(2)), int(match.group(3)))


def insert_sqlite_records(
    conn: sqlite3.Connection,
    records: Iterable[tuple[str, str, str, str, str, int, str]],
    batch_size: int = SQLITE_INSERT_BATCH_SIZE,
    progress_callback: Callable[[int], None] | None = None,
) -> int:
    insert_sql = f"""
        INSERT INTO {SQLITE_TABLE} (
            observed_at,
            intersection,
            approach_direction,
            direction,
            vehicle_type,
            traffic_volume,
            vehicle_code
        )
        VALUES (?, ?, ?, ?, ?, ?, ?)
    """
    inserted = 0
    batch: list[tuple[str, str, str, str, str, int, str]] = []
    for record in records:
        batch.append(record)
        if len(batch) >= batch_size:
            conn.executemany(insert_sql, batch)
            inserted += len(batch)
            batch.clear()
            if progress_callback is not None:
                progress_callback(inserted)

    if batch:
        conn.executemany(insert_sql, batch)
        inserted += len(batch)
        if progress_callback is not None:
            progress_callback(inserted)

    return inserted


def recreate_sqlite_file(output_path: Path) -> sqlite3.Connection:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if output_path.exists():
        output_path.unlink()
    conn = sqlite3.connect(output_path)
    recreate_sqlite_table(conn)
    return conn


def save_sqlite(rows: list[list], output_path: Path) -> None:
    conn = recreate_sqlite_file(output_path)
    try:
        insert_sqlite_records(conn, (to_sqlite_record(row) for row in rows))
        create_sqlite_indexes(conn)
        conn.commit()
    finally:
        conn.close()


def iter_sqlite_records_from_cursor(
    cursor: Any,
    batch_size: int,
    day_completed_callback: Callable[[date], None] | None = None,
) -> Iterator[tuple[str, str, str, str, str, int, str]]:
    previous_observed_date: date | None = None
    while True:
        rows = cursor.fetchmany(batch_size)
        if not rows:
            break
        for row in rows:
            if day_completed_callback is not None:
                current_observed_date = observed_date_from_value(row[0])
                if (
                    previous_observed_date is not None
                    and current_observed_date != previous_observed_date
                ):
                    day_completed_callback(previous_observed_date)
                previous_observed_date = current_observed_date
            yield normalize_sqlite_record(row)


def make_chunk_day_progress_callback(
    periods: list[Period],
    day_progress_callback: Callable[[int], None],
) -> tuple[Callable[[date], None], Callable[[], None]]:
    requested_dates = list(iter_period_dates(periods))
    requested_date_set = set(requested_dates)
    completed_dates: set[date] = set()

    def mark_completed(observed_date: date) -> None:
        if observed_date not in requested_date_set or observed_date in completed_dates:
            return
        completed_dates.add(observed_date)
        day_progress_callback(len(completed_dates))

    def finish_chunk() -> None:
        day_progress_callback(count_period_days(periods))

    return mark_completed, finish_chunk


def write_oracle_periods_to_sqlite_chunk(
    oracle_conn: Any,
    chunk_path: Path,
    intersections: list[Intersection],
    periods: list[Period],
    hours: list[int],
    vehicle_kind_names: dict[str, str],
    batch_size: int = SQLITE_INSERT_BATCH_SIZE,
    is_all_intersections: bool = False,
    day_progress_callback: Callable[[int], None] | None = None,
) -> int:
    sqlite_conn = recreate_sqlite_file(chunk_path)
    total_rows = 0
    try:
        if is_all_intersections:
            day_completed_callback = None
            finish_day_progress = None
            if day_progress_callback is not None:
                day_completed_callback, finish_day_progress = make_chunk_day_progress_callback(
                    periods,
                    day_progress_callback,
                )

            sql, bind_params = build_fetch_all_intersections_periods_sql_params(
                periods,
                hours,
                vehicle_kind_names,
            )
            with oracle_conn.cursor() as cur:
                cur.arraysize = batch_size
                cur.execute(sql, bind_params)
                total_rows += insert_sqlite_records(
                    sqlite_conn,
                    iter_sqlite_records_from_cursor(
                        cur,
                        batch_size,
                        day_completed_callback=day_completed_callback,
                    ),
                    batch_size=batch_size,
                )
            if finish_day_progress is not None:
                finish_day_progress()
        else:
            for intersection in intersections:
                for period in periods:
                    sql, bind_params = build_fetch_rows_sql_params(
                        intersection,
                        period,
                        hours,
                        vehicle_kind_names,
                    )
                    with oracle_conn.cursor() as cur:
                        cur.arraysize = batch_size
                        cur.execute(sql, bind_params)
                        total_rows += insert_sqlite_records(
                            sqlite_conn,
                            iter_sqlite_records_from_cursor(cur, batch_size),
                            batch_size=batch_size,
                        )
        sqlite_conn.commit()
        return total_rows
    except Exception:
        sqlite_conn.rollback()
        raise
    finally:
        sqlite_conn.close()


def iter_stored_sqlite_records(
    source_path: Path,
    batch_size: int = SQLITE_INSERT_BATCH_SIZE,
) -> Iterator[tuple[str, str, str, str, str, int, str]]:
    conn = sqlite3.connect(source_path)
    try:
        cursor = conn.execute(
            f"""
            SELECT
                observed_at,
                intersection,
                approach_direction,
                direction,
                vehicle_type,
                traffic_volume,
                vehicle_code
            FROM {SQLITE_TABLE}
            """
        )
        while True:
            rows = cursor.fetchmany(batch_size)
            if not rows:
                break
            for row in rows:
                yield (
                    str(row[0]),
                    str(row[1]),
                    str(row[2]),
                    str(row[3]),
                    str(row[4]),
                    int(row[5] or 0),
                    str(row[6] or ""),
                )
    finally:
        conn.close()


def merge_sqlite_chunk_file(
    accumulator_conn: sqlite3.Connection,
    chunk_path: Path,
    batch_size: int = SQLITE_INSERT_BATCH_SIZE,
    progress_callback: Callable[[int], None] | None = None,
) -> int:
    row_count = insert_sqlite_records(
        accumulator_conn,
        iter_stored_sqlite_records(chunk_path, batch_size=batch_size),
        batch_size=batch_size,
        progress_callback=progress_callback,
    )
    accumulator_conn.commit()
    return row_count


def iter_output_rows_from_sqlite(
    source_path: Path,
    batch_size: int = SQLITE_INSERT_BATCH_SIZE,
) -> Iterator[list[object]]:
    conn = sqlite3.connect(source_path)
    try:
        cursor = conn.execute(
            f"""
            SELECT
                observed_at,
                intersection,
                approach_direction,
                direction,
                vehicle_type,
                traffic_volume
            FROM {SQLITE_TABLE}
            ORDER BY
                observed_at,
                intersection,
                approach_direction,
                direction,
                vehicle_type
            """
        )
        while True:
            rows = cursor.fetchmany(batch_size)
            if not rows:
                break
            for row in rows:
                yield [
                    str(row[0]),
                    str(row[1]),
                    str(row[2]),
                    str(row[3]),
                    str(row[4]),
                    int(row[5] or 0),
                ]
    finally:
        conn.close()


def read_output_rows_from_sqlite(source_path: Path) -> list[list]:
    return list(iter_output_rows_from_sqlite(source_path))


def export_sqlite_to_csv(
    source_path: Path,
    output_path: Path,
    batch_size: int = SQLITE_INSERT_BATCH_SIZE,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(CSV_HEADER)
        for row in iter_output_rows_from_sqlite(source_path, batch_size=batch_size):
            writer.writerow(row)


def export_sqlite_to_xlsx(
    source_path: Path,
    output_path: Path,
    batch_size: int = SQLITE_INSERT_BATCH_SIZE,
) -> None:
    try:
        from openpyxl import Workbook
    except ImportError as exc:  # pragma: no cover - 실행 환경 안내용
        raise RuntimeError(
            "XLSX 저장에는 openpyxl 패키지가 필요합니다. "
            "설치 후 다시 실행해 주세요: python -m pip install openpyxl"
        ) from exc

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("교통량")
    sheet.append(CSV_HEADER)
    for row in iter_output_rows_from_sqlite(source_path, batch_size=batch_size):
        sheet.append(row)
    workbook.save(output_path)


def copy_sqlite_database(source_path: Path, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if output_path.exists():
        output_path.unlink()
    source_conn = sqlite3.connect(source_path)
    output_conn = sqlite3.connect(output_path)
    try:
        source_conn.backup(output_conn)
        output_conn.commit()
    finally:
        output_conn.close()
        source_conn.close()


def make_temporary_sqlite_output_path(output_path: Path) -> Path:
    stamp = int(time.time() * 1000)
    return output_path.with_name(f".{output_path.name}.{os.getpid()}.{stamp}.tmp")


def replace_sqlite_database(source_path: Path, output_path: Path) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    source_path.replace(output_path)
    return output_path


def remove_file_if_exists(path: Path) -> None:
    try:
        path.unlink()
    except FileNotFoundError:
        pass


def save_accumulator_output(
    accumulator_path: Path,
    output_path: Path,
    output_format: str,
    batch_size: int = SQLITE_INSERT_BATCH_SIZE,
) -> None:
    if output_format == "csv":
        export_sqlite_to_csv(accumulator_path, output_path, batch_size=batch_size)
    elif output_format == "xlsx":
        export_sqlite_to_xlsx(accumulator_path, output_path, batch_size=batch_size)
    elif output_format == "db":
        replace_sqlite_database(accumulator_path, output_path)
    else:
        raise ValueError(f"지원하지 않는 출력 방식입니다: {output_format}")


def save_rows(rows: list[list], output_path: Path, output_format: str) -> None:
    if output_format == "csv":
        save_csv(rows, output_path)
    elif output_format == "xlsx":
        save_xlsx(rows, output_path)
    elif output_format == "db":
        save_sqlite(rows, output_path)
    else:
        raise ValueError(f"지원하지 않는 출력 방식입니다: {output_format}")


def stream_vehicle_traffic_to_sqlite(
    oracle_conn: Any,
    output_path: Path,
    vehicle_kind_names: dict[str, str],
    batch_size: int = SQLITE_INSERT_BATCH_SIZE,
) -> int:
    sql, bind_params = build_all_intersections_vehicle_sql(vehicle_kind_names)
    ensure_select_sql(sql)

    sqlite_conn = recreate_sqlite_file(output_path)
    try:
        with oracle_conn.cursor() as cur:
            cur.arraysize = batch_size
            cur.execute(sql, bind_params)
            row_count = insert_sqlite_records(
                sqlite_conn,
                iter_sqlite_records_from_cursor(cur, batch_size),
                batch_size=batch_size,
            )
        create_sqlite_indexes(
            sqlite_conn,
            progress_callback=print_sqlite_index_progress,
        )
        sqlite_conn.commit()
        return row_count
    except Exception:
        sqlite_conn.rollback()
        raise
    finally:
        sqlite_conn.close()


def validate_sqlite_output(output_path: Path) -> dict[str, object]:
    conn = sqlite3.connect(output_path)
    try:
        row = conn.execute(
            f"""
            SELECT
                COUNT(*) AS row_count,
                MIN(observed_at) AS min_observed_at,
                MAX(observed_at) AS max_observed_at,
                SUM(CASE WHEN vehicle_code IN ('0', '1') THEN 1 ELSE 0 END) AS excluded_code_rows,
                SUM(CASE WHEN traffic_volume <= 0 THEN 1 ELSE 0 END) AS non_positive_rows
            FROM {SQLITE_TABLE}
            """
        ).fetchone()
    finally:
        conn.close()

    summary = {
        "row_count": int(row[0] or 0),
        "min_observed_at": row[1],
        "max_observed_at": row[2],
        "excluded_code_rows": int(row[3] or 0),
        "non_positive_rows": int(row[4] or 0),
    }
    if summary["excluded_code_rows"]:
        raise RuntimeError("SQLite 검증 실패: vehicle_code 0 또는 1 행이 있습니다.")
    if summary["non_positive_rows"]:
        raise RuntimeError("SQLite 검증 실패: 교통량이 0 이하인 행이 있습니다.")
    if summary["min_observed_at"] and summary["min_observed_at"] < FIXED_START_TEXT:
        raise RuntimeError("SQLite 검증 실패: 시작 시간 이전 행이 있습니다.")
    if summary["max_observed_at"] and summary["max_observed_at"] >= FIXED_END_TEXT:
        raise RuntimeError("SQLite 검증 실패: 종료 시간 이후 행이 있습니다.")
    return summary


def export_all_intersections_vehicle_sqlite(
    output_path: Path = FIXED_SQLITE_OUTPUT_PATH,
    batch_size: int = SQLITE_INSERT_BATCH_SIZE,
) -> dict[str, object]:
    started = time.perf_counter()
    print("=" * 60)
    print("  전체교차로 차종별 1시간 교통량 SQLite 추출")
    print("=" * 60)

    vehicle_kind_names = run_instant_step(
        "[1/5] 차종 JSON 로드",
        lambda: load_vehicle_kind_names(),
    )

    try:
        conn = connect_db()
    except Exception as exc:
        raise RuntimeError(f"DB 연결 실패: {exc}") from exc

    temp_output_path = make_temporary_sqlite_output_path(output_path)
    try:
        try:
            with conn.cursor() as cur:
                run_instant_step(
                    "[2/5] 필수 테이블/컬럼 검증",
                    lambda: validate_vehicle_extract_tables(cur),
                )

            row_count = run_instant_step(
                "[3/5] Oracle 조회 및 SQLite 배치 저장",
                lambda: stream_vehicle_traffic_to_sqlite(
                    conn,
                    temp_output_path,
                    vehicle_kind_names,
                    batch_size=batch_size,
                ),
            )
        finally:
            conn.close()

        summary = run_instant_step(
            "[4/5] SQLite 결과 검증",
            lambda: validate_sqlite_output(temp_output_path),
        )
        run_instant_step(
            "[5/5] 최종 DB 교체",
            lambda: replace_sqlite_database(temp_output_path, output_path),
        )
        elapsed = time.perf_counter() - started
        summary["row_count"] = row_count
        summary["output_path"] = str(output_path.resolve())
        summary["elapsed_seconds"] = round(elapsed, 1)
        return summary
    finally:
        remove_file_if_exists(temp_output_path)


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="차종별 교차로 교통량 추출")
    parser.add_argument(
        "--all-intersections-2505-2605-db",
        action="store_true",
        help="2025-05-01 이상 2026-06-01 미만 전체교차로 차종별 1시간 교통량을 SQLite DB로 저장합니다.",
    )
    parser.add_argument(
        "--output",
        default=None,
        help="--all-intersections-2505-2605-db 사용 시 SQLite 저장 경로를 지정합니다.",
    )
    parser.add_argument(
        "--batch-size",
        type=int,
        default=SQLITE_INSERT_BATCH_SIZE,
        help=f"SQLite 배치 insert 크기입니다. 기본값: {SQLITE_INSERT_BATCH_SIZE}",
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    if args.all_intersections_2505_2605_db:
        output_path = Path(args.output) if args.output else FIXED_SQLITE_OUTPUT_PATH
        try:
            summary = export_all_intersections_vehicle_sqlite(
                output_path,
                batch_size=max(args.batch_size, 1),
            )
        except Exception as exc:
            print(f"\n실행을 중단합니다: {exc}", file=sys.stderr)
            return 1

        print(f"\n저장 완료: {summary['output_path']}")
        print(f"레코드 수: {summary['row_count']:,}건")
        print(f"시간 범위: {summary['min_observed_at']} ~ {summary['max_observed_at']}")
        print(f"소요 시간: {summary['elapsed_seconds']:.1f}초")
        return 0

    started = time.perf_counter()
    print("=" * 60)
    print("  차종별 교차로 교통량 추출 프로그램")
    print("=" * 60)
    output_format = input_output_format()

    try:
        conn = connect_db()
    except Exception as exc:
        print(f"\nDB 연결 실패: {exc}")
        return 1

    try:
        all_intersections, _drct_names, vknd_names = run_instant_step(
            "[1/4] 코드/교차로 정보 로드",
            lambda: load_reference_data(conn),
        )
    finally:
        conn.close()

    selected_intersections, is_all = input_intersections(all_intersections)
    periods = input_periods()
    hours = input_hours()
    print_step_done("[2/4] 입력값 해석 및 조회 조건 확정")

    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    base_filename = make_filename(
        selected_intersections,
        is_all,
        (period.token for period in periods),
    )
    filename = base_filename + OUTPUT_FORMAT_EXTENSIONS[output_format]
    final_output_path = RESULT_DIR / filename
    save_label = f"[4/4] {OUTPUT_FORMAT_LABELS[output_format]}"

    if output_format == "db":
        temp_output_path = make_temporary_sqlite_output_path(final_output_path)
        try:
            row_count = fetch_all_rows_parallel_to_sqlite(
                selected_intersections,
                periods,
                hours,
                vknd_names,
                temp_output_path,
                is_all_intersections=is_all,
            )
            output_path = run_instant_step(
                save_label,
                lambda: replace_sqlite_database(temp_output_path, final_output_path),
            )
        finally:
            remove_file_if_exists(temp_output_path)
    else:
        with TemporaryDirectory(prefix="accumulator_", dir=RESULT_DIR) as temp_dir:
            accumulator_path = Path(temp_dir) / "traffic_accumulator.db"
            row_count = fetch_all_rows_parallel_to_sqlite(
                selected_intersections,
                periods,
                hours,
                vknd_names,
                accumulator_path,
                is_all_intersections=is_all,
            )

            def _save_work():
                save_accumulator_output(accumulator_path, final_output_path, output_format)
                return final_output_path

            output_path = run_instant_step(save_label, _save_work)

    elapsed = time.perf_counter() - started
    print(f"\n저장 완료: {output_path}")
    print(f"레코드 수: {row_count:,}건")
    print(f"소요 시간: {elapsed:.1f}초")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
