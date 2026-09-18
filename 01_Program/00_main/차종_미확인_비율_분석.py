#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""미확인 차종(0·1) 교통량 비율을 읽기 전용으로 XLSX에 저장한다."""

from __future__ import annotations

import argparse
import os
import re
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

try:
    import oracledb
except ImportError:  # pragma: no cover - 실행 환경 안내용
    oracledb = None

try:
    from dotenv import load_dotenv
except ImportError:  # pragma: no cover - 실행 환경 안내용
    load_dotenv = None


BASE_DIR = Path(__file__).resolve().parents[2]
ENV_PATH = BASE_DIR / "00_Data" / ".env"
RESULT_DIR = BASE_DIR / "02_Result" / "차종_미확인_파악"
# 운영 DB의 1시간 차종별 원천 테이블명이다. (S_CRSRD_VKND_1HH 뷰는 없음)
TRAFFIC_TABLE = "S_CRSRD_VKND_TRF_1HH"
DEFAULT_WORKERS = 1
BLOCKED_SQL = re.compile(
    r"\b(INSERT|UPDATE|DELETE|MERGE|DROP|TRUNCATE|ALTER|CREATE|COMMIT|"
    r"ROLLBACK|GRANT|REVOKE|SET|BEGIN|DECLARE|CALL|EXECUTE)\b",
    re.IGNORECASE,
)


@dataclass(frozen=True)
class Period:
    start: date
    end: date

    @property
    def label(self) -> str:
        if self.start == self.end:
            return f"{self.start:%Y.%m.%d}"
        return f"{self.start:%Y.%m.%d}-{self.end:%Y.%m.%d}"


@dataclass(frozen=True)
class Target:
    intersection: str | None = None
    direction: str | None = None

    @property
    def label(self) -> str:
        if self.intersection is None:
            return "전체 교차로"
        return (
            self.intersection if self.direction is None else f"{self.intersection}-{self.direction}"
        )


@dataclass(frozen=True)
class AnalysisRow:
    period: str
    intersection: str
    direction: str
    approach: str
    unidentified: int
    total: int


def parse_yymmdd(value: str) -> date:
    if not re.fullmatch(r"\d{6}", value):
        raise ValueError(f"날짜 형식이 올바르지 않습니다: {value}")
    return date(2000 + int(value[:2]), int(value[2:4]), int(value[4:6]))


def month_end(year: int, month: int) -> date:
    return date(year + (month == 12), month % 12 + 1, 1) - timedelta(days=1)


def parse_periods(raw: str) -> list[Period]:
    periods: list[Period] = []
    for token in (part.strip() for part in raw.split(",")):
        if not token:
            continue
        compact = re.sub(r"\s+", "", token)
        if re.fullmatch(r"\d{6}~\d{6}", compact):
            start, end = map(parse_yymmdd, compact.split("~"))
        elif re.fullmatch(r"\d{4}", compact):
            year, month = 2000 + int(compact[:2]), int(compact[2:])
            start, end = date(year, month, 1), month_end(year, month)
        elif re.fullmatch(r"\d{6}", compact):
            start = end = parse_yymmdd(compact)
        else:
            raise ValueError(f"기간 형식이 올바르지 않습니다: {token}")
        if start > end:
            raise ValueError(f"시작일이 종료일보다 늦습니다: {token}")
        periods.append(Period(start, end))
    if not periods:
        raise ValueError("기간을 입력해 주세요.")
    return periods


def parse_target(raw: str) -> Target:
    value = raw.strip()
    if value in {"*", "전체", "all", "ALL"}:
        return Target()
    if not value:
        raise ValueError("대상 교차로를 입력해 주세요.")
    if "-" not in value:
        return Target(intersection=value)
    intersection, direction = (part.strip() for part in value.rsplit("-", 1))
    if not intersection or not direction:
        raise ValueError("방향 대상은 교차로명-북 (동향) 형식으로 입력해 주세요.")
    return Target(intersection=intersection, direction=direction)


def parse_hours(raw: str) -> list[int]:
    aliases = {"첨두시": "07~09", "일반시간": "07~10"}
    slots: set[int] = set()
    for token in (part.strip().lower() for part in raw.split(",")):
        token = aliases.get(token, token)
        match = re.fullmatch(r"(\d{1,2})(?::00)?\s*~\s*(\d{1,2})(?::00)?", token)
        if not match:
            raise ValueError(f"시간 형식이 올바르지 않습니다: {token}")
        start, end = map(int, match.groups())
        if not 0 <= start < end <= 24:
            raise ValueError(f"시간 범위가 올바르지 않습니다: {token}")
        slots.update(range(start, end))
    if not slots:
        raise ValueError("시간을 입력해 주세요.")
    return sorted(slots)


def hours_label(hours: list[int]) -> str:
    ranges: list[tuple[int, int]] = []
    for hour in sorted(set(hours)):
        if ranges and hour == ranges[-1][1]:
            ranges[-1] = (ranges[-1][0], hour + 1)
        else:
            ranges.append((hour, hour + 1))
    return ", ".join(f"{start:02d}:00-{end:02d}:00" for start, end in ranges)


def ensure_select_sql(sql: str) -> None:
    normalized = re.sub(r"/\*.*?\*/|--[^\r\n]*", " ", sql, flags=re.DOTALL).strip()
    if not normalized.upper().startswith("SELECT") or ";" in normalized:
        raise RuntimeError("읽기 전용 보호: 단일 SELECT 문만 실행할 수 있습니다.")
    if BLOCKED_SQL.search(normalized):
        raise RuntimeError("읽기 전용 보호: 쓰기·DDL·트랜잭션 SQL은 실행할 수 없습니다.")


def execute_select(cursor: Any, sql: str, params: dict[str, Any] | None = None) -> list[tuple]:
    """가드된 단일 SELECT만 실행한다."""
    ensure_select_sql(sql)
    cursor.execute(sql, params or {})
    return list(cursor.fetchall())


def connect_read_only():
    if oracledb is None or load_dotenv is None:
        raise RuntimeError("oracledb 및 python-dotenv 패키지가 필요합니다.")
    load_dotenv(ENV_PATH)
    keys = ("DB_USER", "DB_PASSWORD", "DB_HOST", "DB_SERVICE")
    if missing := [key for key in keys if not os.getenv(key)]:
        raise RuntimeError(f"DB 접속 환경변수가 없습니다: {', '.join(missing)}")
    connection = oracledb.connect(
        user=os.getenv("DB_USER", "").strip(),
        password=os.getenv("DB_PASSWORD", "").strip(),
        host=os.getenv("DB_HOST", "").strip(),
        port=int(os.getenv("DB_PORT", "1521")),
        service_name=os.getenv("DB_SERVICE", "").strip(),
    )
    with connection.cursor() as cursor:
        cursor.execute("SET TRANSACTION READ ONLY")
    return connection


def build_day_sql(target: Target, hours: list[int]) -> tuple[str, dict[str, Any]]:
    params: dict[str, Any] = {f"hour{index}": hour for index, hour in enumerate(hours)}
    filters = [
        "v.TOT_DT >= :start_dt",
        "v.TOT_DT < :end_dt",
        "TO_NUMBER(TO_CHAR(v.TOT_DT, 'HH24')) IN ("
        + ", ".join(f":hour{i}" for i in range(len(hours)))
        + ")",
    ]
    if target.intersection is not None:
        params["intersection_name"] = target.intersection
        filters.append("c.CRSRD_NM = :intersection_name")
    if target.direction is not None:
        params["direction_name"] = target.direction
        filters.append("a.ACSR_NM = :direction_name")
    code = "TRIM(TO_CHAR(v.VKND_CD))"
    sql = f"""
        SELECT c.CRSRD_NM, NVL(a.ACSR_NM, ''), NVL(d.CD_NM, TRIM(TO_CHAR(v.DRCT_CD))),
               SUM(CASE WHEN {code} IN ('0', '1') THEN NVL(v.TRF_QNTY, 0) ELSE 0 END),
               SUM(NVL(v.TRF_QNTY, 0))
          FROM {TRAFFIC_TABLE} v
          JOIN M_CRSRD_INF c ON c.NODE_ID = v.NODE_ID
          LEFT JOIN M_CRSRD_ACSR_INF a ON a.NODE_ID = v.NODE_ID AND a.ACSR_ID = v.ACSR_ID
          LEFT JOIN M_CD_INF d ON d.GRP_CD = 'DRCT_CD' AND TRIM(TO_CHAR(d.CD)) = TRIM(TO_CHAR(v.DRCT_CD))
         WHERE {" AND ".join(filters)}
         GROUP BY c.CRSRD_NM, a.ACSR_NM, d.CD_NM, TRIM(TO_CHAR(v.DRCT_CD))
    """
    ensure_select_sql(sql)
    return sql, params


def day_chunks(periods: Iterable[Period]) -> list[tuple[Period, date]]:
    return [
        (period, period.start + timedelta(days=offset))
        for period in periods
        for offset in range((period.end - period.start).days + 1)
    ]


def fetch_day(period: Period, day: date, target: Target, hours: list[int]) -> list[AnalysisRow]:
    sql, params = build_day_sql(target, hours)
    params.update(
        start_dt=datetime.combine(day, datetime.min.time()),
        end_dt=datetime.combine(day + timedelta(days=1), datetime.min.time()),
    )
    connection = connect_read_only()
    try:
        with connection.cursor() as cursor:
            return [
                AnalysisRow(period.label, str(a), str(b), str(c), int(d or 0), int(e or 0))
                for a, b, c, d, e in execute_select(cursor, sql, params)
            ]
    finally:
        connection.close()


def combine_rows(rows: Iterable[AnalysisRow]) -> list[AnalysisRow]:
    totals: dict[tuple[str, str, str, str], list[int]] = {}
    for row in rows:
        values = totals.setdefault(
            (row.period, row.intersection, row.direction, row.approach), [0, 0]
        )
        values[0] += row.unidentified
        values[1] += row.total
    return [AnalysisRow(*key, *values) for key, values in sorted(totals.items())]


def fetch_parallel(
    periods: list[Period], target: Target, hours: list[int], workers: int = DEFAULT_WORKERS
) -> list[AnalysisRow]:
    chunks = day_chunks(periods)
    if not chunks:
        return []
    completed: list[AnalysisRow] = []
    with ThreadPoolExecutor(max_workers=max(1, min(workers, len(chunks)))) as executor:
        futures = [executor.submit(fetch_day, period, day, target, hours) for period, day in chunks]
        for future in as_completed(futures):
            completed.extend(future.result())
    return combine_rows(completed)


def save_xlsx(
    rows: list[AnalysisRow], periods: list[Period], target: Target, hours: list[int], path: Path
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "미확인 분석"
    sheet.merge_cells("A1:H1")
    sheet["A1"] = "차종 미확인 비율 분석"
    sheet["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    sheet["A1"].alignment = Alignment(horizontal="center")
    conditions = [
        ("입력 조건", ""),
        ("기간", ", ".join(period.label for period in periods)),
        ("시간대", hours_label(hours)),
        ("대상", target.label),
        ("계산 기준", "미확인=차종코드 0 또는 1의 교통량 합계 / 전체=모든 차종 교통량 합계"),
        ("표시 기준", "전체 교통량이 0이면 0.00%"),
    ]
    for row_no, (label, value) in enumerate(conditions, 3):
        sheet.cell(row_no, 1, label).font = Font(bold=True)
        sheet.cell(row_no, 2, value)
        sheet.merge_cells(start_row=row_no, start_column=2, end_row=row_no, end_column=8)
    header_row = 10
    headers = [
        "기간",
        "시간",
        "교차로",
        "방향",
        "접근로",
        "미확인 교통량",
        "전체 교통량",
        "미확인 교통량 비율(%)",
    ]
    for col, header in enumerate(headers, 1):
        cell = sheet.cell(header_row, col, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="5B9BD5")
        cell.alignment = Alignment(horizontal="center")
    for row_no, row in enumerate(rows, header_row + 1):
        sheet.append(
            [
                row.period,
                hours_label(hours),
                row.intersection,
                row.direction,
                row.approach,
                row.unidentified,
                row.total,
                f"=IF(G{row_no}=0,0,F{row_no}/G{row_no})",
            ]
        )
        sheet.cell(row_no, 6).number_format = "#,##0"
        sheet.cell(row_no, 7).number_format = "#,##0"
        sheet.cell(row_no, 8).number_format = "0.00%"
    end_row = max(header_row + 1, header_row + len(rows))
    table = Table(displayName="UnidentifiedVehicleAnalysis", ref=f"A{header_row}:H{end_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.auto_filter.ref = f"A{header_row}:H{end_row}"
    for column, width in {
        "A": 24,
        "B": 24,
        "C": 24,
        "D": 20,
        "E": 14,
        "F": 18,
        "G": 16,
        "H": 24,
    }.items():
        sheet.column_dimensions[column].width = width
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def choose_worker_count(durations: dict[int, float]) -> int:
    fastest = min(durations.values())
    return min(workers for workers, duration in durations.items() if duration <= fastest * 1.05)


def run_benchmark() -> int:
    connection = connect_read_only()
    try:
        with connection.cursor() as cursor:
            latest = execute_select(cursor, f"SELECT MAX(TRUNC(TOT_DT)) FROM {TRAFFIC_TABLE}")[0][0]
    finally:
        connection.close()
    periods = [Period(latest - timedelta(days=4), latest)]
    durations: dict[int, float] = {}
    for workers in range(1, 6):
        started = time.perf_counter()
        fetch_parallel(periods, Target(), [7, 8], workers)
        durations[workers] = time.perf_counter() - started
        print(f"{workers} worker: {durations[workers]:.2f}초")
    print(f"권장 기본 워커: {choose_worker_count(durations)}")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--benchmark", action="store_true")
    args = parser.parse_args()
    if args.benchmark:
        return run_benchmark()
    try:
        periods = parse_periods(input("기간 (YYMMDD~YYMMDD, YYMM, 단일일): "))
        target = parse_target(input("대상 (*, 교차로명, 교차로명-북 (동향)): "))
        hours = parse_hours(input("시간 (첨두시, 일반시간, 07~09, 07~09, 17~19): "))
        rows = fetch_parallel(periods, target, hours)
        output = RESULT_DIR / f"차종_미확인_비율_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
        save_xlsx(rows, periods, target, hours, output)
        print(f"저장 완료: {output} ({len(rows)}건)")
    except (RuntimeError, ValueError) as exc:
        print(f"오류: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
