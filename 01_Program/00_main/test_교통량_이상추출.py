from __future__ import annotations

import importlib.util
import json
import sys
import urllib.error
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook


MODULE_PATH = Path(__file__).with_name("교통량_이상추출.py")
SPEC = importlib.util.spec_from_file_location("traffic_anomaly_export", MODULE_PATH)
assert SPEC and SPEC.loader
exporter = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = exporter
SPEC.loader.exec_module(exporter)


def notion_page(**overrides):
    values = {
        "이상발생일": "2026-02-01",
        "요청일": "2026-02-02",
        "조치 완료일": "2026-02-28",
        "조치상태": "완료",
        "이상유형": ["통신", "결측"],
        "교차로": "테스트교차로",
        "요청대상": "담당자",
        "원인": "장비 통신 오류",
        "조치사항": "조치함",
        "비고": "확인",
    }
    values.update(overrides)
    properties = {}
    for name in ("이상발생일", "요청일", "조치 완료일"):
        properties[name] = (
            {"type": "date", "date": {"start": values[name]}}
            if values[name]
            else {"type": "date", "date": None}
        )
    properties["조치상태"] = {"type": "status", "status": {"name": values["조치상태"]}}
    properties["이상유형"] = {
        "type": "multi_select",
        "multi_select": [{"name": name} for name in values["이상유형"]],
    }
    for name in ("교차로", "요청대상", "원인", "조치사항", "비고"):
        properties[name] = {"type": "rich_text", "rich_text": [{"plain_text": values[name]}]}
    return {"properties": properties}


@pytest.mark.parametrize(
    ("raw", "expected"),
    [("2507", (2025, 7)), ("2607", (2026, 7)), ("2701", (2027, 1))],
)
def test_parse_year_month_converts_two_digit_year_and_month(raw, expected):
    assert exporter.parse_year_month(raw) == expected


@pytest.mark.parametrize("raw", ["", "문자", "250", "2500", "2513"])
def test_parse_year_month_rejects_invalid_format_or_month(raw):
    with pytest.raises(ValueError):
        exporter.parse_year_month(raw)


def test_select_records_for_period_filters_completed_by_selected_year_and_month():
    records = exporter.extract_records(
        [
            notion_page(이상발생일="2025-07-03", **{"조치 완료일": "2025-07-31"}),
            notion_page(이상발생일="2025-07-03", **{"조치 완료일": "2025-06-30"}),
            notion_page(이상발생일="2025-07-03", **{"조치 완료일": "2026-07-01"}),
        ]
    )
    completed, incomplete = exporter.select_records_for_period(records, 2025, 7)
    assert [item.completed_on for item in completed] == [date(2025, 7, 31)]
    assert incomplete == []


def test_select_records_for_period_includes_all_prior_incomplete_records():
    records = exporter.extract_records(
        [
            notion_page(이상발생일="2024-12-31", 조치상태="미완료", **{"조치 완료일": None}),
            notion_page(이상발생일="2026-07-31", 조치상태="미완료", **{"조치 완료일": None}),
            notion_page(이상발생일="2026-08-01", 조치상태="미완료", **{"조치 완료일": None}),
        ]
    )
    _, incomplete = exporter.select_records_for_period(records, 2026, 7)
    assert [item.occurred_on for item in incomplete] == [
        date(2024, 12, 31),
        date(2026, 7, 31),
    ]


def test_date_properties_preserve_date_only_and_local_datetime_values():
    record = exporter.extract_record(
        notion_page(
            이상발생일="2026-02-01T09:30:45+09:00",
            요청일="2026-02-02",
            **{"조치 완료일": None},
        )
    )

    assert record is not None
    assert record.occurred_on == datetime(2026, 2, 1, 9, 30, 45)
    assert record.occurred_on.tzinfo is None
    assert record.requested_on == date(2026, 2, 2)
    assert record.completed_on is None


def test_select_records_sorts_same_day_datetime_values_by_time():
    records = exporter.extract_records(
        [
            notion_page(
                이상발생일="2026-02-03T15:00:00+09:00",
                조치상태="미완료",
                **{"조치 완료일": None},
            ),
            notion_page(
                이상발생일="2026-02-03T09:00:00+09:00",
                조치상태="미완료",
                **{"조치 완료일": None},
            ),
            notion_page(
                이상발생일="2026-02-03",
                조치상태="미완료",
                **{"조치 완료일": None},
            ),
        ]
    )

    _, incomplete = exporter.select_records_for_period(records, 2026, 2)

    assert [item.occurred_on for item in incomplete] == [
        date(2026, 2, 3),
        datetime(2026, 2, 3, 9),
        datetime(2026, 2, 3, 15),
    ]


def test_run_writes_expected_workbook(tmp_path):
    class FakeQuery:
        def fetch_pages(self):
            return [
                notion_page(이상발생일="2026-02-02"),
                notion_page(
                    이상발생일="2026-01-01",
                    조치상태="미완료",
                    **{"조치 완료일": None},
                ),
            ]

    output_path = exporter.run("2602", FakeQuery(), tmp_path)
    workbook = load_workbook(output_path)
    assert workbook.sheetnames == ["완료", "미완료"]
    assert [cell.value for cell in workbook["완료"][1]] == list(exporter.COMPLETED_HEADERS)
    assert [cell.value for cell in workbook["미완료"][1]] == list(exporter.INCOMPLETE_HEADERS)
    assert workbook["완료"]["A2"].value.date() == date(2026, 2, 2)
    assert workbook["완료"]["E2"].value == "통신, 결측"
    assert workbook["완료"]["H2"].value == "장비 통신 오류"
    assert workbook["완료"]["I2"].value == "조치함"
    assert workbook["완료"]["A2"].number_format == "yyyy-mm-dd"
    assert workbook["미완료"]["A2"].value.date() == date(2026, 1, 1)
    for worksheet in workbook.worksheets:
        assert worksheet.auto_filter.ref is None
        assert worksheet.freeze_panes == "A2"
    assert output_path.name == "2026년 2월 교통량 이상목록.xlsx"


def test_run_writes_date_and_datetime_formats_without_converting_time(tmp_path):
    class FakeQuery:
        def fetch_pages(self):
            return [
                notion_page(
                    이상발생일="2026-02-02T09:15:30+09:00",
                    요청일="2026-02-02",
                    **{"조치 완료일": "2026-02-28T17:45:00+09:00"},
                )
            ]

    output_path = exporter.run("2602", FakeQuery(), tmp_path)
    worksheet = load_workbook(output_path)["완료"]

    assert worksheet["A2"].value == datetime(2026, 2, 2, 9, 15, 30)
    assert worksheet["B2"].value == datetime(2026, 2, 2)
    assert worksheet["C2"].value == datetime(2026, 2, 28, 17, 45)
    assert worksheet["A2"].number_format == "yyyy-mm-dd hh:mm"
    assert worksheet["B2"].number_format == "yyyy-mm-dd"
    assert worksheet["C2"].number_format == "yyyy-mm-dd hh:mm"
    assert worksheet.column_dimensions["A"].width == 18


def test_workbook_uses_sheet_specific_empty_values_and_auto_row_heights(tmp_path):
    class FakeQuery:
        def fetch_pages(self):
            return [
                notion_page(원인="길이가 긴 원인을\n확인하기 위한 텍스트입니다", 비고=""),
                notion_page(
                    이상발생일="2026-01-01",
                    조치상태="미완료",
                    요청대상="없음",
                    원인="",
                    비고="",
                    **{"조치 완료일": None},
                ),
            ]

    output_path = exporter.run("2602", FakeQuery(), tmp_path)
    workbook = load_workbook(output_path)
    completed = workbook["완료"]
    incomplete = workbook["미완료"]

    assert completed["J2"].value is None
    assert incomplete["G2"].value == "-"
    assert incomplete["H2"].value == "-"
    assert incomplete["F2"].value == "-"
    assert "조치 완료일" not in [cell.value for cell in incomplete[1]]
    assert "조치사항" not in [cell.value for cell in incomplete[1]]
    assert completed["A2"].alignment.horizontal == "left"
    assert completed["D2"].alignment.horizontal == "center"
    assert completed["H2"].alignment.horizontal == "left"
    assert incomplete["F2"].alignment.horizontal == "center"
    assert incomplete["G2"].alignment.horizontal == "center"
    assert incomplete["H2"].alignment.horizontal == "center"
    assert completed.column_dimensions["H"].width > 10
    for worksheet in (completed, incomplete):
        for row_index in range(2, worksheet.max_row + 1):
            assert worksheet.row_dimensions[row_index].height is None
    assert completed["H2"].alignment.wrap_text is True


def test_run_rejects_empty_selection(tmp_path):
    class FakeQuery:
        def fetch_pages(self):
            return []

    with pytest.raises(LookupError, match="해당하는"):
        exporter.run("2602", FakeQuery(), tmp_path)


def test_notion_client_collects_multiple_pages():
    bodies = [
        {"data_sources": [{"id": "data-source", "name": "통합 목록"}]},
        {"results": [{"id": "one"}], "has_more": True, "next_cursor": "cursor"},
        {"results": [{"id": "two"}], "has_more": False, "next_cursor": None},
    ]

    class Response:
        def __init__(self, body):
            self.body = body

        def __enter__(self):
            return self

        def __exit__(self, *args):
            return None

        def read(self):
            return json.dumps(self.body).encode()

    urls = []
    payloads = []

    def opener(request, timeout):
        urls.append(request.full_url)
        if request.data:
            payloads.append(json.loads(request.data.decode()))
        return Response(bodies.pop(0))

    client = exporter.NotionDataSourceClient("secret", "database-id", opener)
    assert [page["id"] for page in client.fetch_pages()] == ["one", "two"]
    assert payloads == [{"page_size": 100}, {"page_size": 100, "start_cursor": "cursor"}]
    assert urls[0].endswith("/v1/databases/database-id")


@pytest.mark.parametrize("status_code", [403, 404, 429])
def test_notion_client_reports_http_errors(status_code):
    def opener(*_args, **_kwargs):
        raise urllib.error.HTTPError("url", status_code, "error", {}, None)

    with pytest.raises(exporter.NotionRequestError):
        exporter.NotionDataSourceClient("secret", "data-source", opener).fetch_pages()


def test_notion_client_reports_network_error():
    def opener(*_args, **_kwargs):
        raise urllib.error.URLError("offline")

    with pytest.raises(exporter.NotionRequestError, match="연결"):
        exporter.NotionDataSourceClient("secret", "data-source", opener).fetch_pages()


def test_main_reports_missing_token(monkeypatch, capsys):
    monkeypatch.setattr(exporter, "load_notion_token", lambda: None)
    monkeypatch.setattr("builtins.input", lambda _prompt: "2507")

    assert exporter.main() == exporter.EXIT_TOKEN_ERROR
    assert "NOTION_API_TOKEN" in capsys.readouterr().err
