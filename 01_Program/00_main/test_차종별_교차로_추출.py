from __future__ import annotations

import importlib.util
import csv
import json
import re
import sqlite3
import sys
from datetime import datetime
from pathlib import Path

import pytest


MODULE_PATH = Path(__file__).with_name("차종별_교차로_추출.py")
SPEC = importlib.util.spec_from_file_location("vehicle_intersection_extract", MODULE_PATH)
assert SPEC is not None and SPEC.loader is not None
vehicle = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = vehicle
SPEC.loader.exec_module(vehicle)


def test_load_vehicle_kind_names_keeps_unidentified_codes(tmp_path):
    path = tmp_path / "vehicle-kinds.json"
    path.write_text(
        json.dumps({"0": "미확인", "1": "미확인", "2": "세단"}),
        encoding="utf-8",
    )

    assert vehicle.load_vehicle_kind_names(path) == {
        "0": "미확인",
        "1": "미확인",
        "2": "세단",
    }


@pytest.mark.parametrize(
    "aggregation_unit",
    [vehicle.ONE_HOUR_AGGREGATION, vehicle.FIFTEEN_MINUTE_AGGREGATION],
)
def test_common_traffic_sql_preserves_all_vehicle_codes(aggregation_unit):
    sql, params = vehicle.build_fetch_all_intersections_periods_sql_params(
        vehicle.parse_periods("260501"),
        vehicle.parse_time_range("07~09", aggregation_unit),
        {"0": "미확인", "1": "미확인", "2": "세단"},
        aggregation_unit,
    )
    compact_sql = re.sub(r"\s+", " ", sql)

    assert "TRIM(TO_CHAR(v.VKND_CD)) AS VKND_CD" in compact_sql
    assert "TRIM(TO_CHAR(v.VKND_CD)) IN" not in compact_sql
    assert "0/1" not in params.values()
    assert "ELSE '미확인' END" in compact_sql
    assert "TRIM(TO_CHAR(v.VKND_CD))" in compact_sql.split("GROUP BY")[1]
    assert "HAVING SUM(NVL(v.TRF_QNTY, 0)) > 0" in compact_sql


def test_sqlite_validation_accepts_original_unidentified_codes(tmp_path):
    output_path = tmp_path / "traffic.db"
    conn = vehicle.recreate_sqlite_file(output_path)
    try:
        vehicle.insert_sqlite_records(
            conn,
            [
                ("2025-05-01 00:00:00", "A", "", "직진", "미확인", 3, code)
                for code in ("0", "1", "99")
            ],
        )
        conn.commit()
    finally:
        conn.close()

    summary = vehicle.validate_sqlite_output(output_path)

    assert summary["raw_unidentified_code_rows"] == 2
    assert summary["row_count"] == 3


@pytest.mark.parametrize(
    "aggregation_unit",
    [vehicle.ONE_HOUR_AGGREGATION, vehicle.FIFTEEN_MINUTE_AGGREGATION],
)
def test_generated_query_groups_by_original_code(aggregation_unit):
    # Execute the shared query against local sample tables without an Oracle connection.
    conn = sqlite3.connect(":memory:")
    conn.create_function("TO_CHAR", 1, str)
    conn.create_function(
        "TO_CHAR", 2, lambda value, fmt: value[11:16] if fmt == "HH24:MI" else value[11:13]
    )
    conn.create_function("TO_NUMBER", 1, int)
    conn.create_function("NVL", 2, lambda value, fallback: fallback if value is None else value)
    try:
        conn.executescript(f"""
            CREATE TABLE {aggregation_unit.traffic_table}
                (TOT_DT TEXT, NODE_ID INTEGER, ACSR_ID TEXT, DRCT_CD TEXT, VKND_CD TEXT, TRF_QNTY INTEGER);
            CREATE TABLE M_CRSRD_INF (NODE_ID INTEGER, CRSRD_NM TEXT);
            CREATE TABLE M_CRSRD_ACSR_INF (NODE_ID INTEGER, ACSR_ID TEXT, ACSR_NM TEXT);
            CREATE TABLE M_CD_INF (GRP_CD TEXT, CD TEXT, CD_NM TEXT);
            INSERT INTO M_CRSRD_INF VALUES (7, 'A');
            INSERT INTO M_CRSRD_ACSR_INF VALUES (7, 'N', '북');
            INSERT INTO M_CD_INF VALUES ('DRCT_CD', '1', '직진');
        """)
        conn.executemany(
            f"INSERT INTO {aggregation_unit.traffic_table} VALUES (?, 7, 'N', '1', ?, ?)",
            [
                ("2026-05-01 07:00:00", code, volume)
                for code, volume in [
                    ("0", 3),
                    ("1", 5),
                    ("02", 7),
                    ("03", 11),
                    ("99", 13),
                    ("99", 2),
                    ("100", 17),
                    ("101", 0),
                ]
            ],
        )
        sql, params = vehicle.build_fetch_rows_sql_params(
            vehicle.Intersection(7, "A"),
            vehicle.parse_periods("260501")[0],
            vehicle.parse_time_range("07~09", aggregation_unit),
            {"0": "미확인", "1": "미확인", "02": "세단", "03": "세단", "100": ""},
            aggregation_unit,
        )
        params = {
            key: str(value) if isinstance(value, datetime) else value
            for key, value in params.items()
        }
        rows = conn.execute(sql, params).fetchall()
        assert {row[5]: (row[6], row[7]) for row in rows} == {
            "0": ("미확인", 3),
            "1": ("미확인", 5),
            "02": ("세단", 7),
            "03": ("세단", 11),
            "99": ("미확인", 15),
            "100": ("미확인", 17),
        }
    finally:
        conn.close()


@pytest.mark.parametrize(
    "observed_at, traffic_volume, vehicle_code",
    [
        ("2025-05-01 00:00:00", 0, "2"),
        ("2025-04-30 23:00:00", 3, "2"),
    ],
)
def test_sqlite_validation_keeps_raw_code_volume_and_period_checks(
    tmp_path, observed_at, traffic_volume, vehicle_code
):
    output_path = tmp_path / "traffic.db"
    conn = vehicle.recreate_sqlite_file(output_path)
    try:
        vehicle.insert_sqlite_records(
            conn,
            [(observed_at, "A", "", "직진", "차종", traffic_volume, vehicle_code)],
        )
        conn.commit()
    finally:
        conn.close()

    with pytest.raises(RuntimeError):
        vehicle.validate_sqlite_output(output_path)


def test_parse_time_range_supports_hourly_and_quarter_hour_slots():
    assert vehicle.parse_time_range("07~09") == [420, 480]
    assert vehicle.parse_time_range("07:00~09:00") == [420, 480]
    assert vehicle.parse_time_range("07:15~09:30", vehicle.FIFTEEN_MINUTE_AGGREGATION) == [
        435,
        450,
        465,
        480,
        495,
        510,
        525,
        540,
        555,
    ]
    assert vehicle.parse_time_range("24시간", vehicle.FIFTEEN_MINUTE_AGGREGATION) == list(
        range(0, 24 * 60, 15)
    )


def test_blank_vehicle_names_and_empty_mapping_use_unidentified(tmp_path):
    path = tmp_path / "kinds.json"
    path.write_text(json.dumps({"02": "세단", "99": " ", "100": None}), encoding="utf-8")
    names = vehicle.load_vehicle_kind_names(path)
    assert names == {"02": "세단", "99": "미확인", "100": "미확인"}
    params = {}
    expression = vehicle.build_vehicle_name_case_expression({"99": " "}, "code", params)
    assert params["vknd_case_name0"] == "미확인"
    assert "ELSE '미확인' END" in expression
    sql, _ = vehicle.build_all_intersections_vehicle_sql({})
    assert "'미확인' AS VKND_NM" in sql
    path.write_text("{}", encoding="utf-8")
    assert vehicle.load_vehicle_kind_names(path) == {}


def test_all_output_paths_preserve_codes_names_and_volumes(tmp_path):
    from openpyxl import load_workbook

    source_rows = [
        (datetime(2026, 5, 1, 7), "A", "북", "1", "직진", code, name, volume)
        for code, name, volume in [
            ("0", "미확인", 3),
            ("1", "미확인", 5),
            ("02", "세단", 7),
            ("03", "세단", 11),
            ("99", None, 13),
            ("100", "", 17),
        ]
    ]
    rows = [vehicle.normalize_result_row(row) for row in source_rows]
    records = [vehicle.normalize_sqlite_record(row) for row in source_rows]
    assert [tuple(row) for row in rows] == records
    assert [row[4] for row in rows] == ["미확인", "미확인", "세단", "세단", "미확인", "미확인"]
    db_path = tmp_path / "traffic.db"
    vehicle.save_sqlite(rows, db_path)
    stored = vehicle.read_output_rows_from_sqlite(db_path)
    assert sorted(stored) == sorted(rows)
    assert sum(row[5] for row in stored) == 56
    assert vehicle.validate_sqlite_output(db_path)["row_count"] == 6
    for export in (False, True):
        csv_path = tmp_path / f"traffic-{export}.csv"
        xlsx_path = tmp_path / f"traffic-{export}.xlsx"
        if export:
            vehicle.export_sqlite_to_csv(db_path, csv_path, batch_size=1)
            vehicle.export_sqlite_to_xlsx(db_path, xlsx_path, batch_size=1)
        else:
            vehicle.save_csv(rows, csv_path)
            vehicle.save_xlsx(rows, xlsx_path)
        with csv_path.open(encoding="utf-8-sig", newline="") as stream:
            csv_rows = list(csv.reader(stream))
        assert csv_rows[0] == [
            "시간",
            "교차로",
            "교차로 방향",
            "방향",
            "차종",
            "교통량",
            "차종 코드",
        ]
        assert sorted(csv_rows[1:]) == sorted([[str(value) for value in row] for row in rows])
        workbook = load_workbook(xlsx_path, read_only=True)
        try:
            xlsx_rows = list(workbook.active.values)
            assert xlsx_rows[0] == tuple(csv_rows[0])
            assert sorted(xlsx_rows[1:]) == sorted(records)
        finally:
            workbook.close()


@pytest.mark.parametrize(
    ("raw", "aggregation_unit", "message"),
    [
        ("07:15~09:00", vehicle.ONE_HOUR_AGGREGATION, ":00 경계"),
        ("07:10~09:00", vehicle.FIFTEEN_MINUTE_AGGREGATION, ":15, :30, :45 경계"),
    ],
)
def test_parse_time_range_rejects_unaligned_boundaries(raw, aggregation_unit, message):
    with pytest.raises(ValueError, match=re.escape(message)):
        vehicle.parse_time_range(raw, aggregation_unit)


def test_fifteen_minute_sql_uses_selected_table_and_minute_slots():
    sql, params = vehicle.build_fetch_rows_sql_params(
        vehicle.Intersection(node_id=7, name="A교차로"),
        vehicle.parse_periods("260501")[0],
        vehicle.parse_time_range("07:15~09:30", vehicle.FIFTEEN_MINUTE_AGGREGATION),
        {"2": "세단"},
        vehicle.FIFTEEN_MINUTE_AGGREGATION,
    )
    compact_sql = re.sub(r"\s+", " ", sql)

    assert "FROM S_CRSRD_VKND_TRF_15MI v" in compact_sql
    assert "TRIM(TO_CHAR(v.VKND_CD)) AS VKND_CD" in compact_sql
    assert "TO_CHAR(v.TOT_DT, 'HH24:MI') IN" in compact_sql
    assert [params[f"time_slot{index}"] for index in range(9)] == [
        "07:15",
        "07:30",
        "07:45",
        "08:00",
        "08:15",
        "08:30",
        "08:45",
        "09:00",
        "09:15",
    ]


def test_hourly_sql_keeps_hour_source_and_hour_filter():
    sql, params = vehicle.build_fetch_all_intersections_periods_sql_params(
        vehicle.parse_periods("260501"),
        vehicle.parse_time_range("07~09"),
        {"2": "세단"},
        vehicle.ONE_HOUR_AGGREGATION,
    )
    compact_sql = re.sub(r"\s+", " ", sql)

    assert f"FROM {vehicle.TRAFFIC_TABLE} v" in compact_sql
    assert "TO_NUMBER(TO_CHAR(v.TOT_DT, 'HH24')) IN" in compact_sql
    assert params["hour0"] == 7
    assert params["hour1"] == 8


def test_required_column_validation_uses_selected_traffic_table(monkeypatch):
    checked_tables: list[str] = []

    def fake_columns(_cursor, table_name):
        checked_tables.append(table_name)
        return {
            "NODE_ID",
            "TOT_DT",
            "ACSR_ID",
            "DRCT_CD",
            "VKND_CD",
            "TRF_QNTY",
            "CRSRD_NM",
            "ACSR_NM",
            "GRP_CD",
            "CD",
            "CD_NM",
        }

    monkeypatch.setattr(vehicle, "get_table_columns", fake_columns)

    vehicle.validate_vehicle_extract_tables(object(), vehicle.FIFTEEN_MINUTE_AGGREGATION)

    assert checked_tables[0] == "S_CRSRD_VKND_TRF_15MI"


@pytest.mark.parametrize("sql", ["DELETE FROM traffic", "SELECT 1; SELECT 2"])
def test_ensure_select_sql_rejects_non_read_only_or_multiple_statements(sql):
    with pytest.raises(RuntimeError):
        vehicle.ensure_select_sql(sql)


def test_gogang_west_raw_sql_reads_only_requested_source_rows():
    sql, params = vehicle.build_gogang_west_raw_15m_sql_params()
    compact_sql = re.sub(r"\s+", " ", sql)

    assert "FROM S_CRSRD_VKND_TRF_15MI v" in compact_sql
    assert "SUM(" not in compact_sql
    assert "GROUP BY" not in compact_sql
    assert "HAVING" not in compact_sql
    assert params["approach_id"] == "ACSR000003"
    assert [params[f"date_start{index}"].date() for index in range(2)] == [
        vehicle.GOGANG_WEST_RAW_DATES[0],
        vehicle.GOGANG_WEST_RAW_DATES[1],
    ]
    assert [params[f"time_slot{index}"] for index in range(16)] == [
        "07:00",
        "07:15",
        "07:30",
        "07:45",
        "08:00",
        "08:15",
        "08:30",
        "08:45",
        "17:00",
        "17:15",
        "17:30",
        "17:45",
        "18:00",
        "18:15",
        "18:30",
        "18:45",
    ]


def test_gogang_west_raw_export_writes_requested_header_and_preserves_quantity(
    monkeypatch, tmp_path
):
    class Cursor:
        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return False

    class Connection:
        def cursor(self):
            return Cursor()

        def close(self):
            return None

    rows = [
        (datetime(2026, 9, 4, 7), "고강지하차도사거리", "직진", "SUV", "07", 38),
        (datetime(2026, 9, 4, 7), "고강지하차도사거리", "좌회전", "세단", "02", 10),
    ]
    monkeypatch.setattr(vehicle, "connect_db", lambda: Connection())
    monkeypatch.setattr(vehicle, "validate_vehicle_extract_tables", lambda *_args: None)
    monkeypatch.setattr(vehicle, "execute_select", lambda *_args: rows)
    output_path = tmp_path / "raw.csv"

    assert vehicle.export_gogang_west_raw_15m(output_path) == 2
    content = output_path.read_text(encoding="utf-8-sig").splitlines()

    assert content[0].split(",") == vehicle.GOGANG_WEST_RAW_CSV_HEADER
    assert content[1].split(",") == [
        "2026-09-04 07:00:00",
        "고강지하차도사거리",
        "서",
        "좌",
        "세단",
        "02",
        "10",
    ]
    assert content[2].endswith(",07,38")


def test_gogang_peak_raw_sql_uses_all_approaches_and_requested_slots():
    sql, params = vehicle.build_gogang_peak_raw_15m_sql_params()
    compact_sql = re.sub(r"\s+", " ", sql)

    assert "FROM S_CRSRD_VKND_TRF_15MI v" in compact_sql
    assert "a.ACSR_NM" in compact_sql
    assert "TRIM(TO_CHAR(v.VKND_CD)) AS VKND_CD" in compact_sql
    assert "SUM(" not in compact_sql
    assert "GROUP BY" not in compact_sql
    assert "HAVING" not in compact_sql
    assert params["intersection_name"] == "고강지하차도사거리"
    assert params["date_start"].date() == vehicle.GOGANG_PEAK_RAW_DATE
    assert params["date_end"].date() == vehicle.GOGANG_PEAK_RAW_DATE.replace(day=5)
    assert [params[f"time_slot{index}"] for index in range(16)] == [
        "07:00",
        "07:15",
        "07:30",
        "07:45",
        "08:00",
        "08:15",
        "08:30",
        "08:45",
        "17:00",
        "17:15",
        "17:30",
        "17:45",
        "18:00",
        "18:15",
        "18:30",
        "18:45",
    ]


def test_gogang_peak_raw_xlsx_has_one_sheet_header_and_vehicle_code(tmp_path):
    from openpyxl import load_workbook

    output_path = tmp_path / "gogang-peak.xlsx"
    vehicle.save_gogang_peak_raw_xlsx(
        [["2026-09-04 07:00:00", "고강지하차도사거리", "동(서향)", "좌", "SUV", "07", 38]],
        output_path,
    )

    workbook = load_workbook(output_path, read_only=True, data_only=True)
    try:
        assert workbook.sheetnames == ["고강지하차도사거리"]
        assert list(workbook.active.values) == [
            tuple(vehicle.GOGANG_WEST_RAW_CSV_HEADER),
            ("2026-09-04 07:00:00", "고강지하차도사거리", "동(서향)", "좌", "SUV", "07", 38),
        ]
    finally:
        workbook.close()


def test_ten_intersections_raw_sql_uses_source_rows_and_requested_conditions():
    intersections = [
        vehicle.Intersection(node_id=index, name=name)
        for index, name in enumerate(vehicle.TEN_INTERSECTIONS_RAW_NAMES, start=1)
    ]
    sql, params = vehicle.build_ten_intersections_raw_15m_sql_params(intersections)
    compact_sql = re.sub(r"\s+", " ", sql)

    assert "FROM S_CRSRD_VKND_TRF_15MI v" in compact_sql
    assert "SUM(" not in compact_sql
    assert "GROUP BY" not in compact_sql
    assert "HAVING" not in compact_sql
    assert "a.ACSR_NM" in compact_sql
    assert "TRIM(TO_CHAR(v.VKND_CD)) AS VKND_CD" in compact_sql
    assert [params[f"date_start{index}"].date() for index in range(2)] == list(
        vehicle.GOGANG_WEST_RAW_DATES
    )
    assert [params[f"time_slot{index}"] for index in range(16)] == [
        "07:00",
        "07:15",
        "07:30",
        "07:45",
        "08:00",
        "08:15",
        "08:30",
        "08:45",
        "17:00",
        "17:15",
        "17:30",
        "17:45",
        "18:00",
        "18:15",
        "18:30",
        "18:45",
    ]


def test_ten_intersections_xlsx_has_requested_sheets_headers_and_separated_rows(tmp_path):
    from openpyxl import load_workbook

    first_name, second_name, *_ = vehicle.TEN_INTERSECTIONS_RAW_NAMES
    output_path = tmp_path / "ten-intersections.xlsx"
    vehicle.save_ten_intersections_raw_xlsx(
        {
            first_name: [["2026-09-04 07:00:00", first_name, "북", "좌", "세단", "02", 10]],
            second_name: [["2026-09-04 07:15:00", second_name, "남", "직", "SUV", "07", 38]],
            **{
                name: []
                for name in vehicle.TEN_INTERSECTIONS_RAW_NAMES
                if name not in {first_name, second_name}
            },
        },
        output_path,
    )

    workbook = load_workbook(output_path, read_only=True, data_only=True)
    try:
        assert workbook.sheetnames == list(vehicle.TEN_INTERSECTIONS_RAW_NAMES)
        assert list(workbook[first_name].values) == [
            tuple(vehicle.GOGANG_WEST_RAW_CSV_HEADER),
            ("2026-09-04 07:00:00", first_name, "북", "좌", "세단", "02", 10),
        ]
        assert list(workbook[second_name].values)[1][5] == "07"
        assert list(workbook[vehicle.TEN_INTERSECTIONS_RAW_NAMES[2]].values) == [
            tuple(vehicle.GOGANG_WEST_RAW_CSV_HEADER)
        ]
    finally:
        workbook.close()
