#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""2026년 교통량 이상 요청 목록을 Notion에서 Excel로 저장한다."""

from __future__ import annotations

import json
import sys
import urllib.error
import urllib.request
from calendar import monthrange
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Callable, Iterable, Mapping, Protocol

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parents[2]
ENV_PATH = PROJECT_ROOT / "00_Data" / ".env"
OUTPUT_DIR = PROJECT_ROOT / "02_Result" / "교통량_이상추출"
NOTION_DATABASE_URL = "https://api.notion.com/v1/databases/{database_id}"
NOTION_API_URL = "https://api.notion.com/v1/data_sources/{data_source_id}/query"
NOTION_VERSION = "2025-09-03"

# `2026년 교통량 이상 요청 목록`의 Notion 데이터베이스 ID.
# 실행 시 이 DB에서 실제 데이터 소스 ID를 조회해 사용한다.
TRAFFIC_ANOMALY_DATABASE_ID = "3b86e6899fbc8003a1efd704239737a8"
TRAFFIC_ANOMALY_DATA_SOURCE_NAME = "2026년 교통량 이상 요청 목록"

HEADERS = (
    "이상발생일",
    "요청일",
    "조치 완료일",
    "조치상태",
    "이상유형",
    "교차로",
    "요청대상",
    "조치사항",
    "비고",
)
DATE_HEADERS = {"이상발생일", "요청일", "조치 완료일"}
DateValue = date | datetime | None

EXIT_INPUT_ERROR = 2
EXIT_TOKEN_ERROR = 3
EXIT_NOTION_ERROR = 4
EXIT_EMPTY_RESULT = 5
EXIT_SAVE_ERROR = 6


class NotionQuery(Protocol):
    """Notion 데이터 소스의 페이지 목록을 제공한다."""

    def fetch_pages(self) -> list[Mapping[str, Any]]:
        """모든 페이지를 반환한다."""


class NotionRequestError(RuntimeError):
    """사용자에게 표시할 수 있는 Notion 통신 오류."""


@dataclass(frozen=True)
class AnomalyRecord:
    occurred_on: DateValue
    requested_on: DateValue
    completed_on: DateValue
    status: str
    anomaly_type: str
    intersection: str
    request_target: str
    action: str
    note: str

    def values(self) -> tuple[DateValue | str, ...]:
        return (
            self.occurred_on,
            self.requested_on,
            self.completed_on,
            self.status,
            self.anomaly_type,
            self.intersection,
            self.request_target,
            self.action,
            self.note,
        )


class NotionDataSourceClient:
    """Notion API 데이터 소스 페이지네이션 구현체."""

    def __init__(
        self,
        token: str,
        database_id: str = TRAFFIC_ANOMALY_DATABASE_ID,
        opener: Callable[..., Any] = urllib.request.urlopen,
    ) -> None:
        if not database_id:
            raise NotionRequestError("대상 Notion 데이터베이스 ID가 설정되지 않았습니다.")
        self._token = token
        self._database_id = database_id
        self._opener = opener

    def fetch_pages(self) -> list[Mapping[str, Any]]:
        data_source_id = self._fetch_data_source_id()
        pages: list[Mapping[str, Any]] = []
        cursor: str | None = None
        while True:
            payload: dict[str, Any] = {"page_size": 100}
            if cursor:
                payload["start_cursor"] = cursor
            request = urllib.request.Request(
                NOTION_API_URL.format(data_source_id=data_source_id),
                data=json.dumps(payload).encode("utf-8"),
                headers={
                    "Authorization": f"Bearer {self._token}",
                    "Content-Type": "application/json",
                    "Notion-Version": NOTION_VERSION,
                },
                method="POST",
            )
            try:
                with self._opener(request, timeout=30) as response:
                    body = response.read().decode("utf-8")
            except urllib.error.HTTPError as exc:
                raise _notion_http_error(exc.code) from exc
            except urllib.error.URLError as exc:
                raise NotionRequestError(
                    "Notion 서버에 연결할 수 없습니다. 네트워크를 확인해 주세요."
                ) from exc
            except TimeoutError as exc:
                raise NotionRequestError(
                    "Notion 요청 시간이 초과되었습니다. 잠시 후 다시 시도해 주세요."
                ) from exc

            try:
                response_body = json.loads(body)
            except json.JSONDecodeError as exc:
                raise NotionRequestError("Notion 응답을 해석할 수 없습니다.") from exc
            results = response_body.get("results")
            if not isinstance(results, list):
                raise NotionRequestError("Notion 응답에 페이지 목록이 없습니다.")
            pages.extend(item for item in results if isinstance(item, Mapping))
            if not response_body.get("has_more"):
                return pages
            cursor = response_body.get("next_cursor")
            if not isinstance(cursor, str) or not cursor:
                raise NotionRequestError("Notion 다음 페이지 정보를 확인할 수 없습니다.")

    def _fetch_data_source_id(self) -> str:
        request = urllib.request.Request(
            NOTION_DATABASE_URL.format(database_id=self._database_id),
            headers={
                "Authorization": f"Bearer {self._token}",
                "Notion-Version": NOTION_VERSION,
            },
            method="GET",
        )
        try:
            with self._opener(request, timeout=30) as response:
                body = response.read().decode("utf-8")
        except urllib.error.HTTPError as exc:
            raise _notion_http_error(exc.code) from exc
        except urllib.error.URLError as exc:
            raise NotionRequestError(
                "Notion 서버에 연결할 수 없습니다. 네트워크를 확인해 주세요."
            ) from exc
        except TimeoutError as exc:
            raise NotionRequestError(
                "Notion 요청 시간이 초과되었습니다. 잠시 후 다시 시도해 주세요."
            ) from exc

        try:
            response_body = json.loads(body)
        except json.JSONDecodeError as exc:
            raise NotionRequestError("Notion 응답을 해석할 수 없습니다.") from exc
        data_sources = response_body.get("data_sources")
        if not isinstance(data_sources, list):
            raise NotionRequestError("Notion 응답에서 데이터 소스 목록을 찾을 수 없습니다.")
        matching_ids = [
            item.get("id")
            for item in data_sources
            if isinstance(item, Mapping) and item.get("name") == TRAFFIC_ANOMALY_DATA_SOURCE_NAME
        ]
        if len(matching_ids) != 1 or not isinstance(matching_ids[0], str):
            raise NotionRequestError("대상 Notion 데이터 소스를 하나로 식별할 수 없습니다.")
        return matching_ids[0]


def parse_month(raw: str) -> int:
    """1~12월 입력값을 검증한다."""

    value = raw.strip()
    if not value:
        raise ValueError("월을 입력해 주세요. (1~12)")
    if not value.isdigit():
        raise ValueError("월은 숫자로 입력해 주세요. (1~12)")
    month = int(value)
    if not 1 <= month <= 12:
        raise ValueError("월은 1부터 12 사이여야 합니다.")
    return month


def load_notion_token(env_path: Path = ENV_PATH) -> str | None:
    """.env에서 토큰만 읽는다. 값은 절대로 출력하지 않는다."""

    try:
        lines = env_path.read_text(encoding="utf-8-sig").splitlines()
    except FileNotFoundError:
        return None
    for line in lines:
        key, separator, value = line.partition("=")
        if separator and key.strip() == "NOTION_API_TOKEN":
            return value.strip().strip('"').strip("'") or None
    return None


def extract_records(pages: Iterable[Mapping[str, Any]]) -> list[AnomalyRecord]:
    return [record for page in pages if (record := extract_record(page)) is not None]


def extract_record(page: Mapping[str, Any]) -> AnomalyRecord | None:
    properties = page.get("properties")
    if not isinstance(properties, Mapping):
        return None
    return AnomalyRecord(
        occurred_on=_property_date(properties.get("이상발생일")),
        requested_on=_property_date(properties.get("요청일")),
        completed_on=_property_date(properties.get("조치 완료일")),
        status=_property_text(properties.get("조치상태")),
        anomaly_type=_property_text(properties.get("이상유형")),
        intersection=_property_text(properties.get("교차로")),
        request_target=_property_text(properties.get("요청대상")),
        action=_property_text(properties.get("조치사항")),
        note=_property_text(properties.get("비고")),
    )


def select_records(
    records: Iterable[AnomalyRecord], month: int
) -> tuple[list[AnomalyRecord], list[AnomalyRecord]]:
    """완료 월별 목록과 연초부터 해당 월까지의 미완료 목록을 선별한다."""

    end_date = date(2026, month, monthrange(2026, month)[1])
    completed = [
        record
        for record in records
        if record.status == "완료"
        and record.completed_on
        and record.completed_on.year == 2026
        and record.completed_on.month == month
    ]
    incomplete = [
        record
        for record in records
        if record.status == "미완료"
        and record.occurred_on
        and date(2026, 1, 1) <= _calendar_date(record.occurred_on) <= end_date
    ]
    return _sort_by_occurred_on(completed), _sort_by_occurred_on(incomplete)


def save_workbook(
    completed: list[AnomalyRecord], incomplete: list[AnomalyRecord], output_path: Path
) -> None:
    workbook = Workbook()
    completed_sheet = workbook.active
    completed_sheet.title = "완료"
    incomplete_sheet = workbook.create_sheet("미완료")
    _write_sheet(completed_sheet, completed)
    _write_sheet(incomplete_sheet, incomplete)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def build_output_path(month: int, output_dir: Path = OUTPUT_DIR) -> Path:
    return output_dir / f"2026년_{month}월_교통량_이상목록.xlsx"


def run(
    raw_month: str,
    query: NotionQuery,
    output_dir: Path = OUTPUT_DIR,
) -> Path:
    month = parse_month(raw_month)
    records = extract_records(query.fetch_pages())
    completed, incomplete = select_records(records, month)
    if not completed and not incomplete:
        raise LookupError("선택한 기간에 해당하는 완료 또는 미완료 목록이 없습니다.")
    output_path = build_output_path(month, output_dir)
    save_workbook(completed, incomplete, output_path)
    return output_path


def main() -> int:
    try:
        month = parse_month(input("추출할 월을 입력해 주세요 (1~12): "))
    except (EOFError, ValueError) as exc:
        print(f"입력 오류: {exc}", file=sys.stderr)
        return EXIT_INPUT_ERROR

    token = load_notion_token()
    if not token:
        print("NOTION_API_TOKEN을 00_Data/.env에 설정해 주세요.", file=sys.stderr)
        return EXIT_TOKEN_ERROR

    try:
        output_path = run(str(month), NotionDataSourceClient(token))
    except NotionRequestError as exc:
        print(f"Notion 조회 오류: {exc}", file=sys.stderr)
        return EXIT_NOTION_ERROR
    except LookupError as exc:
        print(str(exc), file=sys.stderr)
        return EXIT_EMPTY_RESULT
    except PermissionError:
        print(
            "결과 파일이 열려 있어 저장할 수 없습니다. Excel에서 파일을 닫고 다시 실행해 주세요.",
            file=sys.stderr,
        )
        return EXIT_SAVE_ERROR
    except OSError as exc:
        print(f"파일 저장 중 오류가 발생했습니다: {exc}", file=sys.stderr)
        return EXIT_SAVE_ERROR

    print(f"저장 완료: {output_path}")
    return 0


def _property_date(property_value: Any) -> DateValue:
    if not isinstance(property_value, Mapping):
        return None
    value = property_value.get("date")
    if not isinstance(value, Mapping) or not isinstance(value.get("start"), str):
        return None
    raw_value = value["start"]
    try:
        if "T" not in raw_value:
            return date.fromisoformat(raw_value)
        return datetime.fromisoformat(raw_value.replace("Z", "+00:00")).replace(tzinfo=None)
    except ValueError:
        return None


def _property_text(property_value: Any) -> str:
    if not isinstance(property_value, Mapping):
        return ""
    property_type = property_value.get("type")
    if property_type in {"select", "status"}:
        value = property_value.get(property_type)
        return value.get("name", "") if isinstance(value, Mapping) else ""
    if property_type == "multi_select":
        values = property_value.get("multi_select")
        return (
            ", ".join(
                item["name"] for item in values if isinstance(item, Mapping) and item.get("name")
            )
            if isinstance(values, list)
            else ""
        )
    if property_type in {"title", "rich_text"}:
        values = property_value.get(property_type)
        return (
            "".join(item.get("plain_text", "") for item in values if isinstance(item, Mapping))
            if isinstance(values, list)
            else ""
        )
    if property_type == "formula":
        formula = property_value.get("formula")
        if isinstance(formula, Mapping):
            return str(formula.get("string") or formula.get("number") or "")
    return ""


def _sort_by_occurred_on(records: list[AnomalyRecord]) -> list[AnomalyRecord]:
    return sorted(records, key=lambda record: _date_time_sort_key(record.occurred_on))


def _calendar_date(value: date | datetime) -> date:
    return value.date() if isinstance(value, datetime) else value


def _date_time_sort_key(value: DateValue) -> datetime:
    if value is None:
        return datetime.max
    if isinstance(value, datetime):
        return value
    return datetime.combine(value, time.min)


def _write_sheet(worksheet: Any, records: Iterable[AnomalyRecord]) -> None:
    worksheet.append(HEADERS)
    for record in records:
        worksheet.append(record.values())
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.freeze_panes = "A2"
    for column_index, header in enumerate(HEADERS, start=1):
        worksheet.column_dimensions[get_column_letter(column_index)].width = (
            18 if header in DATE_HEADERS else 20
        )
    for row in worksheet.iter_rows(min_row=2, max_col=len(HEADERS)):
        for cell, header in zip(row, HEADERS, strict=True):
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if header in DATE_HEADERS:
                if isinstance(cell.value, datetime):
                    cell.number_format = "yyyy-mm-dd hh:mm"
                elif isinstance(cell.value, date):
                    cell.number_format = "yyyy-mm-dd"


def _notion_http_error(status_code: int) -> NotionRequestError:
    messages = {
        401: "Notion 토큰이 올바르지 않습니다.",
        403: "대상 데이터 소스에 대한 Notion Integration 읽기 권한이 없습니다.",
        404: "대상 Notion 데이터 소스를 찾을 수 없습니다. ID와 연결 상태를 확인해 주세요.",
        429: "Notion 요청 한도를 초과했습니다. 잠시 후 다시 시도해 주세요.",
    }
    return NotionRequestError(
        messages.get(status_code, f"Notion 요청이 실패했습니다. (HTTP {status_code})")
    )


if __name__ == "__main__":
    raise SystemExit(main())
