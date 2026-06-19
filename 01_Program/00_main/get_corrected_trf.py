#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Local API based corrected direction traffic extractor.

This script intentionally reads traffic data only from the running API server.
It does not read `.env` or connect to the database directly.
"""

from __future__ import annotations

import json
import re
import sys
import urllib.error
import urllib.request
from calendar import monthrange
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover - exercised only when dependency is absent.
    Workbook = None
    Alignment = Border = Font = PatternFill = Side = None
    get_column_letter = None


API_BASE_URL = "http://localhost:8000"
PROJECT_ROOT = Path(__file__).resolve().parents[2]
OUTPUT_DIR = PROJECT_ROOT / "02_Result" / "API_추출"
ENDPOINT = "/corrected-traffic-drct"
REQUIRED_OPENAPI_PATHS = (ENDPOINT, "/intersections")
TRAFFIC_NUMBER_FORMAT = "#,##0"

OUTPUT_MODE_INTERSECTION = "intersection"
OUTPUT_MODE_DIRECTION = "direction"

HANGUL_RE = re.compile(r"[가-힣]")
INVALID_FILENAME_RE = re.compile(r'[\\/:*?"<>|]+')


@dataclass(frozen=True)
class Intersection:
    node_id: str | int
    name: str
    order: int


@dataclass(frozen=True)
class Period:
    start: date
    end: date
    label: str


@dataclass(frozen=True)
class IntervalSpec:
    label: str
    api_interval: str
    minutes: int
    rollup_from_5m: bool = False


@dataclass(frozen=True)
class TimeRange:
    start_minute: int
    end_minute: int
    label: str

    @property
    def is_full_day(self) -> bool:
        return self.start_minute == 0 and self.end_minute == 24 * 60


INTERVAL_SPECS = {
    "5분": IntervalSpec("5분", "5m", 5),
    "15분": IntervalSpec("15분", "5m", 15, rollup_from_5m=True),
    "30분": IntervalSpec("30분", "5m", 30, rollup_from_5m=True),
    "1시간": IntervalSpec("1시간", "1h", 60),
    "1일": IntervalSpec("1일", "1d", 24 * 60),
}

INTERVAL_ALIASES = {
    "5": "5분",
    "5m": "5분",
    "5분": "5분",
    "15": "15분",
    "15m": "15분",
    "15분": "15분",
    "30": "30분",
    "30m": "30분",
    "30분": "30분",
    "1h": "1시간",
    "1시간": "1시간",
    "60": "1시간",
    "1d": "1일",
    "1일": "1일",
    "일": "1일",
}


def parse_interval(raw: str) -> IntervalSpec:
    key = INTERVAL_ALIASES.get(raw.strip().lower())
    if key is None:
        allowed = ", ".join(INTERVAL_SPECS)
        raise ValueError(f"집계 간격은 {allowed} 중 하나여야 합니다.")
    return INTERVAL_SPECS[key]


def parse_periods(raw: str) -> list[Period]:
    parts = [part.strip() for part in raw.split(",") if part.strip()]
    if not parts:
        raise ValueError("기간을 입력해 주세요.")

    periods: list[Period] = []
    for part in parts:
        if re.fullmatch(r"\d{6}~\d{6}", part):
            start_raw, end_raw = part.split("~", 1)
            start = _parse_yymmdd(start_raw)
            end = _parse_yymmdd(end_raw)
            if start > end:
                raise ValueError(f"시작 날짜가 종료 날짜보다 늦습니다: {part}")
            periods.append(Period(start, end, part))
            continue

        if re.fullmatch(r"\d{6}", part):
            day = _parse_yymmdd(part)
            periods.append(Period(day, day, part))
            continue

        if re.fullmatch(r"\d{4}", part) and not part.startswith("20"):
            year = 2000 + int(part[:2])
            month = int(part[2:])
            if not 1 <= month <= 12:
                raise ValueError(f"월 값이 올바르지 않습니다: {part}")
            last_day = monthrange(year, month)[1]
            periods.append(Period(date(year, month, 1), date(year, month, last_day), part))
            continue

        if re.fullmatch(r"20\d{2}", part):
            year = int(part)
            periods.append(Period(date(year, 1, 1), date(year, 12, 31), part))
            continue

        raise ValueError(
            f"기간 형식이 올바르지 않습니다: {part} (예: 260412~260415, 260412, 2605, 2026)"
        )

    return periods


def _parse_yymmdd(raw: str) -> date:
    try:
        return date(2000 + int(raw[:2]), int(raw[2:4]), int(raw[4:6]))
    except ValueError as exc:
        raise ValueError(f"날짜 값이 올바르지 않습니다: {raw}") from exc


def parse_time_range(raw: str, interval: IntervalSpec | str) -> TimeRange:
    spec = parse_interval(interval) if isinstance(interval, str) else interval
    value = raw.strip()
    if value in {"24시간", "24", "all", "ALL", "*", "전체"}:
        time_range = TimeRange(0, 24 * 60, "24시간")
    else:
        match = re.fullmatch(r"(\d{1,2}:\d{2})~(\d{1,2}:\d{2})", value)
        if not match:
            raise ValueError("시간 형식은 24시간 또는 HH:MM~HH:MM 입니다.")
        start = _parse_clock(match.group(1), is_end=False)
        end = _parse_clock(match.group(2), is_end=True)
        if start >= end:
            raise ValueError("시간 범위는 시작 포함, 종료 제외이며 시작이 종료보다 빨라야 합니다.")
        time_range = TimeRange(start, end, value)

    _validate_time_range_alignment(time_range, spec)
    return time_range


def _parse_clock(raw: str, is_end: bool) -> int:
    hour_raw, minute_raw = raw.split(":", 1)
    hour = int(hour_raw)
    minute = int(minute_raw)
    if minute < 0 or minute > 59:
        raise ValueError(f"분 값이 올바르지 않습니다: {raw}")
    if hour == 24:
        if not is_end or minute != 0:
            raise ValueError("24:00은 종료 시각으로만 사용할 수 있습니다.")
        return 24 * 60
    if hour < 0 or hour > 23:
        raise ValueError(f"시 값이 올바르지 않습니다: {raw}")
    return hour * 60 + minute


def _validate_time_range_alignment(time_range: TimeRange, interval: IntervalSpec) -> None:
    if interval.label == "1일":
        if not time_range.is_full_day:
            raise ValueError("1일 집계는 24시간 범위만 사용할 수 있습니다.")
        return

    if (
        time_range.start_minute % interval.minutes != 0
        or time_range.end_minute % interval.minutes != 0
    ):
        raise ValueError(
            f"시간 범위가 {interval.label} 집계 간격과 맞지 않습니다. "
            "예: 1시간 집계에는 17:15~17:45를 사용할 수 없습니다."
        )

    if time_range.end_minute - time_range.start_minute < interval.minutes:
        raise ValueError("시간 범위가 선택한 집계 간격보다 짧습니다.")


def hours_for_api(time_range: TimeRange) -> list[int]:
    if time_range.is_full_day:
        return list(range(24))
    first_hour = time_range.start_minute // 60
    last_hour = (time_range.end_minute - 1) // 60
    return list(range(first_hour, last_hour + 1))


def resolve_intersection_input(raw: str, items: list[dict]) -> list[Intersection]:
    normalized_items = []
    for order, item in enumerate(items):
        node_id = item.get("node_id", item.get("id"))
        name = item.get("name", item.get("node_name"))
        if node_id is None or not name:
            continue
        normalized_items.append(Intersection(node_id=node_id, name=str(name), order=order))

    if not normalized_items:
        raise ValueError("API에서 교차로 목록을 찾을 수 없습니다.")

    value = raw.strip()
    if value in {"*", "all", "ALL", "전체"}:
        return normalized_items

    by_name = {item.name: item for item in normalized_items}
    selected: list[Intersection] = []
    seen: set[str] = set()
    for name in [part.strip() for part in value.split(",") if part.strip()]:
        item = by_name.get(name)
        if item is None:
            raise ValueError(f"교차로를 찾을 수 없습니다: {name}")
        key = _node_key(item.node_id)
        if key not in seen:
            selected.append(Intersection(item.node_id, item.name, len(selected)))
            seen.add(key)

    if not selected:
        raise ValueError("교차로를 입력해 주세요.")
    return selected


def verify_api_contract(base_url: str = API_BASE_URL) -> None:
    schema = api_get("/openapi.json", base_url=base_url)
    paths = schema.get("paths") or {}
    missing = [path for path in REQUIRED_OPENAPI_PATHS if path not in paths]
    if missing:
        joined = ", ".join(missing)
        raise RuntimeError(f"서버 버전이 오래되어 필요한 API가 없습니다: {joined}")


def fetch_intersection_items(base_url: str = API_BASE_URL) -> list[dict]:
    data = api_get("/intersections", base_url=base_url)
    items = data.get("items")
    if not isinstance(items, list):
        raise RuntimeError("/intersections 응답에 items 목록이 없습니다.")
    return items


def api_get(path: str, base_url: str = API_BASE_URL) -> dict:
    return _api_request_json("GET", path, base_url=base_url)


def api_post(path: str, payload: dict, base_url: str = API_BASE_URL) -> dict:
    return _api_request_json("POST", path, payload=payload, base_url=base_url)


def _api_request_json(
    method: str,
    path: str,
    payload: dict | None = None,
    base_url: str = API_BASE_URL,
) -> dict:
    url = _join_url(base_url, path)
    data = None
    headers = {}
    if payload is not None:
        data = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        headers["Content-Type"] = "application/json"

    req: urllib.request.Request | str
    if method == "GET":
        req = url
    else:
        req = urllib.request.Request(url, data=data, headers=headers, method=method)

    try:
        with urllib.request.urlopen(req, timeout=60) as response:
            body = response.read().decode("utf-8")
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"API 요청 실패 ({path}, HTTP {exc.code}): {detail}") from exc
    except urllib.error.URLError as exc:
        raise RuntimeError(f"API 서버에 연결할 수 없습니다: {base_url}") from exc

    if not body:
        return {}
    parsed = json.loads(body)
    if not isinstance(parsed, dict):
        raise RuntimeError(f"API 응답이 객체 형식이 아닙니다: {path}")
    return parsed


def _join_url(base_url: str, path: str) -> str:
    return f"{base_url.rstrip('/')}/{path.lstrip('/')}"


def fetch_corrected_direction_slots(
    intersections: list[Intersection],
    periods: list[Period],
    interval: IntervalSpec,
    time_range: TimeRange,
    base_url: str = API_BASE_URL,
) -> list[dict]:
    order_map = {_node_key(item.node_id): item.order for item in intersections}
    name_map = {_node_key(item.node_id): item.name for item in intersections}
    rows: list[dict] = []

    for period in periods:
        payload = {
            "node_ids": [item.node_id for item in intersections],
            "date_start": period.start.isoformat(),
            "date_end": period.end.isoformat(),
            "hours": hours_for_api(time_range),
            "interval": interval.api_interval,
        }
        response = api_post(ENDPOINT, payload, base_url=base_url)
        for slot in response.get("drct_slots", []):
            if not isinstance(slot, dict):
                continue
            row = normalize_api_slot(slot, order_map, name_map)
            if not is_valid_direction_slot(row):
                continue
            if not slot_in_time_range(row, time_range):
                continue
            rows.append(row)

    if interval.rollup_from_5m:
        rows = aggregate_slots_by_interval(rows, interval)

    return sort_slots(rows)


def normalize_api_slot(
    slot: dict,
    order_map: dict[str, int] | None = None,
    name_map: dict[str, str] | None = None,
) -> dict:
    order_map = order_map or {}
    name_map = name_map or {}
    node_id = slot.get("node_id")
    node_key = _node_key(node_id)
    timestamp = slot_datetime(slot)
    node_name = str(slot.get("node_name") or name_map.get(node_key) or node_id or "")
    approach_id = slot.get("approach_id", slot.get("acsr_id"))
    drct_cd = normalize_drct_cd(slot.get("drct_cd"))
    value = slot_value(slot)

    return {
        "interval": slot.get("interval"),
        "timestamp": timestamp.isoformat(),
        "date": timestamp.date().isoformat(),
        "hour": timestamp.hour,
        "minute": timestamp.minute,
        "node_id": node_id,
        "node_key": node_key,
        "node_name": node_name,
        "node_order": order_map.get(node_key, 10**9),
        "approach_id": approach_id,
        "approach_name": str(slot.get("approach_name") or approach_id or ""),
        "drct_cd": drct_cd,
        "drct_name": str(slot.get("drct_name") or ""),
        "traffic_volume": slot.get("traffic_volume"),
        "corrected_value": slot.get("corrected_value"),
        "value": value,
    }


def slot_datetime(slot: dict) -> datetime:
    timestamp = slot.get("timestamp")
    if timestamp:
        return datetime.fromisoformat(str(timestamp))

    slot_date = date.fromisoformat(str(slot["date"]))
    hour = int(slot.get("hour") or 0)
    minute = slot.get("minute")
    minute = 0 if minute is None else int(minute)
    return datetime(slot_date.year, slot_date.month, slot_date.day, hour, minute)


def slot_value(slot: dict) -> int | float | None:
    value = slot.get("corrected_value")
    if value is None:
        value = slot.get("traffic_volume")
    return numeric_or_none(value)


def numeric_or_none(value: Any) -> int | float | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, int | float):
        return value
    try:
        parsed = float(value)
    except (TypeError, ValueError):
        return None
    return int(parsed) if parsed.is_integer() else parsed


def normalize_drct_cd(value: Any) -> str:
    if value is None:
        return ""
    raw = str(value).strip()
    if raw.isdigit():
        return raw.zfill(2)
    return raw


def is_valid_direction_slot(row: dict) -> bool:
    if normalize_drct_cd(row.get("drct_cd")) == "00":
        return False
    name = str(row.get("drct_name") or "").strip()
    if not name or name.isdigit():
        return False
    return bool(HANGUL_RE.search(name))


def slot_in_time_range(row: dict, time_range: TimeRange) -> bool:
    if time_range.is_full_day:
        return True
    ts = datetime.fromisoformat(row["timestamp"])
    minute_of_day = ts.hour * 60 + ts.minute
    return time_range.start_minute <= minute_of_day < time_range.end_minute


def aggregate_slots_by_interval(rows: list[dict], interval: IntervalSpec) -> list[dict]:
    if not interval.rollup_from_5m:
        return sort_slots(rows)

    aggregate: dict[tuple, dict] = {}
    for row in rows:
        ts = datetime.fromisoformat(row["timestamp"])
        bucket_minute = (ts.minute // interval.minutes) * interval.minutes
        bucket_ts = ts.replace(minute=bucket_minute, second=0, microsecond=0)
        key = (
            bucket_ts.isoformat(),
            row["node_key"],
            row.get("approach_id"),
            row.get("drct_cd"),
        )
        if key not in aggregate:
            item = dict(row)
            item["interval"] = f"{interval.minutes}m"
            item["timestamp"] = bucket_ts.isoformat()
            item["date"] = bucket_ts.date().isoformat()
            item["hour"] = bucket_ts.hour
            item["minute"] = bucket_ts.minute
            item["traffic_volume"] = None
            item["corrected_value"] = None
            item["value"] = None
            aggregate[key] = item
        value = numeric_or_none(row.get("value"))
        if value is not None:
            aggregate[key]["value"] = sum_optional(aggregate[key]["value"], value)

    result = []
    for item in aggregate.values():
        item["traffic_volume"] = item["value"]
        item["corrected_value"] = item["value"]
        result.append(item)
    return sort_slots(result)


def sum_optional(left: int | float | None, right: int | float | None) -> int | float | None:
    if right is None:
        return left
    if left is None:
        return right
    return left + right


def sort_slots(rows: list[dict]) -> list[dict]:
    return sorted(
        rows,
        key=lambda row: (
            row["timestamp"],
            row.get("node_order", 10**9),
            sort_id(row.get("approach_id")),
            sort_id(row.get("drct_cd")),
        ),
    )


def sort_id(value: Any) -> tuple[int, Any]:
    if value is None:
        return (1, "")
    raw = str(value)
    if raw.isdigit():
        return (0, int(raw))
    return (1, raw)


def _node_key(value: Any) -> str:
    return "" if value is None else str(value)


def format_slot_label(row: dict, interval: IntervalSpec) -> str:
    ts = datetime.fromisoformat(row["timestamp"])
    if interval.label == "1일":
        return ts.strftime("%Y-%m-%d")

    end = ts + timedelta(minutes=interval.minutes)
    end_clock = "24:00" if end.date() > ts.date() and end.hour == 0 else end.strftime("%H:%M")
    return f"{ts:%Y-%m-%d} {ts:%H:%M}~{end_clock}"


def create_workbook(rows: list[dict], interval: IntervalSpec, output_mode: str):
    require_openpyxl()
    wb = Workbook()
    ws = wb.active
    if output_mode == OUTPUT_MODE_INTERSECTION:
        ws.title = "교차로"
        write_intersection_sheet(ws, rows, interval)
    elif output_mode == OUTPUT_MODE_DIRECTION:
        ws.title = "방향"
        write_direction_sheet(ws, rows, interval)
    else:
        raise ValueError(f"알 수 없는 출력 방식입니다: {output_mode}")
    return wb


def save_workbook(
    rows: list[dict],
    interval: IntervalSpec,
    output_mode: str,
    output_path: Path,
) -> None:
    wb = create_workbook(rows, interval, output_mode)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def require_openpyxl() -> None:
    if Workbook is None:
        raise RuntimeError(
            "XLSX 저장에는 openpyxl 패키지가 필요합니다. "
            "설치 후 다시 실행해 주세요: python -m pip install openpyxl"
        )


def write_intersection_sheet(ws, rows: list[dict], interval: IntervalSpec) -> None:
    groups = build_intersection_groups(rows)
    row1 = ["시간대", "교차로 이름"]
    row2 = ["", ""]
    for group in groups:
        span = len(group["turns"]) + 1
        row1.extend([group["display_name"]] + [""] * (span - 1))
        row2.extend([turn["name"] for turn in group["turns"]])
        row2.append("합계")

    ws.append(row1)
    ws.append(row2)
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
    ws.merge_cells(start_row=1, start_column=2, end_row=2, end_column=2)

    col = 3
    for group in groups:
        span = len(group["turns"]) + 1
        ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + span - 1)
        col += span

    value_map = build_value_map(rows)
    time_nodes = build_time_node_rows(rows)
    for timestamp, node_key, node_name in time_nodes:
        output_row = [format_slot_label({"timestamp": timestamp}, interval), node_name]
        numeric_cols: list[int] = []
        for group in groups:
            values = []
            for turn in group["turns"]:
                if group["node_key"] == node_key:
                    value = value_map.get((timestamp, node_key, group["approach_id"], turn["code"]))
                else:
                    value = None
                values.append(value)
                output_row.append(display_traffic_value(value))
                if is_positive_number(value):
                    numeric_cols.append(len(output_row))
            total = positive_total(values)
            output_row.append(display_traffic_value(total))
            if is_positive_number(total):
                numeric_cols.append(len(output_row))
        ws.append(output_row)
        apply_numeric_format(ws, ws.max_row, numeric_cols)

    style_sheet(ws, header_rows=2)
    ws.freeze_panes = "C3"


def write_direction_sheet(ws, rows: list[dict], interval: IntervalSpec) -> None:
    turns = build_turn_columns(rows)
    header = ["시간대", "방향 이름"] + [turn["name"] for turn in turns] + ["합계"]
    ws.append(header)

    value_map = build_value_map(rows)
    direction_rows = build_time_direction_rows(rows)
    for timestamp, node_key, approach_id, approach_name in direction_rows:
        values = [value_map.get((timestamp, node_key, approach_id, turn["code"])) for turn in turns]
        total = positive_total(values)
        output_row = [format_slot_label({"timestamp": timestamp}, interval), approach_name]
        output_row.extend(display_traffic_value(value) for value in values)
        output_row.append(display_traffic_value(total))
        ws.append(output_row)
        numeric_cols = [
            index
            for index, value in enumerate(values + [total], start=3)
            if is_positive_number(value)
        ]
        apply_numeric_format(ws, ws.max_row, numeric_cols)

    style_sheet(ws, header_rows=1)
    ws.freeze_panes = "C2"


def build_intersection_groups(rows: list[dict]) -> list[dict]:
    groups: dict[tuple, dict] = {}
    for row in rows:
        key = (row["node_key"], row.get("approach_id"))
        if key not in groups:
            groups[key] = {
                "node_key": row["node_key"],
                "node_order": row.get("node_order", 10**9),
                "approach_id": row.get("approach_id"),
                "approach_name": row.get("approach_name", ""),
                "display_name": compact_approach_name(
                    row.get("node_name", ""),
                    row.get("approach_name", ""),
                ),
                "turn_map": {},
            }
        groups[key]["turn_map"][row["drct_cd"]] = row["drct_name"]

    result = []
    for group in groups.values():
        turns = [
            {"code": code, "name": name}
            for code, name in sorted(group["turn_map"].items(), key=lambda item: sort_id(item[0]))
        ]
        result.append({**group, "turns": turns})

    return sorted(
        result,
        key=lambda item: (item["node_order"], sort_id(item["approach_id"])),
    )


def build_turn_columns(rows: list[dict]) -> list[dict]:
    turn_map = {row["drct_cd"]: row["drct_name"] for row in rows}
    return [
        {"code": code, "name": name}
        for code, name in sorted(turn_map.items(), key=lambda item: sort_id(item[0]))
    ]


def compact_approach_name(node_name: str, approach_name: str) -> str:
    node = str(node_name).strip()
    approach = str(approach_name).strip()
    if node and approach.startswith(f"{node}-"):
        return approach[len(node) + 1 :].strip()
    if node and approach.startswith(node):
        compact = approach[len(node) :].lstrip(" -_")
        if compact:
            return compact
    return approach


def build_value_map(rows: list[dict]) -> dict[tuple, int | float | None]:
    values: dict[tuple, int | float | None] = {}
    for row in rows:
        key = (row["timestamp"], row["node_key"], row.get("approach_id"), row["drct_cd"])
        values[key] = sum_optional(values.get(key), numeric_or_none(row.get("value")))
    return values


def build_time_node_rows(rows: list[dict]) -> list[tuple[str, str, str]]:
    seen: dict[tuple, tuple[str, str, str]] = {}
    for row in rows:
        key = (row["timestamp"], row.get("node_order", 10**9), row["node_key"])
        seen[key] = (row["timestamp"], row["node_key"], row["node_name"])
    return [seen[key] for key in sorted(seen)]


def build_time_direction_rows(rows: list[dict]) -> list[tuple[str, str, Any, str]]:
    seen: dict[tuple, tuple[str, str, Any, str]] = {}
    for row in rows:
        key = (
            row["timestamp"],
            row.get("node_order", 10**9),
            sort_id(row.get("approach_id")),
            row["node_key"],
        )
        seen[key] = (
            row["timestamp"],
            row["node_key"],
            row.get("approach_id"),
            row.get("approach_name", ""),
        )
    return [seen[key] for key in sorted(seen)]


def display_traffic_value(value: int | float | None) -> int | float | str:
    value = numeric_or_none(value)
    if value is None or value <= 0:
        return "-"
    return value


def is_positive_number(value: Any) -> bool:
    numeric = numeric_or_none(value)
    return numeric is not None and numeric > 0


def positive_total(values: list[int | float | None]) -> int | float | None:
    total: int | float | None = None
    for value in values:
        numeric = numeric_or_none(value)
        if numeric is not None and numeric > 0:
            total = sum_optional(total, numeric)
    return total


def apply_numeric_format(ws, row_index: int, columns: list[int]) -> None:
    for column in columns:
        ws.cell(row=row_index, column=column).number_format = TRAFFIC_NUMBER_FORMAT


def style_sheet(ws, header_rows: int) -> None:
    require_openpyxl()
    dark_fill = PatternFill("solid", fgColor="1F4E78")
    light_fill = PatternFill("solid", fgColor="D9EAF7")
    odd_fill = PatternFill("solid", fgColor="F7FBFF")
    even_fill = PatternFill("solid", fgColor="FFFFFF")
    header_font = Font(color="FFFFFF", bold=True)
    subheader_font = Font(color="1F2937", bold=True)
    center = Alignment(horizontal="center", vertical="center")
    border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = center
            cell.border = border

    for row_idx in range(1, header_rows + 1):
        fill = dark_fill if row_idx == 1 else light_fill
        font = header_font if row_idx == 1 else subheader_font
        for cell in ws[row_idx]:
            cell.fill = fill
            cell.font = font

    for row_idx in range(header_rows + 1, ws.max_row + 1):
        fill = odd_fill if (row_idx - header_rows) % 2 else even_fill
        for cell in ws[row_idx]:
            cell.fill = fill

    auto_size_columns(ws)
    ws.sheet_view.showGridLines = False


def auto_size_columns(ws) -> None:
    if get_column_letter is None:
        return
    for column_idx in range(1, ws.max_column + 1):
        max_len = 0
        for row_idx in range(1, ws.max_row + 1):
            value = ws.cell(row=row_idx, column=column_idx).value
            if value is None:
                continue
            max_len = max(max_len, display_width(str(value)))
        width = max(10, min(max_len + 3, 36))
        ws.column_dimensions[get_column_letter(column_idx)].width = width


def display_width(value: str) -> int:
    return sum(2 if HANGUL_RE.match(char) else 1 for char in value)


def parse_output_mode(raw: str) -> str:
    value = raw.strip()
    if value in {"1", "교차로", "교차로 교통량"}:
        return OUTPUT_MODE_INTERSECTION
    if value in {"2", "방향", "방향별", "방향별 교통량"}:
        return OUTPUT_MODE_DIRECTION
    raise ValueError("출력방식은 1 또는 2를 선택해 주세요.")


def make_output_path(
    intersections: list[Intersection],
    periods: list[Period],
    interval: IntervalSpec,
    output_mode: str,
    output_dir: Path = OUTPUT_DIR,
) -> Path:
    if len(intersections) == 1:
        intersection_label = intersections[0].name
    elif len(intersections) <= 3:
        intersection_label = "_".join(item.name for item in intersections)
    else:
        intersection_label = f"{len(intersections)}개교차로"

    period_label = "_".join(period.label for period in periods)
    mode_label = "교차로" if output_mode == OUTPUT_MODE_INTERSECTION else "방향"
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = (
        f"보정교통량_{intersection_label}_{period_label}_{interval.label}_{mode_label}_{stamp}.xlsx"
    )
    return output_dir / safe_filename(filename)


def safe_filename(value: str) -> str:
    sanitized = INVALID_FILENAME_RE.sub("_", value).strip()
    if len(sanitized) > 180:
        suffix = Path(sanitized).suffix
        stem = Path(sanitized).stem[: 180 - len(suffix)]
        sanitized = f"{stem}{suffix}"
    return sanitized


def input_intersections(items: list[dict]) -> list[Intersection]:
    print("\n[1단계] 교차로 입력")
    print("  예: 계남고가사거리,중동IC사거리")
    print("  전체 선택: *")
    print(f"  API 교차로 수: {len(items)}")
    while True:
        raw = input("교차로: ").strip()
        try:
            selected = resolve_intersection_input(raw, items)
        except ValueError as exc:
            print(f"  {exc}")
            continue
        print(f"  선택: {', '.join(item.name for item in selected[:5])}", end="")
        if len(selected) > 5:
            print(f" 외 {len(selected) - 5}개", end="")
        print()
        return selected


def input_periods() -> list[Period]:
    print("\n[2단계] 기간 입력")
    print("  예: 260412~260415, 260412, 2605, 2026")
    while True:
        raw = input("기간: ").strip()
        try:
            return parse_periods(raw)
        except ValueError as exc:
            print(f"  {exc}")


def input_interval() -> IntervalSpec:
    print("\n[3단계] 집계 간격 입력")
    print("  선택: 5분, 15분, 30분, 1시간, 1일")
    while True:
        raw = input("집계 간격: ").strip()
        try:
            return parse_interval(raw)
        except ValueError as exc:
            print(f"  {exc}")


def input_time_range(interval: IntervalSpec) -> TimeRange:
    print("\n[4단계] 시간 입력")
    print("  예: 24시간, 18:00~24:00, 17:15~17:45")
    while True:
        raw = input("시간: ").strip()
        try:
            return parse_time_range(raw, interval)
        except ValueError as exc:
            print(f"  {exc}")


def input_output_mode() -> str:
    print("\n[5단계] 출력방식 입력")
    print("  1. 교차로 교통량")
    print("  2. 방향별 교통량")
    while True:
        raw = input("출력방식: ").strip()
        try:
            return parse_output_mode(raw)
        except ValueError as exc:
            print(f"  {exc}")


def main() -> int:
    try:
        print("보정 교통량 추출기를 시작합니다.")
        print(f"API 서버: {API_BASE_URL}")
        verify_api_contract(API_BASE_URL)
        items = fetch_intersection_items(API_BASE_URL)
        intersections = input_intersections(items)
        periods = input_periods()
        interval = input_interval()
        time_range = input_time_range(interval)
        output_mode = input_output_mode()

        print("\nAPI 조회 중...")
        rows = fetch_corrected_direction_slots(
            intersections,
            periods,
            interval,
            time_range,
            base_url=API_BASE_URL,
        )
        output_path = make_output_path(intersections, periods, interval, output_mode)
        save_workbook(rows, interval, output_mode, output_path)
        print(f"\n저장 완료: {output_path}")
        print(f"행 수: {len(rows):,}")
        return 0
    except KeyboardInterrupt:
        print("\n사용자에 의해 중단되었습니다.")
        return 130
    except Exception as exc:
        print(f"\n오류: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
