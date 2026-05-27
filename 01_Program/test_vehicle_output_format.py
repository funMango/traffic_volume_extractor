import csv
import importlib.util
import sqlite3
import sys
import unittest
import builtins
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch


MODULE_PATH = Path(__file__).with_name("차종별_교차로_추출.py")
SPEC = importlib.util.spec_from_file_location("vehicle_intersection_extract", MODULE_PATH)
vehicle = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = vehicle
SPEC.loader.exec_module(vehicle)


SAMPLE_ROWS = [
    ["구지사거리", "2026-04-10 07:00:00", "북", "직진", "승용차", 10],
    ["구지사거리", "2026-04-10 08:00:00", "남", "좌회전", "버스", 20],
]


class TestOutputFormatInput(unittest.TestCase):
    def test_accepts_text_and_number_choices(self):
        cases = {
            "": "csv",
            "csv": "csv",
            "1": "csv",
            "xlsx": "xlsx",
            "2": "xlsx",
            "db": "db",
            "3": "db",
        }
        for raw, expected in cases.items():
            with self.subTest(raw=raw), patch("builtins.input", return_value=raw):
                self.assertEqual(vehicle.input_output_format(), expected)

    def test_retries_invalid_input(self):
        with patch("builtins.input", side_effect=["bad", "2"]):
            self.assertEqual(vehicle.input_output_format(), "xlsx")


class TestOutputFiles(unittest.TestCase):
    def test_make_filename_returns_base_name_without_extension(self):
        intersections = [vehicle.Intersection(node_id=1, name="구지사거리")]
        filename = vehicle.make_filename(intersections, False, ["260410~260412"])
        self.assertEqual(filename, "구지사거리_260410~260412")

    def test_save_csv_writes_header_and_rows(self):
        with TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "traffic.csv"
            vehicle.save_csv(SAMPLE_ROWS, output_path)

            with output_path.open("r", newline="", encoding="utf-8-sig") as f:
                rows = list(csv.reader(f))

        self.assertEqual(rows[0], vehicle.CSV_HEADER)
        self.assertEqual(rows[1], [str(value) for value in SAMPLE_ROWS[0]])

    def test_save_sqlite_recreates_traffic_table(self):
        with TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "traffic.db"
            vehicle.save_sqlite(SAMPLE_ROWS[:1], output_path)
            vehicle.save_sqlite(SAMPLE_ROWS, output_path)

            conn = sqlite3.connect(output_path)
            try:
                count = conn.execute("SELECT COUNT(*) FROM traffic_by_vehicle").fetchone()[0]
                columns = [
                    row[1]
                    for row in conn.execute("PRAGMA table_info(traffic_by_vehicle)").fetchall()
                ]
            finally:
                conn.close()

        self.assertEqual(count, 2)
        self.assertEqual(
            columns,
            [
                "id",
                "intersection",
                "observed_at",
                "approach_direction",
                "direction",
                "vehicle_type",
                "traffic_volume",
            ],
        )

    def test_save_xlsx_writes_single_traffic_sheet(self):
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest("openpyxl is not installed")

        with TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "traffic.xlsx"
            vehicle.save_xlsx(SAMPLE_ROWS, output_path)
            workbook = load_workbook(output_path)
            try:
                sheet = workbook["교통량"]
                rows = list(sheet.iter_rows(values_only=True))
            finally:
                workbook.close()

        self.assertEqual(list(rows[0]), vehicle.CSV_HEADER)
        self.assertEqual(list(rows[1]), SAMPLE_ROWS[0])

    def test_save_xlsx_reports_missing_openpyxl(self):
        original_import = builtins.__import__

        def fake_import(name, *args, **kwargs):
            if name == "openpyxl":
                raise ImportError("missing openpyxl")
            return original_import(name, *args, **kwargs)

        with TemporaryDirectory() as temp_dir, patch("builtins.__import__", fake_import):
            output_path = Path(temp_dir) / "traffic.xlsx"
            with self.assertRaisesRegex(RuntimeError, "openpyxl"):
                vehicle.save_xlsx(SAMPLE_ROWS, output_path)


if __name__ == "__main__":
    unittest.main()
