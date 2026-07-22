#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Fill the Samjeong remicon Sheet1 traffic summary from SQLite data."""

from __future__ import annotations

import argparse
from collections.abc import Iterable
from copy import copy
from dataclasses import dataclass
from datetime import date, timedelta
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
import re
import sqlite3
from tempfile import NamedTemporaryFile
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_DB_PATH = PROJECT_ROOT / "00_Data" / "삼정동_레미콘" / "삼정동_레미콘_데이터.db"
DEFAULT_WORKBOOK_PATH = PROJECT_ROOT / "02_Result" / "09_기타" / "삼정동_레미콘_교통량_수정_v1.xlsx"
SHEET_NAME = "Sheet1"
TABLE_NAME = "vehicle_detection"
SPREADSHEETML_NAMESPACE = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
HOLIDAYS = {date(2026, 5, 1), date(2026, 5, 5), date(2026, 5, 25), date(2026, 6, 3)}
PEAK_HOURS = (7, 8, 17, 18)
REMICON_PATTERNS = (
    re.compile(r"^014[가-힣][0-9]\*{3}$"),
    re.compile(r"^[가-힣]{2,}14[가-힣][0-9]\*{3}$"),
)
LOCATION_BY_DIRECTION = {
    ("박촌교 삼거리", "남향"): "박촌교 삼거리[남향]",
    ("박촌교 삼거리", "북향"): "박촌교 삼거리[북향]",
    ("박촌교 삼거리", "서향(순방향)"): "박촌교 삼거리[서향(순방향)]",
    ("박촌교 삼거리", "서향(역방향)"): "박촌교 삼거리[서향(역방향)]",
    ("봉오고가교 사거리", "남향"): "봉오고가교 사거리[남향]",
    ("봉오고가교 사거리", "서향"): "봉오고가교 사거리[서향]",
    ("삼정고가 삼거리", "동향"): "삼정고가 삼거리[동향]",
    ("삼정고가 삼거리", "서향"): "삼정고가 삼거리[서향]",
    ("삼정교 사거리", "북동향"): "삼정교 사거리[북동향]",
    ("산업길 사거리", "남향"): "산업길 사거리[남향]",
    ("산업길 사거리", "북향"): "산업길 사거리[북향]",
    ("봉오대로 사거리", "동향"): "봉오대로 사거리[동향]",
    ("봉오대로 사거리", "서향"): "봉오대로 사거리[서향]",
    ("자동차검사소", "-"): "자동차검사소",
    ("삼정동 320-1", "-"): "삼정동320-1 부천IC",
}


@dataclass(frozen=True)
class Period:
    start: date
    end: date
    daily_divisor: int
    weekday_divisor: int


@dataclass(frozen=True)
class AverageDivisors:
    daily: int
    weekday: int


@dataclass(frozen=True)
class Counts:
    total: int
    remicon: int

    def rounded_average(self, divisor: int) -> tuple[int, int, float]:
        total_average = Decimal(self.total) / divisor
        remicon_average = Decimal(self.remicon) / divisor
        ratio = Decimal("0") if self.total == 0 else Decimal(self.remicon) * 100 / self.total
        return (
            int(total_average.quantize(Decimal("1"), rounding=ROUND_HALF_UP)),
            int(remicon_average.quantize(Decimal("1"), rounding=ROUND_HALF_UP)),
            float(ratio.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)),
        )


PERIOD_BY_MONTH = {
    "5월": Period(date(2026, 5, 1), date(2026, 6, 1), 31, 18),
    "6월": Period(date(2026, 6, 1), date(2026, 7, 1), 30, 21),
}
AVERAGE_DIVISORS_BY_MONTH_AND_LOCATION = {
    ("5월", "자동차검사소"): AverageDivisors(daily=25, weekday=15),
}
MULTI_DIRECTION_INTERSECTIONS = {
    "박촌교 삼거리",
    "봉오고가교 사거리",
    "삼정고가 삼거리",
    "산업길 사거리",
    "봉오대로 사거리",
}


@dataclass
class DirectionGroup:
    month: str
    intersection: str
    direction_rows: list[int]
    total_row: int | None = None

    @property
    def start_row(self) -> int:
        return self.direction_rows[0]

    @property
    def end_row(self) -> int:
        return self.direction_rows[-1]

    @property
    def is_multi_direction(self) -> bool:
        return len(self.direction_rows) >= 2


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--db", type=Path, default=DEFAULT_DB_PATH)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK_PATH)
    parser.add_argument("--verify-only", action="store_true")
    return parser.parse_args()


def is_remicon(vehicle_number: object) -> int:
    return int(any(pattern.fullmatch(str(vehicle_number or "")) for pattern in REMICON_PATTERNS))


def connect_readonly(db_path: Path) -> sqlite3.Connection:
    connection = sqlite3.connect(f"file:{db_path.resolve().as_posix()}?mode=ro", uri=True)
    connection.create_function("IS_REMICON", 1, is_remicon)
    return connection


def weekday_dates(period: Period) -> list[str]:
    dates: list[str] = []
    current = period.start
    while current < period.end:
        if current.weekday() < 5 and current not in HOLIDAYS:
            dates.append(current.isoformat())
        current += timedelta(days=1)
    if len(dates) != period.weekday_divisor:
        raise RuntimeError(f"Unexpected weekday count: {len(dates)}")
    return dates


def average_divisors(month: str, location: str, period: Period) -> AverageDivisors:
    return AVERAGE_DIVISORS_BY_MONTH_AND_LOCATION.get(
        (month, location),
        AverageDivisors(daily=period.daily_divisor, weekday=period.weekday_divisor),
    )


def query_counts(
    connection: sqlite3.Connection, locations: Iterable[str], period: Period
) -> tuple[dict[str, Counts], dict[str, Counts], dict[str, Counts]]:
    location_list = tuple(locations)
    placeholders = ", ".join("?" for _ in location_list)
    weekday_list = weekday_dates(period)
    weekday_placeholders = ", ".join("?" for _ in weekday_list)
    sql = f"""
        SELECT location_name,
               COUNT(*) AS monthly_total,
               SUM(IS_REMICON(vehicle_number)) AS monthly_remicon,
               SUM(CASE WHEN date(collected_at) IN ({weekday_placeholders}) THEN 1 ELSE 0 END),
               SUM(CASE WHEN date(collected_at) IN ({weekday_placeholders})
                         THEN IS_REMICON(vehicle_number) ELSE 0 END),
               SUM(CASE WHEN date(collected_at) IN ({weekday_placeholders})
                            AND CAST(strftime('%H', collected_at) AS INTEGER) IN (7, 8, 17, 18)
                         THEN 1 ELSE 0 END),
               SUM(CASE WHEN date(collected_at) IN ({weekday_placeholders})
                            AND CAST(strftime('%H', collected_at) AS INTEGER) IN (7, 8, 17, 18)
                         THEN IS_REMICON(vehicle_number) ELSE 0 END)
        FROM {TABLE_NAME}
        WHERE collected_at >= ? AND collected_at < ?
          AND location_name IN ({placeholders})
        GROUP BY location_name
    """
    values = (
        *weekday_list,
        *weekday_list,
        *weekday_list,
        *weekday_list,
        period.start.isoformat(),
        period.end.isoformat(),
        *location_list,
    )
    monthly: dict[str, Counts] = {}
    weekday: dict[str, Counts] = {}
    peak: dict[str, Counts] = {}
    for row in connection.execute(sql, values):
        location, *counts = row
        monthly[str(location)] = Counts(int(counts[0] or 0), int(counts[1] or 0))
        weekday[str(location)] = Counts(int(counts[2] or 0), int(counts[3] or 0))
        peak[str(location)] = Counts(int(counts[4] or 0), int(counts[5] or 0))
    missing = set(location_list) - monthly.keys()
    if missing:
        raise RuntimeError(f"No DB data for locations: {sorted(missing)}")
    return monthly, weekday, peak


def find_direction_groups(worksheet: object) -> list[DirectionGroup]:
    current_month = ""
    current_intersection = ""
    groups: list[DirectionGroup] = []
    group_by_key: dict[tuple[str, str], DirectionGroup] = {}
    for row in range(4, worksheet.max_row + 1):
        month = worksheet.cell(row=row, column=1).value
        if month:
            current_month = str(month)
        intersection = worksheet.cell(row=row, column=2).value
        if intersection:
            current_intersection = str(intersection)
        direction = str(worksheet.cell(row=row, column=3).value or "")
        if not direction:
            continue
        if direction == "합계":
            if groups:
                groups[-1].total_row = row
            continue
        if current_month not in PERIOD_BY_MONTH:
            raise RuntimeError(f"Unexpected month at worksheet row {row}: {current_month!r}")
        if not current_intersection:
            raise RuntimeError(f"Missing intersection at worksheet row {row}")
        key = (current_month, current_intersection)
        group = group_by_key.get(key)
        if group is None:
            group = DirectionGroup(current_month, current_intersection, [])
            groups.append(group)
            group_by_key[key] = group
        group.direction_rows.append(row)
    return groups


def row_locations(worksheet: object) -> dict[int, str]:
    result: dict[int, str] = {}
    for group in find_direction_groups(worksheet):
        for row in group.direction_rows:
            direction = str(worksheet.cell(row=row, column=3).value)
            try:
                result[row] = LOCATION_BY_DIRECTION[(group.intersection, direction)]
            except KeyError as error:
                raise RuntimeError(
                    f"Unmapped worksheet row {row}: {group.intersection} / {direction}"
                ) from error
    return result


def expected_values(workbook_path: Path, db_path: Path) -> dict[int, tuple[object, ...]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        worksheet = workbook[SHEET_NAME]
        row_map = row_locations(worksheet)
        groups = find_direction_groups(worksheet)
    finally:
        workbook.close()
    results: dict[int, tuple[object, ...]] = {}
    connection = connect_readonly(db_path)
    try:
        for month, period in PERIOD_BY_MONTH.items():
            month_rows = [
                row for group in groups if group.month == month for row in group.direction_rows
            ]
            month_locations = [row_map[row] for row in month_rows]
            monthly, weekday, peak = query_counts(connection, month_locations, period)
            for row in month_rows:
                location = row_map[row]
                divisors = average_divisors(month, location, period)
                results[row] = (
                    *monthly[location].rounded_average(divisors.daily),
                    *weekday[location].rounded_average(divisors.weekday),
                    *peak[location].rounded_average(divisors.weekday),
                )
    finally:
        connection.close()
    return results


def unmerge_data_labels(worksheet: object) -> None:
    for merged_range in list(worksheet.merged_cells.ranges):
        if merged_range.max_row >= 4 and merged_range.min_col <= 2:
            worksheet.unmerge_cells(str(merged_range))


def copy_row_style(worksheet: object, source_row: int, target_row: int) -> None:
    worksheet.row_dimensions[target_row].height = worksheet.row_dimensions[source_row].height
    for column in range(1, 15):
        source = worksheet.cell(row=source_row, column=column)
        target = worksheet.cell(row=target_row, column=column)
        target._style = copy(source._style)
        if source.has_style:
            target.number_format = source.number_format
        if source.hyperlink:
            target._hyperlink = copy(source.hyperlink)


def rebuild_group_labels_and_merges(worksheet: object) -> None:
    groups = find_direction_groups(worksheet)
    for row in range(4, worksheet.max_row + 1):
        worksheet.cell(row=row, column=1).value = None
        worksheet.cell(row=row, column=2).value = None

    month_groups: dict[str, list[DirectionGroup]] = {}
    for group in groups:
        month_groups.setdefault(group.month, []).append(group)
        worksheet.cell(row=group.start_row, column=2).value = group.intersection
        group_end = group.total_row or group.end_row
        if group_end > group.start_row:
            worksheet.merge_cells(
                start_row=group.start_row, start_column=2, end_row=group_end, end_column=2
            )

    for month, month_group_list in month_groups.items():
        start_row = month_group_list[0].start_row
        end_row = month_group_list[-1].total_row or month_group_list[-1].end_row
        worksheet.cell(row=start_row, column=1).value = month
        worksheet.merge_cells(start_row=start_row, start_column=1, end_row=end_row, end_column=1)


def configure_total_row(worksheet: object, group: DirectionGroup) -> None:
    if group.total_row is None:
        raise RuntimeError(f"Missing total row for {group.month} {group.intersection}")
    row = group.total_row
    start_row = group.start_row
    end_row = group.end_row
    worksheet.cell(row=row, column=3).value = "합계"
    worksheet.cell(row=row, column=4).value = None
    worksheet.cell(row=row, column=5).value = None
    for total_column, remicon_column in ((6, 7), (9, 10), (12, 13)):
        worksheet.cell(
            row=row, column=total_column
        ).value = f"=SUM({worksheet.cell(start_row, total_column).coordinate}:{worksheet.cell(end_row, total_column).coordinate})"
        worksheet.cell(
            row=row, column=remicon_column
        ).value = f"=SUM({worksheet.cell(start_row, remicon_column).coordinate}:{worksheet.cell(end_row, remicon_column).coordinate})"
        worksheet.cell(row=row, column=total_column).number_format = "#,##0"
        worksheet.cell(row=row, column=remicon_column).number_format = "#,##0"
    for total_column, remicon_column, ratio_column in ((6, 7, 8), (9, 10, 11), (12, 13, 14)):
        worksheet.cell(row=row, column=ratio_column).value = (
            f"=IF({worksheet.cell(row, total_column).coordinate}=0,0,"
            f"{worksheet.cell(row, remicon_column).coordinate}/{worksheet.cell(row, total_column).coordinate}*100)"
        )
        worksheet.cell(row=row, column=ratio_column).number_format = "0.00"
    for column in range(1, 15):
        cell = worksheet.cell(row=row, column=column)
        font = copy(cell.font)
        font.bold = True
        cell.font = font
        border = copy(cell.border)
        border.top = copy(worksheet.cell(row=end_row, column=column).border.top)
        cell.border = border


def ensure_total_rows(worksheet: object) -> None:
    unmerge_data_labels(worksheet)
    groups = find_direction_groups(worksheet)
    missing_groups = [
        group
        for group in groups
        if group.is_multi_direction
        and (
            group.end_row == worksheet.max_row
            or worksheet.cell(group.end_row + 1, 3).value != "합계"
        )
    ]
    for group in reversed(missing_groups):
        total_row = group.end_row + 1
        worksheet.insert_rows(total_row)
        copy_row_style(worksheet, total_row - 1, total_row)
        worksheet.cell(row=total_row, column=3).value = "합계"

    groups = find_direction_groups(worksheet)
    for group in groups:
        if group.is_multi_direction:
            configure_total_row(worksheet, group)
    rebuild_group_labels_and_merges(worksheet)


def verify_total_rows(worksheet: object) -> None:
    groups = find_direction_groups(worksheet)
    total_groups = [group for group in groups if group.total_row is not None]
    if len(total_groups) != 10:
        raise RuntimeError(f"Expected 10 total rows, found {len(total_groups)}")
    total_group_keys = {(group.month, group.intersection) for group in total_groups}
    expected_total_group_keys = {
        (month, intersection)
        for month in PERIOD_BY_MONTH
        for intersection in MULTI_DIRECTION_INTERSECTIONS
    }
    if total_group_keys != expected_total_group_keys:
        raise RuntimeError(f"Unexpected total groups: {sorted(total_group_keys)}")
    for group in groups:
        if group.is_multi_direction != (group.total_row is not None):
            raise RuntimeError(f"Incorrect total row for {group.month} {group.intersection}")
        if not group.is_multi_direction:
            continue
        total_row = group.total_row
        if total_row is None:
            raise RuntimeError("Unreachable missing total row")
        if (
            worksheet.cell(total_row, 4).value is not None
            or worksheet.cell(total_row, 5).value is not None
        ):
            raise RuntimeError(f"Total row D:E must be blank at row {total_row}")
        for total_column, remicon_column, ratio_column in ((6, 7, 8), (9, 10, 11), (12, 13, 14)):
            total_formula = f"=SUM({worksheet.cell(group.start_row, total_column).coordinate}:{worksheet.cell(group.end_row, total_column).coordinate})"
            remicon_formula = f"=SUM({worksheet.cell(group.start_row, remicon_column).coordinate}:{worksheet.cell(group.end_row, remicon_column).coordinate})"
            ratio_formula = (
                f"=IF({worksheet.cell(total_row, total_column).coordinate}=0,0,"
                f"{worksheet.cell(total_row, remicon_column).coordinate}/{worksheet.cell(total_row, total_column).coordinate}*100)"
            )
            if worksheet.cell(total_row, total_column).value != total_formula:
                raise RuntimeError(f"Unexpected total formula at row {total_row}")
            if worksheet.cell(total_row, remicon_column).value != remicon_formula:
                raise RuntimeError(f"Unexpected remicon formula at row {total_row}")
            if worksheet.cell(total_row, ratio_column).value != ratio_formula:
                raise RuntimeError(f"Unexpected ratio formula at row {total_row}")
            if worksheet.cell(total_row, ratio_column).number_format != "0.00":
                raise RuntimeError(f"Unexpected ratio format at row {total_row}")
        if not worksheet.cell(total_row, 3).font.bold:
            raise RuntimeError(f"Total row must be bold at row {total_row}")
        if worksheet.cell(total_row, 3).border.top.style is None:
            raise RuntimeError(f"Total row must have a top border at row {total_row}")
        expected_merge = f"B{group.start_row}:B{total_row}"
        if expected_merge not in {str(merged) for merged in worksheet.merged_cells.ranges}:
            raise RuntimeError(f"Missing group merge {expected_merge}")

    expected_month_merges = {
        f"A{min(group.start_row for group in groups if group.month == month)}:"
        f"A{max((group.total_row or group.end_row) for group in groups if group.month == month)}"
        for month in PERIOD_BY_MONTH
    }
    actual_month_merges = {
        str(merged)
        for merged in worksheet.merged_cells.ranges
        if merged.min_col == 1 and merged.max_col == 1 and merged.min_row >= 4
    }
    if actual_month_merges != expected_month_merges:
        raise RuntimeError(f"Unexpected month merges: {sorted(actual_month_merges)}")


def suppress_total_formula_indicators(workbook_path: Path, total_rows: list[int]) -> None:
    """Suppress Excel's inconsistent-formula indicator only for total formulas."""
    worksheet_xml_path = "xl/worksheets/sheet1.xml"
    namespace = f"{{{SPREADSHEETML_NAMESPACE}}}"
    formula_cells = " ".join(f"F{row}:N{row}" for row in total_rows)
    with ZipFile(workbook_path, "r") as source:
        worksheet_xml = source.read(worksheet_xml_path)
        entries = [(entry, source.read(entry.filename)) for entry in source.infolist()]

    root = ElementTree.fromstring(worksheet_xml)
    ignored_errors = root.find(f"{namespace}ignoredErrors")
    if ignored_errors is not None:
        root.remove(ignored_errors)
    ignored_errors = ElementTree.Element(f"{namespace}ignoredErrors")
    ElementTree.SubElement(
        ignored_errors,
        f"{namespace}ignoredError",
        {"sqref": formula_cells, "formula": "1"},
    )
    insert_before = {
        f"{namespace}{tag}"
        for tag in (
            "smartTags",
            "drawing",
            "legacyDrawing",
            "legacyDrawingHF",
            "picture",
            "oleObjects",
            "controls",
            "webPublishItems",
            "tableParts",
            "extLst",
        )
    }
    insert_index = next(
        (index for index, child in enumerate(root) if child.tag in insert_before),
        len(root),
    )
    root.insert(insert_index, ignored_errors)

    with NamedTemporaryFile(dir=workbook_path.parent, suffix=".xlsx", delete=False) as temporary:
        temporary_path = Path(temporary.name)
    try:
        with ZipFile(temporary_path, "w", ZIP_DEFLATED) as destination:
            for entry, content in entries:
                if entry.filename == worksheet_xml_path:
                    content = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)
                destination.writestr(entry, content)
        temporary_path.replace(workbook_path)
    finally:
        temporary_path.unlink(missing_ok=True)


def apply_or_verify(
    workbook_path: Path, results: dict[int, tuple[object, ...]], verify_only: bool
) -> None:
    workbook = load_workbook(workbook_path, data_only=False)
    try:
        worksheet = workbook[SHEET_NAME]
        for row, values in results.items():
            actual = tuple(worksheet.cell(row=row, column=column).value for column in range(6, 15))
            if verify_only:
                if actual != values:
                    raise RuntimeError(f"Mismatch in Sheet1 row {row}: {actual!r} != {values!r}")
                continue
            for column in (4, 5):
                worksheet.cell(row=row, column=column).value = None
            for column, value in enumerate(values, start=6):
                cell = worksheet.cell(row=row, column=column)
                cell.value = value
                cell.number_format = "0.00" if column in (8, 11, 14) else "#,##0"
        if verify_only:
            verify_total_rows(worksheet)
        else:
            ensure_total_rows(worksheet)
        if not verify_only:
            total_rows = [
                group.total_row
                for group in find_direction_groups(worksheet)
                if group.total_row is not None
            ]
            workbook.save(workbook_path)
            suppress_total_formula_indicators(workbook_path, total_rows)
    finally:
        workbook.close()


def main() -> int:
    args = parse_args()
    results = expected_values(args.workbook, args.db)
    apply_or_verify(args.workbook, results, args.verify_only)
    for row, values in results.items():
        print(row, *values, sep="\t")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
