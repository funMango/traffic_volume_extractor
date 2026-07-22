#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Compare Samjeong edge detections with smart-intersection raw traffic.

The domain calculation is independent of SQLite, Oracle, and Excel.  The
adapters in this module are deliberately read-only; the presentation layer is
responsible for applying the already-verified values to a workbook.
"""

from __future__ import annotations

import argparse
import json
import shutil
import sqlite3
import sys
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Iterable, Protocol

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parents[2]
EDGE_DB_PATH = BASE_DIR / "00_Data" / "삼정동_레미콘" / "삼정동_레미콘_데이터.db"
TEMPLATE_PATH = BASE_DIR / "00_Data" / "삼정동_레미콘" / "삼정동_레미콘_교통량_작성서식.xlsx"
DEFAULT_OUTPUT_PATH = BASE_DIR / "02_Result" / "09_기타" / "삼정동_교차로_비교_작성서식.xlsx"
V2_WORKBOOK_PATH = BASE_DIR / "02_Result" / "09_기타" / "삼정동_레미콘_교통량_v2.xlsx"
HOLIDAYS = {date(2026, 5, 1), date(2026, 5, 5), date(2026, 5, 25), date(2026, 6, 3)}
MONTHS = (5, 6)
PEAK_HOURS = {7, 8, 17, 18}
MONTHLY_DIVISORS = {5: 31, 6: 30}
WEEKDAY_DIVISORS = {5: 18, 6: 21}
V2_SHEET_NAME = "스마트교차로 비교"
V2_TARGET_COLUMNS = {"monthly": 4, "weekday": 7, "peak": 10}
V2_TARGET_ROWS = range(4, 16)


@dataclass(frozen=True)
class ComparisonTarget:
    label: str
    edge_locations: tuple[str, ...]
    smart_intersection: str
    smart_approaches: tuple[str, ...]


TARGETS = (
    ComparisonTarget(
        "박촌교 삼거리",
        ("박촌교 삼거리[남향]", "박촌교 삼거리[북향]"),
        "박촌교삼거리",
        ("박촌교삼거리-북 (남향)", "박촌교삼거리-남 (북향)"),
    ),
    ComparisonTarget(
        "봉오고가교사거리",
        ("봉오고가교 사거리[남향]", "봉오고가교 사거리[서향]"),
        "봉오고가교사거리",
        ("봉오고가교사거리-북 (남향)", "봉오고가교사거리-동 (서향)"),
    ),
    ComparisonTarget(
        "삼정고가삼거리",
        ("삼정고가 삼거리[동향]", "삼정고가 삼거리[서향]"),
        "삼정고가교삼거리",
        ("삼정고가교삼거리-서 (동향)", "삼정고가교삼거리-동 (서향)"),
    ),
    ComparisonTarget(
        "삼정교사거리", ("삼정교 사거리[북동향]",), "삼정교사거리", ("삼정교사거리-서 (동향)",)
    ),
    ComparisonTarget(
        "산업길사거리",
        ("산업길 사거리[남향]", "산업길 사거리[북향]"),
        "산업길사거리",
        ("산업길사거리-북 (남향)", "산업길사거리-남 (북향)"),
    ),
    ComparisonTarget(
        "봉오대로사거리",
        ("봉오대로 사거리[동향]", "봉오대로 사거리[서향]"),
        "봉오대로사거리",
        ("봉오대로사거리-서 (동향)", "봉오대로4R-동 (서향)"),
    ),
)


@dataclass(frozen=True)
class DailyTraffic:
    volume: int
    collected: bool


@dataclass(frozen=True)
class Metrics:
    monthly: int
    weekday: int
    weekday_peak: int
    monthly_days: int
    weekday_days: int
    weekday_peak_days: int


class EdgeDailyTrafficReader(Protocol):
    def load(
        self, target: ComparisonTarget, start: date, end: date
    ) -> dict[date, DailyTraffic]: ...


class SmartDailyTrafficReader(Protocol):
    def load(
        self, target: ComparisonTarget, start: date, end: date
    ) -> dict[date, DailyTraffic]: ...


def calendar_days(start: date, end: date) -> Iterable[date]:
    current = start
    while current <= end:
        yield current
        current += timedelta(days=1)


def is_weekday(day: date) -> bool:
    return day.weekday() < 5 and day not in HOLIDAYS


def round_half_up(value: Decimal) -> int:
    return int(value.quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def average_daily(
    values: dict[date, DailyTraffic], days: Iterable[date], require_collected: bool
) -> tuple[int, int]:
    selected = list(days)
    if require_collected:
        selected = [day for day in selected if values.get(day, DailyTraffic(0, False)).collected]
    total = sum(values.get(day, DailyTraffic(0, False)).volume for day in selected)
    return (round_half_up(Decimal(total) / len(selected)), len(selected)) if selected else (0, 0)


def calculate_metrics(
    values: dict[date, DailyTraffic], month: int, require_collected: bool
) -> Metrics:
    start = date(2026, month, 1)
    end = date(2026, month + 1, 1) - timedelta(days=1) if month < 12 else date(2026, 12, 31)
    all_days = list(calendar_days(start, end))
    weekdays = [day for day in all_days if is_weekday(day)]
    monthly, monthly_days = average_daily(values, all_days, require_collected)
    weekday, weekday_days = average_daily(values, weekdays, require_collected)
    peak, peak_days = average_daily(values, weekdays, require_collected)
    return Metrics(monthly, weekday, peak, monthly_days, weekday_days, peak_days)


class SQLiteEdgeReader:
    """Reads all selected edge records; absent dates deliberately stay in the denominator."""

    def __init__(self, database_path: Path):
        self._database_path = database_path

    def load(self, target: ComparisonTarget, start: date, end: date) -> dict[date, DailyTraffic]:
        database = self._database_path.resolve()
        conn = sqlite3.connect(f"file:{database.as_posix()}?mode=ro", uri=True)
        try:
            placeholders = ", ".join("?" for _ in target.edge_locations)
            rows = conn.execute(
                f"""
                SELECT substr(collected_at, 1, 10),
                       SUM(CASE WHEN substr(collected_at, 12, 2) IN ('07', '08', '17', '18')
                                THEN 1 ELSE 0 END),
                       COUNT(*)
                FROM vehicle_detection
                WHERE location_name IN ({placeholders})
                  AND collected_at >= ? AND collected_at < ?
                GROUP BY substr(collected_at, 1, 10)
                """,
                (*target.edge_locations, start.isoformat(), (end + timedelta(days=1)).isoformat()),
            ).fetchall()
        finally:
            conn.close()
        result = {day: DailyTraffic(0, False) for day in calendar_days(start, end)}
        for raw_day, peak_count, total_count in rows:
            day = date.fromisoformat(str(raw_day))
            result[day] = DailyTraffic(int(total_count or 0), True)
            result[(day, "peak")] = DailyTraffic(int(peak_count or 0), True)  # type: ignore[index]
        return result


class OracleSmartReader:
    """Reads Oracle raw hourly direction records inside a read-only transaction."""

    def __init__(self) -> None:
        sys.path.insert(0, str(BASE_DIR / "01_Program" / "09_else"))
        from traffic_api.infrastructure import oracle_analysis

        self._oracle = oracle_analysis

    def _resolve(self, conn, target: ComparisonTarget) -> tuple[str, tuple[str, ...]]:
        with conn.cursor() as cursor:
            cursor.execute(
                "SELECT NODE_ID FROM M_CRSRD_INF WHERE CRSRD_NM = :name",
                name=target.smart_intersection,
            )
            nodes = cursor.fetchall()
            if len(nodes) != 1:
                raise RuntimeError(
                    f"스마트교차로명은 하나여야 합니다: {target.smart_intersection} ({len(nodes)})"
                )
            node_id = str(nodes[0][0]).strip()
            cursor.execute(
                "SELECT ACSR_ID, ACSR_NM FROM M_CRSRD_ACSR_INF WHERE NODE_ID = :node_id",
                node_id=node_id,
            )
            approaches = [(str(row[0]).strip(), str(row[1]).strip()) for row in cursor.fetchall()]
        ids: list[str] = []
        for name in target.smart_approaches:
            matches = [
                approach_id for approach_id, approach_name in approaches if approach_name == name
            ]
            if len(matches) != 1:
                raise RuntimeError(
                    f"접근로명은 하나여야 합니다: {target.smart_intersection} {name} ({len(matches)})"
                )
            ids.append(matches[0])
        return node_id, tuple(ids)

    def load(self, target: ComparisonTarget, start: date, end: date) -> dict[date, DailyTraffic]:
        conn = self._oracle.connect_db()
        try:
            with conn.cursor() as cursor:
                cursor.execute("SET TRANSACTION READ ONLY")
            node_id, approach_ids = self._resolve(conn, target)
            bind = {
                "node_id": node_id,
                "start_time": datetime.combine(start, datetime.min.time()),
                "end_time": datetime.combine(end + timedelta(days=1), datetime.min.time()),
            }
            names = []
            for index, approach_id in enumerate(approach_ids):
                key = f"approach_{index}"
                bind[key] = approach_id
                names.append(f":{key}")
            sql = f"""
                SELECT TRUNC(TOT_DT), ACSR_ID,
                       SUM(NVL(TRF_QNTY, 0)),
                       SUM(CASE WHEN TO_NUMBER(TO_CHAR(TOT_DT, 'HH24')) IN (7, 8, 17, 18)
                                THEN NVL(TRF_QNTY, 0) ELSE 0 END)
                FROM S_CRSRD_DRCT_TRF_1HH
                WHERE NODE_ID = :node_id
                  AND ACSR_ID IN ({", ".join(names)})
                  AND TOT_DT >= :start_time AND TOT_DT < :end_time
                GROUP BY TRUNC(TOT_DT), ACSR_ID
            """
            with conn.cursor() as cursor:
                cursor.execute(sql, bind)
                rows = cursor.fetchall()
        finally:
            conn.close()
        grouped: dict[date, dict[str, tuple[int, int]]] = defaultdict(dict)
        for raw_day, approach_id, total, peak in rows:
            day = raw_day.date() if isinstance(raw_day, datetime) else raw_day
            grouped[day][str(approach_id).strip()] = (int(total or 0), int(peak or 0))
        result: dict[date, DailyTraffic] = {}
        for day in calendar_days(start, end):
            day_rows = grouped.get(day, {})
            collected = all(approach_id in day_rows for approach_id in approach_ids)
            result[day] = DailyTraffic(
                sum(day_rows[item][0] for item in approach_ids) if collected else 0, collected
            )
            result[(day, "peak")] = DailyTraffic(
                sum(day_rows[item][1] for item in approach_ids) if collected else 0, collected
            )  # type: ignore[index]
        return result


class OracleIntersectionHourlyReader:
    """Reads selected smart-intersection approaches by day and hour."""

    def __init__(self) -> None:
        sys.path.insert(0, str(BASE_DIR / "01_Program" / "09_else"))
        from traffic_api.infrastructure import oracle_analysis

        self._oracle = oracle_analysis

    def load(self) -> dict[str, dict[date, dict[int, int]]]:
        names = {target.label: target.smart_intersection for target in TARGETS}
        conn = self._oracle.connect_db()
        try:
            with conn.cursor() as cursor:
                cursor.execute("SET TRANSACTION READ ONLY")
                bind = {f"name_{index}": name for index, name in enumerate(names.values())}
                placeholders = ", ".join(f":{key}" for key in bind)
                cursor.execute(
                    f"""
                    SELECT NODE_ID, CRSRD_NM
                    FROM M_CRSRD_INF
                    WHERE CRSRD_NM IN ({placeholders})
                    """,
                    bind,
                )
                nodes = [
                    (str(node_id).strip(), str(name).strip()) for node_id, name in cursor.fetchall()
                ]
                by_name = {name: node_id for node_id, name in nodes}
                if set(by_name) != set(names.values()) or len(nodes) != len(names):
                    raise RuntimeError(f"스마트교차로 NODE_ID 확인 실패: {nodes}")
                selected_pairs: dict[tuple[str, str], str] = {}
                for label, intersection_name in names.items():
                    node_id = by_name[intersection_name]
                    target = next(target for target in TARGETS if target.label == label)
                    cursor.execute(
                        "SELECT ACSR_ID, ACSR_NM FROM M_CRSRD_ACSR_INF WHERE NODE_ID = :node_id",
                        node_id=node_id,
                    )
                    approaches = {
                        str(approach_name).strip(): str(approach_id).strip()
                        for approach_id, approach_name in cursor.fetchall()
                    }
                    for approach_name in target.smart_approaches:
                        if approach_name not in approaches:
                            raise RuntimeError(
                                f"스마트교차로 접근로 확인 실패: {intersection_name} / {approach_name}"
                            )
                        selected_pairs[(node_id, approaches[approach_name])] = label

                query_bind: dict[str, object] = {
                    "start_time": datetime(2026, 5, 1),
                    "end_time": datetime(2026, 7, 1),
                }
                pair_clauses = []
                for index, (node_id, approach_id) in enumerate(selected_pairs):
                    node_key, approach_key = f"node_{index}", f"approach_{index}"
                    query_bind[node_key] = node_id
                    query_bind[approach_key] = approach_id
                    pair_clauses.append(f"(NODE_ID = :{node_key} AND ACSR_ID = :{approach_key})")
                cursor.execute(
                    f"""
                    SELECT NODE_ID,
                           ACSR_ID,
                           TRUNC(TOT_DT),
                           TO_NUMBER(TO_CHAR(TOT_DT, 'HH24')),
                           SUM(NVL(TRF_QNTY, 0))
                    FROM S_CRSRD_ACSR_TRF_1HH
                    WHERE ({" OR ".join(pair_clauses)})
                      AND TOT_DT >= :start_time AND TOT_DT < :end_time
                    GROUP BY NODE_ID, ACSR_ID, TRUNC(TOT_DT), TO_NUMBER(TO_CHAR(TOT_DT, 'HH24'))
                    """,
                    query_bind,
                )
                hourly_rows = cursor.fetchall()
        finally:
            conn.close()
        result = {label: defaultdict(dict) for label in names}
        for node_id, approach_id, raw_day, hour, volume in hourly_rows:
            day = raw_day.date() if isinstance(raw_day, datetime) else raw_day
            label = selected_pairs[(str(node_id).strip(), str(approach_id).strip())]
            hourly = result[label][day]
            hourly[int(hour)] = hourly.get(int(hour), 0) + max(int(volume or 0), 0)
        return result


def is_hourly_day_valid(hours: dict[int, int]) -> bool:
    """A date is usable when fewer than four hourly records are null, zero, or absent."""
    invalid_count = sum(hours.get(hour, 0) <= 0 for hour in range(24))
    return invalid_count < 4


def average_hourly_days(
    values: dict[date, dict[int, int]], days: Iterable[date]
) -> tuple[int, int]:
    valid = [day for day in days if is_hourly_day_valid(values.get(day, {}))]
    total = sum(sum(max(values[day].get(hour, 0), 0) for hour in range(24)) for day in valid)
    return (round_half_up(Decimal(total) / len(valid)), len(valid)) if valid else (0, 0)


def average_weekday_peak_hours(
    values: dict[date, dict[int, int]], days: Iterable[date], divisor: int
) -> tuple[int, int]:
    """Return the fixed-divisor average of the four weekday peak hours."""
    total = sum(max(values.get(day, {}).get(hour, 0), 0) for day in days for hour in PEAK_HOURS)
    return round_half_up(Decimal(total) / divisor), divisor


def average_hourly_values(
    values: dict[date, dict[int, int]], days: Iterable[date], divisor: int
) -> tuple[int, int]:
    total = sum(sum(max(volume, 0) for volume in values.get(day, {}).values()) for day in days)
    return round_half_up(Decimal(total) / divisor), divisor


def smart_v2_payload(
    reader: OracleIntersectionHourlyReader,
) -> dict[tuple[int, str], dict[str, int]]:
    source = reader.load()
    payload: dict[tuple[int, str], dict[str, int]] = {}
    for month in MONTHS:
        start = date(2026, month, 1)
        end = date(2026, month + 1, 1) - timedelta(days=1)
        month_days = list(calendar_days(start, end))
        weekday_days = [day for day in month_days if is_weekday(day)]
        for target in TARGETS:
            values = source[target.label]
            monthly, monthly_days = average_hourly_values(
                values, month_days, MONTHLY_DIVISORS[month]
            )
            weekday, weekday_count = average_hourly_values(
                values, weekday_days, WEEKDAY_DIVISORS[month]
            )
            peak, peak_count = average_weekday_peak_hours(
                values, weekday_days, WEEKDAY_DIVISORS[month]
            )
            payload[(month, target.label)] = {
                "monthly": monthly,
                "weekday": weekday,
                "peak": peak,
                "monthly_days": monthly_days,
                "weekday_days": weekday_count,
                "peak_days": peak_count,
            }
    return payload


def write_v2_workbook(workbook_path: Path, payload: dict[tuple[int, str], dict[str, int]]) -> None:
    workbook = load_workbook(workbook_path)
    try:
        worksheet = workbook[V2_SHEET_NAME]
        for row, month in ((row, 5 if row <= 9 else 6) for row in V2_TARGET_ROWS):
            label = str(worksheet.cell(row, 2).value or "").strip()
            values = payload[(month, label)]
            for metric, column in V2_TARGET_COLUMNS.items():
                cell = worksheet.cell(row, column)
                cell.value = values[metric]
                cell.number_format = "#,##0"
        workbook.save(workbook_path)
    finally:
        workbook.close()


def v2_non_target_snapshot(
    workbook_path: Path,
) -> dict[tuple[str, str], tuple[object, str, int, str]]:
    workbook = load_workbook(workbook_path, data_only=False)
    try:
        target_cells = {
            (V2_SHEET_NAME, f"{chr(64 + column)}{row}")
            for row in V2_TARGET_ROWS
            for column in V2_TARGET_COLUMNS.values()
        }
        return {
            (worksheet.title, cell.coordinate): (
                cell.value,
                cell.data_type,
                cell.style_id,
                cell.number_format,
            )
            for worksheet in workbook.worksheets
            for row in worksheet.iter_rows()
            for cell in row
            if (worksheet.title, cell.coordinate) not in target_cells
        }
    finally:
        workbook.close()


def verify_v2_workbook(
    workbook_path: Path,
    payload: dict[tuple[int, str], dict[str, int]],
    expected_non_target: dict[tuple[str, str], tuple[object, str, int, str]],
) -> None:
    workbook = load_workbook(workbook_path, data_only=False)
    try:
        worksheet = workbook[V2_SHEET_NAME]
        for row, month in ((row, 5 if row <= 9 else 6) for row in V2_TARGET_ROWS):
            label = str(worksheet.cell(row, 2).value or "").strip()
            expected = payload[(month, label)]
            for metric, column in V2_TARGET_COLUMNS.items():
                cell = worksheet.cell(row, column)
                if cell.value != expected[metric] or cell.number_format != "#,##0":
                    raise AssertionError(f"스마트교차로 비교 불일치: {cell.coordinate}")
        actual_non_target = {
            (sheet.title, cell.coordinate): (
                cell.value,
                cell.data_type,
                cell.style_id,
                cell.number_format,
            )
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if not (
                sheet.title == V2_SHEET_NAME
                and cell.column in V2_TARGET_COLUMNS.values()
                and cell.row in V2_TARGET_ROWS
            )
        }
        if actual_non_target != expected_non_target:
            changed = sorted(
                key
                for key in actual_non_target | expected_non_target
                if actual_non_target.get(key) != expected_non_target.get(key)
            )
            raise AssertionError(f"스마트교차로(대) 외 셀 변경: {changed[:10]}")
    finally:
        workbook.close()


def result_payload(
    edge_reader: EdgeDailyTrafficReader, smart_reader: SmartDailyTrafficReader
) -> dict:
    rows = []
    for month in MONTHS:
        start = date(2026, month, 1)
        end = date(2026, month + 1, 1) - timedelta(days=1)
        for target in TARGETS:
            edge = edge_reader.load(target, start, end)
            smart = smart_reader.load(target, start, end)
            edge_metrics = calculate_metrics(edge, month, False)
            smart_metrics = calculate_metrics(smart, month, True)
            # Peak is calculated from the separate 07-09/17-19 daily series.
            edge_peak, edge_peak_days = average_daily(
                {
                    day: edge.get((day, "peak"), DailyTraffic(0, False))
                    for day in calendar_days(start, end)
                },
                [day for day in calendar_days(start, end) if is_weekday(day)],
                False,
            )
            smart_peak, smart_peak_days = average_daily(
                {
                    day: smart.get((day, "peak"), DailyTraffic(0, False))
                    for day in calendar_days(start, end)
                },
                [day for day in calendar_days(start, end) if is_weekday(day)],
                True,
            )
            rows.append(
                {
                    "month": month,
                    "label": target.label,
                    "edge": {
                        "monthly": edge_metrics.monthly,
                        "weekday": edge_metrics.weekday,
                        "peak": edge_peak,
                        "denominators": [
                            edge_metrics.monthly_days,
                            edge_metrics.weekday_days,
                            edge_peak_days,
                        ],
                    },
                    "smart": {
                        "monthly": smart_metrics.monthly,
                        "weekday": smart_metrics.weekday,
                        "peak": smart_peak,
                        "denominators": [
                            smart_metrics.monthly_days,
                            smart_metrics.weekday_days,
                            smart_peak_days,
                        ],
                    },
                }
            )
    return {"rows": rows}


def write_template(template_path: Path, output_path: Path, payload: dict) -> None:
    """Copy the source layout and change only the requested C4:H15 value cells."""
    if not template_path.exists():
        raise FileNotFoundError(template_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if template_path.resolve() != output_path.resolve():
        shutil.copy2(template_path, output_path)
    workbook = load_workbook(output_path)
    try:
        worksheet = workbook["작성서식"]
        rows_by_key = {(int(item["month"]), str(item["label"])): item for item in payload["rows"]}
        for month, row, label in _find_target_rows(worksheet, rows_by_key):
            item = rows_by_key[(month, label)]
            values = (
                item["edge"]["monthly"],
                item["smart"]["monthly"],
                item["edge"]["weekday"],
                item["smart"]["weekday"],
                item["edge"]["peak"],
                item["smart"]["peak"],
            )
            for column, value in enumerate(values, start=3):
                cell = worksheet.cell(row, column)
                cell.value = value
                cell.number_format = "#,##0"
        workbook.save(output_path)
    finally:
        workbook.close()


def verify_template(output_path: Path, payload: dict) -> None:
    workbook = load_workbook(output_path, data_only=True)
    try:
        worksheet = workbook["작성서식"]
        rows_by_key = {(int(item["month"]), str(item["label"])): item for item in payload["rows"]}
        for month, row, label in _find_target_rows(worksheet, rows_by_key):
            item = rows_by_key[(month, label)]
            actual = [worksheet.cell(row, column).value for column in range(3, 9)]
            expected = [
                item["edge"]["monthly"],
                item["smart"]["monthly"],
                item["edge"]["weekday"],
                item["smart"]["weekday"],
                item["edge"]["peak"],
                item["smart"]["peak"],
            ]
            if actual != expected:
                raise AssertionError(f"작성서식 값 불일치: row={row}, {actual} != {expected}")
            if any(worksheet.cell(row, column).number_format != "#,##0" for column in range(3, 9)):
                raise AssertionError(f"작성서식 표시 형식 불일치: row={row}")
    finally:
        workbook.close()


def _find_target_rows(
    worksheet, rows_by_key: dict[tuple[int, str], dict]
) -> list[tuple[int, int, str]]:
    month: int | None = None
    target_rows: list[tuple[int, int, str]] = []
    for row in range(1, worksheet.max_row + 1):
        month_value = str(worksheet.cell(row, 1).value or "")
        if month_value in {"5월", "6월"}:
            month = int(month_value[0])
        label = str(worksheet.cell(row, 2).value or "").rstrip("*").strip()
        if month in MONTHS and (month, label) in rows_by_key:
            target_rows.append((month, row, label))
    if len(target_rows) != len(rows_by_key):
        raise RuntimeError(f"작성서식 대상 행 수가 올바르지 않습니다: {len(target_rows)}")
    return target_rows


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--edge-db", type=Path, default=EDGE_DB_PATH)
    parser.add_argument("--template", type=Path, default=TEMPLATE_PATH)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT_PATH)
    parser.add_argument(
        "--discover", action="store_true", help="List candidate Oracle intersection names."
    )
    parser.add_argument("--discover-approaches", action="store_true")
    parser.add_argument("--inspect-template", action="store_true")
    parser.add_argument(
        "--fill-v2",
        action="store_true",
        help="Fill smart-intersection values in the v2 remicon comparison workbook.",
    )
    parser.add_argument("--v2-workbook", type=Path, default=V2_WORKBOOK_PATH)
    parser.add_argument("--verify-v2", action="store_true")
    args = parser.parse_args()
    if args.discover:
        reader = OracleSmartReader()
        conn = reader._oracle.connect_db()
        try:
            with conn.cursor() as cursor:
                cursor.execute("SET TRANSACTION READ ONLY")
                cursor.execute("SELECT NODE_ID, CRSRD_NM FROM M_CRSRD_INF ORDER BY CRSRD_NM")
                for node_id, name in cursor.fetchall():
                    normalized = str(name).replace(" ", "")
                    if any(
                        key.replace(" ", "") in normalized
                        for key in ("박촌", "봉오", "삼정", "산업")
                    ):
                        print(f"{node_id}\t{name}")
        finally:
            conn.close()
        return 0
    if args.discover_approaches:
        reader = OracleSmartReader()
        conn = reader._oracle.connect_db()
        try:
            with conn.cursor() as cursor:
                cursor.execute("SET TRANSACTION READ ONLY")
                for target in TARGETS:
                    cursor.execute(
                        "SELECT NODE_ID FROM M_CRSRD_INF WHERE CRSRD_NM = :name",
                        name=target.smart_intersection,
                    )
                    node = cursor.fetchone()
                    if not node:
                        print(f"{target.label}\tmissing node")
                        continue
                    cursor.execute(
                        "SELECT ACSR_ID, ACSR_NM FROM M_CRSRD_ACSR_INF WHERE NODE_ID = :node_id ORDER BY ACSR_ID",
                        node_id=node[0],
                    )
                    for approach_id, approach_name in cursor.fetchall():
                        print(f"{target.label}\t{approach_id}\t{approach_name}")
        finally:
            conn.close()
        return 0
    if args.inspect_template:
        workbook = load_workbook(args.template, data_only=True)
        try:
            worksheet = workbook["작성서식"]
            for row in range(1, 21):
                print(f"{row}\t{worksheet.cell(row, 1).value}\t{worksheet.cell(row, 2).value}")
        finally:
            workbook.close()
        return 0
    if args.fill_v2:
        payload = smart_v2_payload(OracleIntersectionHourlyReader())
        non_target = v2_non_target_snapshot(args.v2_workbook)
        write_v2_workbook(args.v2_workbook, payload)
        verify_v2_workbook(args.v2_workbook, payload, non_target)
        print(
            json.dumps(
                {f"{month}월 {label}": values for (month, label), values in payload.items()},
                ensure_ascii=False,
                indent=2,
            )
        )
        return 0
    if args.verify_v2:
        payload = smart_v2_payload(OracleIntersectionHourlyReader())
        non_target = v2_non_target_snapshot(args.v2_workbook)
        verify_v2_workbook(args.v2_workbook, payload, non_target)
        return 0
    payload = result_payload(SQLiteEdgeReader(args.edge_db), OracleSmartReader())
    write_template(args.template, args.output, payload)
    verify_template(args.output, payload)
    print(f"output={args.output.resolve()}")
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
