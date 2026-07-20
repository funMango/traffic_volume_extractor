#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Export Korean-area ``14`` vehicle detections into location-based sheets."""

from __future__ import annotations

import argparse
import re
import sqlite3
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


BASE_DIR = Path(__file__).resolve().parents[2]
DEFAULT_DB_PATH = BASE_DIR / "00_Data" / "삼정동_레미콘" / "삼정동_레미콘_데이터.db"
DEFAULT_OUTPUT_PATH = (
    BASE_DIR / "02_Result" / "09_기타" / "삼정동_레미콘_지역명14차량_교차로별.xlsx"
)

TABLE_NAME = "vehicle_detection"
EXPECTED_TOTAL_ROWS = 47_935
EXPECTED_LOCATION_COUNT = 15
VEHICLE_NUMBER_PATTERN = re.compile(r"^[가-힣]{2}14[가-힣][0-9]\*{3}$")
INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
SOURCE_COLUMNS = [
    "id",
    "collected_at",
    "device_id",
    "location_name",
    "lane",
    "vehicle_number",
    "owner_registered_area",
    "source_file",
    "source_sheet",
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Export Korean-area 14 vehicle-number detections by location."
    )
    parser.add_argument("--db", type=Path, default=DEFAULT_DB_PATH)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT_PATH)
    parser.add_argument(
        "--skip-expected-check",
        action="store_true",
        help="Allow a changed database row or location count.",
    )
    return parser.parse_args()


def connect_readonly(db_path: Path) -> sqlite3.Connection:
    resolved = db_path.resolve()
    if not resolved.is_file():
        raise FileNotFoundError(f"Database file does not exist: {resolved}")
    return sqlite3.connect(f"file:{resolved.as_posix()}?mode=ro", uri=True)


def validate_schema(conn: sqlite3.Connection) -> None:
    table_exists = conn.execute(
        "SELECT 1 FROM sqlite_master WHERE type = 'table' AND name = ?", (TABLE_NAME,)
    ).fetchone()
    if table_exists is None:
        raise RuntimeError(f"Missing required table: {TABLE_NAME}")

    actual_columns = {str(row[1]) for row in conn.execute(f"PRAGMA table_info({TABLE_NAME})")}
    missing_columns = [column for column in SOURCE_COLUMNS if column not in actual_columns]
    if missing_columns:
        raise RuntimeError("Missing required columns: " + ", ".join(missing_columns))


def load_rows(conn: sqlite3.Connection) -> list[dict[str, Any]]:
    conn.row_factory = sqlite3.Row
    sql = f"""
        SELECT {", ".join(SOURCE_COLUMNS)}
        FROM {TABLE_NAME}
        WHERE vehicle_number LIKE '%14%'
        ORDER BY location_name ASC, collected_at ASC, id ASC
    """
    return [
        {column: row[column] for column in SOURCE_COLUMNS}
        for row in conn.execute(sql)
        if VEHICLE_NUMBER_PATTERN.fullmatch(str(row["vehicle_number"] or ""))
    ]


def group_rows_by_location(rows: list[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        location_name = str(row["location_name"] or "location_name_없음")
        grouped[location_name].append(row)
    return dict(sorted(grouped.items()))


def make_sheet_name(location_name: str, used_names: set[str]) -> str:
    base_name = INVALID_SHEET_CHARS.sub("_", location_name).strip() or "Sheet"
    base_name = base_name[:31]
    sheet_name = base_name
    suffix = 1
    while sheet_name in used_names:
        suffix_text = f"_{suffix}"
        sheet_name = base_name[: 31 - len(suffix_text)] + suffix_text
        suffix += 1
    used_names.add(sheet_name)
    return sheet_name


def display_width(value: object) -> int:
    return sum(2 if ord(char) > 127 else 1 for char in str(value or ""))


def format_worksheet(ws: Worksheet) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column_index, column_name in enumerate(SOURCE_COLUMNS, start=1):
        values = (cell.value for cell in ws[get_column_letter(column_index)])
        max_width = max(display_width(value) for value in values)
        ws.column_dimensions[get_column_letter(column_index)].width = min(
            max(max_width + 2, display_width(column_name) + 2, 12), 45
        )


def write_workbook(grouped_rows: dict[str, list[dict[str, Any]]], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    used_names: set[str] = set()

    for location_name, rows in grouped_rows.items():
        worksheet = workbook.create_sheet(make_sheet_name(location_name, used_names))
        worksheet.append(SOURCE_COLUMNS)
        for row in rows:
            worksheet.append([row[column] for column in SOURCE_COLUMNS])
        format_worksheet(worksheet)

    workbook.save(output_path)
    workbook.close()


def validate_source_rows(
    rows: list[dict[str, Any]], grouped_rows: dict[str, list[dict[str, Any]]]
) -> None:
    for row in rows:
        vehicle_number = str(row["vehicle_number"] or "")
        if VEHICLE_NUMBER_PATTERN.fullmatch(vehicle_number) is None:
            raise RuntimeError(f"Unexpected vehicle number: {vehicle_number}")

    for location_name, location_rows in grouped_rows.items():
        sort_keys = [(str(row["collected_at"] or ""), int(row["id"])) for row in location_rows]
        if sort_keys != sorted(sort_keys):
            raise RuntimeError(f"Rows are not sorted: {location_name}")


def validate_saved_workbook(
    output_path: Path, grouped_rows: dict[str, list[dict[str, Any]]]
) -> None:
    workbook = load_workbook(output_path, read_only=False, data_only=True)
    try:
        if len(workbook.sheetnames) != len(grouped_rows):
            raise RuntimeError("Saved workbook sheet count does not match locations")

        actual_total = 0
        expected_counts = sorted(len(rows) for rows in grouped_rows.values())
        actual_counts: list[int] = []
        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            if [cell.value for cell in worksheet[1]] != SOURCE_COLUMNS:
                raise RuntimeError(f"Unexpected header in sheet: {sheet_name}")
            if worksheet.freeze_panes != "A2" or worksheet.auto_filter.ref != worksheet.dimensions:
                raise RuntimeError(f"Missing filter or freeze pane in sheet: {sheet_name}")

            rows = list(worksheet.iter_rows(min_row=2, values_only=True))
            actual_counts.append(len(rows))
            actual_total += len(rows)
            previous_key: tuple[str, int] | None = None
            for row in rows:
                if VEHICLE_NUMBER_PATTERN.fullmatch(str(row[5] or "")) is None:
                    raise RuntimeError(f"Unexpected vehicle number in sheet: {sheet_name}")
                sort_key = (str(row[1] or ""), int(row[0]))
                if previous_key is not None and sort_key < previous_key:
                    raise RuntimeError(f"Unsorted saved rows in sheet: {sheet_name}")
                previous_key = sort_key

        if actual_total != sum(expected_counts) or sorted(actual_counts) != expected_counts:
            raise RuntimeError("Saved workbook row counts do not match source rows")
    finally:
        workbook.close()


def main() -> None:
    args = parse_args()
    conn = connect_readonly(args.db)
    try:
        validate_schema(conn)
        rows = load_rows(conn)
    finally:
        conn.close()

    grouped_rows = group_rows_by_location(rows)
    validate_source_rows(rows, grouped_rows)
    if not args.skip_expected_check and (
        len(rows) != EXPECTED_TOTAL_ROWS or len(grouped_rows) != EXPECTED_LOCATION_COUNT
    ):
        raise RuntimeError(
            f"Expected {EXPECTED_TOTAL_ROWS:,} rows across {EXPECTED_LOCATION_COUNT} locations; "
            f"got {len(rows):,} rows across {len(grouped_rows)} locations."
        )

    write_workbook(grouped_rows, args.output)
    validate_saved_workbook(args.output, grouped_rows)

    print(f"output={args.output.resolve()}")
    print(f"total_rows={len(rows)}")
    print(f"location_sheet_count={len(grouped_rows)}")
    for location_name, location_rows in grouped_rows.items():
        print(f"{location_name}\t{len(location_rows)}")


if __name__ == "__main__":
    main()
