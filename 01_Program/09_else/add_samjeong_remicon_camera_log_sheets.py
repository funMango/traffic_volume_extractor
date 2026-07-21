#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Add grouped Samjeong-dong remicon source-log sheets to the v2 workbook."""

from __future__ import annotations

import argparse
from collections import defaultdict
from datetime import datetime
from pathlib import Path
import re
import sqlite3
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


BASE_DIR = Path(__file__).resolve().parents[2]
DEFAULT_DB_PATH = BASE_DIR / "00_Data" / "삼정동_레미콘" / "삼정동_레미콘_데이터.db"
DEFAULT_WORKBOOK_PATH = BASE_DIR / "02_Result" / "09_기타" / "삼정동_레미콘_교통량_v2.xlsx"
TABLE_NAME = "vehicle_detection"
START_AT = "2026-05-01 00:00:00"
END_AT = "2026-07-01 00:00:00"

SOURCE_COLUMNS = [
    "id",
    "collected_at",
    "registered_at",
    "device_id",
    "location_name",
    "lane",
    "lane_direction_name",
    "vehicle_number",
    "owner_registered_area",
    "source_file",
    "source_sheet",
]
LOCATION_GROUPS = {
    "박촌교 삼거리": [
        "박촌교 삼거리[남향]",
        "박촌교 삼거리[북향]",
        "박촌교 삼거리[서향(순방향)]",
        "박촌교 삼거리[서향(역방향)]",
    ],
    "봉오고가교사거리": ["봉오고가교 사거리[남향]", "봉오고가교 사거리[서향]"],
    "삼정고가삼거리": ["삼정고가 삼거리[동향]", "삼정고가 삼거리[서향]"],
    "삼정교사거리": ["삼정교 사거리[북동향]"],
    "산업길사거리": ["산업길 사거리[남향]", "산업길 사거리[북향]"],
    "봉오대로사거리": ["봉오대로 사거리[동향]", "봉오대로 사거리[서향]"],
    # Excel forbids the ASCII asterisk in worksheet names; use its full-width equivalent.
    "자동차검사소＊": ["자동차검사소"],
    "삼정동 320-1": ["삼정동320-1 부천IC"],
}
REMICON_PATTERNS = (
    re.compile(r"^014[가-힣][0-9]\*{3}$"),
    re.compile(r"^[가-힣]{2,}14[가-힣][0-9]\*{3}$"),
)
DATE_COLUMNS = {"collected_at", "registered_at"}
COLUMN_WIDTHS = (12, 21, 21, 18, 34, 10, 18, 18, 20, 42, 24)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--db", type=Path, default=DEFAULT_DB_PATH)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK_PATH)
    parser.add_argument("--verify-only", action="store_true")
    return parser.parse_args()


def is_remicon(vehicle_number: object) -> bool:
    return any(pattern.fullmatch(str(vehicle_number or "")) for pattern in REMICON_PATTERNS)


def connect_readonly(db_path: Path) -> sqlite3.Connection:
    resolved = db_path.resolve()
    if not resolved.is_file():
        raise FileNotFoundError(f"Database file does not exist: {resolved}")
    connection = sqlite3.connect(f"file:{resolved.as_posix()}?mode=ro", uri=True)
    connection.row_factory = sqlite3.Row
    return connection


def validate_database(connection: sqlite3.Connection) -> None:
    columns = {row[1] for row in connection.execute(f"PRAGMA table_info({TABLE_NAME})")}
    missing = [column for column in SOURCE_COLUMNS if column not in columns]
    if missing:
        raise RuntimeError(f"Missing required columns: {', '.join(missing)}")


def fetch_grouped_rows(connection: sqlite3.Connection) -> dict[str, list[dict[str, Any]]]:
    requested_locations = [
        location for locations in LOCATION_GROUPS.values() for location in locations
    ]
    placeholders = ", ".join("?" for _ in requested_locations)
    sql = f"""
        SELECT {", ".join(SOURCE_COLUMNS)}
        FROM {TABLE_NAME}
        WHERE collected_at >= ? AND collected_at < ?
          AND location_name IN ({placeholders})
          AND vehicle_number LIKE '%14%'
        ORDER BY collected_at ASC, id ASC
    """
    records_by_location: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in connection.execute(sql, (START_AT, END_AT, *requested_locations)):
        if is_remicon(row["vehicle_number"]):
            records_by_location[str(row["location_name"])].append(
                {column: row[column] for column in SOURCE_COLUMNS}
            )

    return {
        sheet_name: sorted(
            (record for location in locations for record in records_by_location[location]),
            key=lambda record: (str(record["collected_at"]), int(record["id"])),
        )
        for sheet_name, locations in LOCATION_GROUPS.items()
    }


def write_log_sheet(workbook: Any, sheet_name: str, rows: list[dict[str, Any]]) -> None:
    worksheet = workbook.create_sheet(sheet_name)
    worksheet.append(SOURCE_COLUMNS)
    for row in rows:
        worksheet.append(
            [
                datetime.fromisoformat(row[column])
                if column in DATE_COLUMNS and row[column]
                else row[column]
                for column in SOURCE_COLUMNS
            ]
        )

    for column_index, width in enumerate(COLUMN_WIDTHS, start=1):
        worksheet.column_dimensions[get_column_letter(column_index)].width = width
    for cell in worksheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
    for row_index in range(2, worksheet.max_row + 1):
        for column_index, column_name in enumerate(SOURCE_COLUMNS, start=1):
            if column_name in DATE_COLUMNS:
                cell = worksheet.cell(row=row_index, column=column_index)
                if isinstance(cell.value, datetime):
                    cell.number_format = "yyyy-mm-dd hh:mm:ss"
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions


def update_workbook(workbook_path: Path, grouped_rows: dict[str, list[dict[str, Any]]]) -> None:
    if not workbook_path.is_file():
        raise FileNotFoundError(f"Workbook file does not exist: {workbook_path.resolve()}")
    workbook = load_workbook(workbook_path)
    try:
        required_sheets = {"작성서식", "월별 일평균"}
        missing = required_sheets - set(workbook.sheetnames)
        if missing:
            raise RuntimeError(f"Missing required sheets: {', '.join(sorted(missing))}")
        for sheet_name in LOCATION_GROUPS:
            if sheet_name in workbook.sheetnames:
                del workbook[sheet_name]
        for sheet_name, rows in grouped_rows.items():
            write_log_sheet(workbook, sheet_name, rows)
        workbook.save(workbook_path)
    finally:
        workbook.close()


def validate_saved_workbook(
    workbook_path: Path, grouped_rows: dict[str, list[dict[str, Any]]]
) -> None:
    workbook = load_workbook(workbook_path, data_only=True)
    try:
        expected_sheets = ["작성서식", "월별 일평균", *LOCATION_GROUPS]
        if workbook.sheetnames != expected_sheets:
            raise RuntimeError(f"Unexpected sheets: {workbook.sheetnames}")
        for sheet_name, expected_rows in grouped_rows.items():
            worksheet = workbook[sheet_name]
            header = [cell.value for cell in worksheet[1]]
            if header != SOURCE_COLUMNS:
                raise RuntimeError(f"Header mismatch: {sheet_name}")
            if worksheet.max_row - 1 != len(expected_rows):
                raise RuntimeError(f"Row-count mismatch: {sheet_name}")
            if worksheet.freeze_panes != "A2" or worksheet.auto_filter.ref != worksheet.dimensions:
                raise RuntimeError(f"Worksheet layout mismatch: {sheet_name}")
            for row_index in range(2, worksheet.max_row + 1):
                for column_index, column_name in enumerate(SOURCE_COLUMNS, start=1):
                    if column_name in DATE_COLUMNS:
                        cell = worksheet.cell(row=row_index, column=column_index)
                        if (
                            not isinstance(cell.value, datetime)
                            or cell.number_format != "yyyy-mm-dd hh:mm:ss"
                        ):
                            raise RuntimeError(
                                f"Date format mismatch: {sheet_name} row={row_index}"
                            )
            actual_rows = list(worksheet.iter_rows(min_row=2, values_only=True))
            for actual, expected in zip(actual_rows, expected_rows, strict=True):
                normalized_actual = tuple(
                    value.strftime("%Y-%m-%d %H:%M:%S")
                    if column in DATE_COLUMNS and isinstance(value, datetime)
                    else ""
                    if value is None
                    else value
                    for column, value in zip(SOURCE_COLUMNS, actual, strict=True)
                )
                normalized_expected = tuple(
                    "" if expected[column] is None else expected[column]
                    for column in SOURCE_COLUMNS
                )
                if normalized_actual != normalized_expected:
                    raise RuntimeError(f"Source-row mismatch: {sheet_name}")
                if not is_remicon(actual[7]):
                    raise RuntimeError(f"Vehicle-number mismatch: {sheet_name}")
                if not (START_AT <= str(actual[1]) < END_AT):
                    raise RuntimeError(f"Date-range mismatch: {sheet_name}")
            keys = [(str(row[1]), int(row[0])) for row in actual_rows]
            if keys != sorted(keys):
                raise RuntimeError(f"Sort mismatch: {sheet_name}")
    finally:
        workbook.close()


def main() -> int:
    args = parse_args()
    connection = connect_readonly(args.db)
    try:
        validate_database(connection)
        grouped_rows = fetch_grouped_rows(connection)
    finally:
        connection.close()
    if not args.verify_only:
        update_workbook(args.workbook, grouped_rows)
    validate_saved_workbook(args.workbook, grouped_rows)
    print(f"workbook={args.workbook.resolve()}")
    print(f"sheet_count={len(grouped_rows)}")
    print(f"total_rows={sum(len(rows) for rows in grouped_rows.values())}")
    for sheet_name, rows in grouped_rows.items():
        print(f"{sheet_name}\t{len(rows)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
