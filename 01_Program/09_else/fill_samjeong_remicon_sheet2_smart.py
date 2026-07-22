#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Fill Sheet2 smart-intersection daily averages from Oracle hourly approach traffic."""

from __future__ import annotations

import argparse
import sys
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_WORKBOOK_PATH = PROJECT_ROOT / "02_Result" / "09_기타" / "삼정동_레미콘_교통량_수정_v1.xlsx"
SHEET_NAME = "스마트교차로 비교"
TRAFFIC_TABLE = "S_CRSRD_ACSR_TRF_1HH"
HOLIDAYS = {date(2026, 5, 1), date(2026, 5, 5), date(2026, 5, 25), date(2026, 6, 3)}
PEAK_HOURS = (7, 8, 17, 18)
TARGET_COLUMNS = (7, 9, 11)


@dataclass(frozen=True)
class Period:
    month: str
    start: date
    end: date
    daily_divisor: int
    weekday_divisor: int


PERIODS = (
    Period("5월", date(2026, 5, 1), date(2026, 6, 1), 31, 18),
    Period("6월", date(2026, 6, 1), date(2026, 7, 1), 30, 21),
)
PERIOD_BY_MONTH = {period.month: period for period in PERIODS}

# Worksheet intersection/direction -> Oracle intersection/approach names.
APPROACH_NAMES = {
    ("박촌교 삼거리", "남향"): ("박촌교삼거리", "박촌교삼거리-북 (남향)"),
    ("박촌교 삼거리", "북향"): ("박촌교삼거리", "박촌교삼거리-남 (북향)"),
    ("박촌교 삼거리", "서향"): ("박촌교삼거리", "박촌교삼거리-서 (동향)"),
    ("봉오고가교 사거리", "남향"): ("봉오고가교사거리", "봉오고가교사거리-북 (남향)"),
    ("봉오고가교 사거리", "서향"): ("봉오고가교사거리", "봉오고가교사거리-동 (서향)"),
    ("삼정고가 삼거리", "동향"): ("삼정고가교삼거리", "삼정고가교삼거리-서 (동향)"),
    ("삼정고가 삼거리", "서향"): ("삼정고가교삼거리", "삼정고가교삼거리-동 (서향)"),
    ("삼정교 사거리", "동향"): ("삼정교사거리", "삼정교사거리-서 (동향)"),
    ("산업길 사거리", "남향"): ("산업길사거리", "산업길사거리-북 (남향)"),
    ("산업길 사거리", "북향"): ("산업길사거리", "산업길사거리-남 (북향)"),
    ("봉오대로 사거리", "동향"): ("봉오대로사거리", "봉오대로사거리-서 (동향)"),
    ("봉오대로 사거리", "서향"): ("봉오대로사거리", "봉오대로4R-동 (서향)"),
}


@dataclass(frozen=True)
class DirectionGroup:
    month: str
    intersection: str
    direction_rows: tuple[int, ...]
    total_row: int | None


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK_PATH)
    parser.add_argument("--verify-only", action="store_true")
    return parser.parse_args()


def round_half_up(value: Decimal) -> int:
    return int(value.quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def weekday_dates(period: Period) -> tuple[date, ...]:
    result: list[date] = []
    current = period.start
    while current < period.end:
        if current.weekday() < 5 and current not in HOLIDAYS:
            result.append(current)
        current += timedelta(days=1)
    if len(result) != period.weekday_divisor:
        raise AssertionError(f"Unexpected weekday divisor for {period.month}: {len(result)}")
    return tuple(result)


def calculate_metrics(
    hourly_values: dict[date, dict[int, int]], period: Period
) -> tuple[int, int, int]:
    monthly_total = sum(sum(hours.values()) for hours in hourly_values.values())
    weekdays = weekday_dates(period)
    weekday_total = sum(sum(hourly_values.get(day, {}).values()) for day in weekdays)
    peak_total = sum(
        hourly_values.get(day, {}).get(hour, 0) for day in weekdays for hour in PEAK_HOURS
    )
    return (
        round_half_up(Decimal(monthly_total) / period.daily_divisor),
        round_half_up(Decimal(weekday_total) / period.weekday_divisor),
        round_half_up(Decimal(peak_total) / period.weekday_divisor),
    )


def find_direction_groups(worksheet: Any) -> tuple[DirectionGroup, ...]:
    current_month = ""
    current_intersection = ""
    grouped: dict[tuple[str, str], list[int]] = {}
    total_rows: dict[tuple[str, str], int] = {}
    ordered_keys: list[tuple[str, str]] = []
    for row in range(4, worksheet.max_row + 1):
        if value := worksheet.cell(row, 1).value:
            current_month = str(value).strip()
        if value := worksheet.cell(row, 2).value:
            current_intersection = str(value).strip()
        direction = str(worksheet.cell(row, 3).value or "").strip()
        if not direction:
            continue
        key = (current_month, current_intersection)
        if direction == "합계":
            total_rows[key] = row
            continue
        if current_month not in PERIOD_BY_MONTH or not current_intersection:
            raise RuntimeError(f"Unexpected Sheet2 label at row {row}")
        if key not in grouped:
            grouped[key] = []
            ordered_keys.append(key)
        grouped[key].append(row)
    groups = tuple(
        DirectionGroup(
            month,
            intersection,
            tuple(grouped[(month, intersection)]),
            total_rows.get((month, intersection)),
        )
        for month, intersection in ordered_keys
    )
    expected = set(APPROACH_NAMES)
    found = {
        (group.intersection, str(worksheet.cell(row, 3).value).strip())
        for group in groups
        for row in group.direction_rows
    }
    if found != expected:
        raise RuntimeError(
            f"Unexpected Sheet2 direction rows: missing={expected - found}, extra={found - expected}"
        )
    if len(groups) != 12 or sum(len(group.direction_rows) for group in groups) != 24:
        raise RuntimeError("Sheet2 must contain 24 mapped direction rows across two months")
    if sum(group.total_row is not None for group in groups) != 10:
        raise RuntimeError("Sheet2 must contain 10 total rows")
    return groups


class OracleApproachTrafficReader:
    """Infrastructure adapter for a read-only Oracle aggregation query."""

    def __init__(self) -> None:
        sys.path.insert(0, str(PROJECT_ROOT / "01_Program" / "09_else"))
        from traffic_api.infrastructure import oracle_analysis

        self._oracle = oracle_analysis

    def load(self) -> dict[tuple[str, str], dict[date, dict[int, int]]]:
        connection = self._oracle.connect_db()
        try:
            with connection.cursor() as cursor:
                cursor.execute("SET TRANSACTION READ ONLY")
                resolved = self._resolve_approaches(cursor)
                return self._load_hourly_values(cursor, resolved)
        finally:
            connection.close()

    def _resolve_approaches(self, cursor: Any) -> dict[tuple[str, str], tuple[str, str]]:
        intersections = sorted({names[0] for names in APPROACH_NAMES.values()})
        binds = {f"name_{index}": name for index, name in enumerate(intersections)}
        placeholders = ", ".join(f":{key}" for key in binds)
        cursor.execute(
            f"SELECT NODE_ID, CRSRD_NM FROM M_CRSRD_INF WHERE CRSRD_NM IN ({placeholders})", binds
        )
        by_name: dict[str, list[str]] = defaultdict(list)
        for node_id, name in cursor.fetchall():
            by_name[str(name).strip()].append(str(node_id).strip())
        node_by_name = {}
        for name in intersections:
            matches = by_name[name]
            if len(matches) != 1:
                raise RuntimeError(f"Expected one NODE_ID for {name}, found {len(matches)}")
            node_by_name[name] = matches[0]

        result: dict[tuple[str, str], tuple[str, str]] = {}
        for key, (intersection, approach_name) in APPROACH_NAMES.items():
            node_id = node_by_name[intersection]
            cursor.execute(
                "SELECT ACSR_ID FROM M_CRSRD_ACSR_INF "
                "WHERE NODE_ID = :node_id AND ACSR_NM = :approach_name",
                {"node_id": node_id, "approach_name": approach_name},
            )
            matches = [str(row[0]).strip() for row in cursor.fetchall()]
            if len(matches) != 1:
                raise RuntimeError(
                    f"Expected one ACSR_ID for {intersection} / {approach_name}, found {len(matches)}"
                )
            result[key] = (node_id, matches[0])
        return result

    def _load_hourly_values(
        self, cursor: Any, resolved: dict[tuple[str, str], tuple[str, str]]
    ) -> dict[tuple[str, str], dict[date, dict[int, int]]]:
        pairs = sorted(set(resolved.values()))
        pair_clauses = []
        binds: dict[str, Any] = {
            "start_time": datetime(2026, 5, 1),
            "end_time": datetime(2026, 7, 1),
        }
        for index, (node_id, approach_id) in enumerate(pairs):
            node_key, approach_key = f"node_{index}", f"approach_{index}"
            binds[node_key] = node_id
            binds[approach_key] = approach_id
            pair_clauses.append(f"(NODE_ID = :{node_key} AND ACSR_ID = :{approach_key})")
        cursor.execute(
            f"""
            SELECT NODE_ID, ACSR_ID, TRUNC(TOT_DT), TO_NUMBER(TO_CHAR(TOT_DT, 'HH24')),
                   SUM(NVL(TRF_QNTY, 0))
            FROM {TRAFFIC_TABLE}
            WHERE ({" OR ".join(pair_clauses)})
              AND TOT_DT >= :start_time AND TOT_DT < :end_time
            GROUP BY NODE_ID, ACSR_ID, TRUNC(TOT_DT), TO_NUMBER(TO_CHAR(TOT_DT, 'HH24'))
            """,
            binds,
        )
        key_by_pair = {pair: key for key, pair in resolved.items()}
        result = {key: defaultdict(dict) for key in resolved}
        for node_id, approach_id, raw_day, hour, traffic in cursor.fetchall():
            pair = (str(node_id).strip(), str(approach_id).strip())
            day = raw_day.date() if isinstance(raw_day, datetime) else raw_day
            result[key_by_pair[pair]][day][int(hour)] = int(traffic or 0)
        return {key: dict(days) for key, days in result.items()}


def expected_values(
    worksheet: Any, reader: OracleApproachTrafficReader
) -> tuple[dict[int, tuple[int, int, int]], tuple[DirectionGroup, ...]]:
    groups = find_direction_groups(worksheet)
    hourly_values = reader.load()
    values: dict[int, tuple[int, int, int]] = {}
    for group in groups:
        period = PERIOD_BY_MONTH[group.month]
        for row in group.direction_rows:
            direction = str(worksheet.cell(row, 3).value).strip()
            values[row] = calculate_metrics(hourly_values[(group.intersection, direction)], period)
    return values, groups


def workbook_snapshot(workbook: Any) -> dict[tuple[str, str], Any]:
    return {
        (worksheet.title, cell.coordinate): cell.value
        for worksheet in workbook.worksheets
        for row in worksheet.iter_rows()
        for cell in row
        if cell.value is not None
    }


def target_coordinates(groups: Iterable[DirectionGroup]) -> set[tuple[str, str]]:
    coordinates = {
        (SHEET_NAME, f"{chr(64 + column)}{row}")
        for group in groups
        for row in group.direction_rows
        for column in TARGET_COLUMNS
    }
    coordinates.update(
        (SHEET_NAME, f"{chr(64 + column)}{group.total_row}")
        for group in groups
        if group.total_row is not None
        for column in TARGET_COLUMNS
    )
    return coordinates


def direction_total(worksheet: Any, direction_rows: tuple[int, ...], column: int) -> int:
    return sum(int(worksheet.cell(row, column).value or 0) for row in direction_rows)


def apply_or_verify(
    workbook_path: Path, values: dict[int, tuple[int, int, int]], verify_only: bool
) -> None:
    workbook = load_workbook(workbook_path, data_only=False)
    try:
        worksheet = workbook[SHEET_NAME]
        groups = find_direction_groups(worksheet)
        targets = target_coordinates(groups)
        before = workbook_snapshot(workbook)
        if not verify_only:
            for row, metrics in values.items():
                for column, metric in zip(TARGET_COLUMNS, metrics, strict=True):
                    cell = worksheet.cell(row, column)
                    cell.value = metric
                    cell.number_format = "#,##0"
            for group in groups:
                if group.total_row is None:
                    continue
                for column in TARGET_COLUMNS:
                    worksheet.cell(group.total_row, column).value = direction_total(
                        worksheet, group.direction_rows, column
                    )
                    worksheet.cell(group.total_row, column).number_format = "#,##0"
            workbook.save(workbook_path)
    finally:
        workbook.close()

    saved = load_workbook(workbook_path, data_only=False)
    try:
        worksheet = saved[SHEET_NAME]
        for row, metrics in values.items():
            actual = tuple(worksheet.cell(row, column).value for column in TARGET_COLUMNS)
            if actual != metrics:
                raise AssertionError(f"Metric mismatch at Sheet2 row {row}: {actual!r}")
        groups = find_direction_groups(worksheet)
        for group in groups:
            if group.total_row is None:
                continue
            for column in TARGET_COLUMNS:
                expected = direction_total(worksheet, group.direction_rows, column)
                if worksheet.cell(group.total_row, column).value != expected:
                    raise AssertionError(
                        f"Total value mismatch at {worksheet.cell(group.total_row, column).coordinate}"
                    )
        if not verify_only:
            after = workbook_snapshot(saved)
            unexpected = sorted(
                key
                for key in before | after
                if before.get(key) != after.get(key) and key not in targets
            )
            if unexpected:
                raise AssertionError(f"Unexpected workbook changes: {unexpected[:10]}")
    finally:
        saved.close()


def main() -> int:
    args = parse_args()
    if not args.workbook.is_file():
        raise FileNotFoundError(args.workbook)
    workbook = load_workbook(args.workbook, read_only=True, data_only=False)
    try:
        values, _groups = expected_values(workbook[SHEET_NAME], OracleApproachTrafficReader())
    finally:
        workbook.close()
    apply_or_verify(args.workbook, values, args.verify_only)
    for row, metrics in sorted(values.items()):
        print(row, *metrics, sep="\t")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
