#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Update the Samjeong-dong remicon workbook from the read-only SQLite source."""

from __future__ import annotations

import argparse
import re
import sqlite3
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, timedelta
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parents[2]
DEFAULT_DB_PATH = BASE_DIR / "00_Data" / "삼정동_레미콘" / "삼정동_레미콘_데이터.db"
DEFAULT_WORKBOOK_PATH = (
    BASE_DIR / "02_Result" / "09_기타" / "삼정동_레미콘_교통량_지역명14포함.xlsx"
)

TABLE_NAME = "vehicle_detection"
TEMPLATE_SHEET = "작성서식"
AVERAGE_SHEET = "월별 일평균"
DETAIL_COLUMNS = [
    "id",
    "collected_at",
    "registered_at",
    "device_id",
    "location_name",
    "lane",
    "lane_direction_name",
    "vehicle_number",
    "owner_registered_area",
    "source_file",
    "source_sheet",
]
REMICON_PATTERNS = (
    re.compile(r"^014[가-힣]"),
    re.compile(r"^[가-힣]{2}14[가-힣][0-9]\*{3}$"),
)
INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
HOLIDAYS = {date(2026, 5, 1), date(2026, 5, 5), date(2026, 5, 25), date(2026, 6, 3)}
MONTHS = (("5월", date(2026, 5, 1), date(2026, 6, 1)), ("6월", date(2026, 6, 1), date(2026, 7, 1)))
LOCATION_GROUPS = {
    "박촌교 삼거리": [
        "박촌교 삼거리[남향]",
        "박촌교 삼거리[북향]",
        "박촌교 삼거리[서향(순방향)]",
        "박촌교 삼거리[서향(역방향)]",
    ],
    "봉오고가교사거리": ["봉오고가교 사거리[남향]", "봉오고가교 사거리[서향]"],
    "삼정고가삼거리": ["삼정고가 삼거리[동향]", "삼정고가 삼거리[서향]"],
    "삼정교사거리": ["삼정교 사거리[북동향]"],
    "산업길사거리": ["산업길 사거리[남향]", "산업길 사거리[북향]"],
    "봉오대로사거리": ["봉오대로 사거리[동향]", "봉오대로 사거리[서향]"],
    "자동차검사소": ["자동차검사소"],
    "삼정동 320-1": ["삼정동320-1 부천IC"],
}


@dataclass(frozen=True)
class Counts:
    total: int
    remicon: int


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--db", type=Path, default=DEFAULT_DB_PATH)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK_PATH)
    parser.add_argument("--verify-only", action="store_true")
    return parser.parse_args()


def is_remicon(value: object) -> int:
    return int(any(pattern.fullmatch(str(value or "")) for pattern in REMICON_PATTERNS))


def detail_sheet_name(location_name: str) -> str:
    return INVALID_SHEET_CHARS.sub("_", location_name).strip()[:31]


def connect_readonly(db_path: Path) -> sqlite3.Connection:
    resolved = db_path.resolve()
    if not resolved.is_file():
        raise FileNotFoundError(f"Database file does not exist: {resolved}")
    conn = sqlite3.connect(f"file:{resolved.as_posix()}?mode=ro", uri=True)
    conn.create_function("IS_REMICON", 1, is_remicon)
    return conn


def validate_database(conn: sqlite3.Connection) -> None:
    columns = {str(row[1]) for row in conn.execute(f"PRAGMA table_info({TABLE_NAME})")}
    required = set(DETAIL_COLUMNS)
    if missing := sorted(required - columns):
        raise RuntimeError("Missing required columns: " + ", ".join(missing))


def day_count(start: date, end: date, criterion: str, group_name: str) -> int:
    dates: list[date] = []
    current = start
    while current < end:
        eligible = criterion == "월별" or (current.weekday() < 5 and current not in HOLIDAYS)
        if group_name == "자동차검사소" and date(2026, 5, 20) <= current <= date(2026, 5, 25):
            eligible = False
        if eligible:
            dates.append(current)
        current += timedelta(days=1)
    return len(dates)


def fetch_counts(
    conn: sqlite3.Connection, locations: list[str], start: date, end: date, criterion: str
) -> Counts:
    placeholders = ", ".join("?" for _ in locations)
    conditions = ""
    params: list[object] = [
        f"{start.isoformat()} 00:00:00",
        f"{end.isoformat()} 00:00:00",
        *locations,
    ]
    if criterion != "월별":
        holiday_placeholders = ", ".join("?" for _ in HOLIDAYS)
        conditions += f" AND CAST(strftime('%w', collected_at) AS INTEGER) NOT IN (0, 6) AND date(collected_at) NOT IN ({holiday_placeholders})"
        params.extend(day.isoformat() for day in sorted(HOLIDAYS))
    if criterion == "평일 첨두시":
        conditions += " AND ((time(collected_at) >= '07:00:00' AND time(collected_at) < '09:00:00') OR (time(collected_at) >= '17:00:00' AND time(collected_at) < '19:00:00'))"
    total, remicon = conn.execute(
        f"SELECT COUNT(*), SUM(IS_REMICON(vehicle_number)) FROM {TABLE_NAME} WHERE collected_at >= ? AND collected_at < ? AND location_name IN ({placeholders}){conditions}",
        params,
    ).fetchone()
    return Counts(int(total or 0), int(remicon or 0))


def round_two(value: Decimal) -> float:
    return float(value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def write_average_sheet(workbook: Any, conn: sqlite3.Connection) -> None:
    if AVERAGE_SHEET in workbook.sheetnames:
        del workbook[AVERAGE_SHEET]
    source = workbook[TEMPLATE_SHEET]
    ws = workbook.copy_worksheet(source)
    ws.title = AVERAGE_SHEET
    workbook._sheets.remove(ws)
    workbook._sheets.insert(1, ws)

    month_rows = ((MONTHS[0], range(4, 12)), (MONTHS[1], range(12, 20)))
    criteria = (("월별", 3), ("평일", 6), ("평일 첨두시", 9))
    for (month_label, start, end), rows in month_rows:
        for row in rows:
            group_name = str(ws.cell(row=row, column=2).value or "").rstrip("*")
            locations = LOCATION_GROUPS[group_name]
            for criterion, column in criteria:
                divisor = day_count(start, end, criterion, group_name)
                counts = fetch_counts(conn, locations, start, end, criterion)
                total_average = round_two(Decimal(counts.total) / Decimal(divisor))
                remicon_average = round_two(Decimal(counts.remicon) / Decimal(divisor))
                ratio = (
                    round_two(Decimal(counts.remicon) * Decimal("100") / Decimal(counts.total))
                    if counts.total
                    else 0.0
                )
                ws.cell(row=row, column=column).value = total_average
                ws.cell(row=row, column=column + 1).value = remicon_average
                ws.cell(row=row, column=column + 2).value = ratio
                ws.cell(row=row, column=column).number_format = "#,##0.00"
                ws.cell(row=row, column=column + 1).number_format = "#,##0.00"
                ws.cell(row=row, column=column + 2).number_format = "0.00"


def load_detail_rows(conn: sqlite3.Connection) -> dict[str, list[dict[str, Any]]]:
    conn.row_factory = sqlite3.Row
    rows = conn.execute(
        f"SELECT {', '.join(DETAIL_COLUMNS)} FROM {TABLE_NAME} WHERE collected_at >= ? AND collected_at < ? AND vehicle_number LIKE '%14%' ORDER BY location_name, collected_at, id",
        ("2026-05-01 00:00:00", "2026-07-01 00:00:00"),
    )
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        if is_remicon(row["vehicle_number"]):
            grouped[str(row["location_name"] or "location_name_없음")].append(
                {column: row[column] for column in DETAIL_COLUMNS}
            )
    return dict(sorted(grouped.items()))


def update_log_sheets(workbook: Any, grouped_rows: dict[str, list[dict[str, Any]]]) -> None:
    expected_names = {detail_sheet_name(location_name) for location_name in grouped_rows}
    actual_names = {
        sheet_name
        for sheet_name in workbook.sheetnames
        if [cell.value for cell in workbook[sheet_name][1]] == DETAIL_COLUMNS
    }
    if actual_names != expected_names:
        raise RuntimeError(
            f"Detail sheet mismatch: expected {sorted(expected_names)}, got {sorted(actual_names)}; "
            f"workbook sheets={workbook.sheetnames}"
        )
    for location_name, rows in grouped_rows.items():
        ws = workbook[detail_sheet_name(location_name)]
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
        for row in rows:
            ws.append([row[column] for column in DETAIL_COLUMNS])
        ws.auto_filter.ref = ws.dimensions


def validate_workbook(workbook_path: Path, grouped_rows: dict[str, list[dict[str, Any]]]) -> None:
    wb = load_workbook(workbook_path, read_only=False, data_only=True)
    try:
        if AVERAGE_SHEET not in wb.sheetnames:
            raise RuntimeError("Monthly average sheet is missing")
        ws = wb[AVERAGE_SHEET]
        template_ws = wb[TEMPLATE_SHEET]
        if ws.max_row != 21 or ws.max_column != 11:
            raise RuntimeError("Monthly average sheet must use the A1:K21 template layout")
        if set(ws.merged_cells.ranges) != set(template_ws.merged_cells.ranges):
            raise RuntimeError("Monthly average sheet merged ranges differ from the template")
        for row in range(1, 22):
            for column in range(1, 12):
                if row >= 4 and row <= 19 and column >= 3:
                    continue
                if (
                    ws.cell(row=row, column=column).value
                    != template_ws.cell(row=row, column=column).value
                ):
                    raise RuntimeError("Monthly average sheet labels differ from the template")
        for location_name, expected_rows in grouped_rows.items():
            sheet_name = detail_sheet_name(location_name)
            ws = wb[sheet_name]
            if ws.auto_filter.ref != ws.dimensions or ws.max_row - 1 != len(expected_rows):
                raise RuntimeError(f"Detail sheet validation failed: {sheet_name}")
            previous: tuple[str, int] | None = None
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not is_remicon(row[7]):
                    raise RuntimeError(f"Unexpected vehicle number in {sheet_name}")
                key = (str(row[1]), int(row[0]))
                if previous is not None and key < previous:
                    raise RuntimeError(f"Unsorted detail sheet: {sheet_name}")
                previous = key
    finally:
        wb.close()


def validate_average_values(workbook_path: Path, conn: sqlite3.Connection) -> None:
    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        ws = wb[AVERAGE_SHEET]
        month_rows = ((MONTHS[0], range(4, 12)), (MONTHS[1], range(12, 20)))
        criteria = (("월별", 3), ("평일", 6), ("평일 첨두시", 9))
        for (_month_label, start, end), rows in month_rows:
            for row in rows:
                group_name = str(ws.cell(row=row, column=2).value or "").rstrip("*")
                for criterion, column in criteria:
                    divisor = day_count(start, end, criterion, group_name)
                    counts = fetch_counts(conn, LOCATION_GROUPS[group_name], start, end, criterion)
                    expected = (
                        round_two(Decimal(counts.total) / Decimal(divisor)),
                        round_two(Decimal(counts.remicon) / Decimal(divisor)),
                        round_two(Decimal(counts.remicon) * Decimal("100") / Decimal(counts.total))
                        if counts.total
                        else 0.0,
                    )
                    actual = tuple(
                        ws.cell(row=row, column=index).value for index in range(column, column + 3)
                    )
                    if actual != expected:
                        raise RuntimeError(
                            f"Monthly average DB reaggregation mismatch: row={row}, criterion={criterion}, {actual} != {expected}"
                        )
                    if (
                        ws.cell(row=row, column=column).number_format != "#,##0.00"
                        or ws.cell(row=row, column=column + 1).number_format != "#,##0.00"
                        or ws.cell(row=row, column=column + 2).number_format != "0.00"
                    ):
                        raise RuntimeError(
                            f"Monthly average number format mismatch: row={row}, criterion={criterion}"
                        )
    finally:
        wb.close()


def main() -> int:
    args = parse_args()
    if not args.workbook.is_file():
        raise FileNotFoundError(f"Workbook file does not exist: {args.workbook.resolve()}")
    conn = connect_readonly(args.db)
    try:
        validate_database(conn)
        grouped_rows = load_detail_rows(conn)
        if len(grouped_rows) != 15:
            raise RuntimeError(f"Expected 15 detail locations, got {len(grouped_rows)}")
        if not args.verify_only:
            workbook = load_workbook(args.workbook)
            try:
                if TEMPLATE_SHEET not in workbook.sheetnames:
                    raise RuntimeError(f"Missing required sheet: {TEMPLATE_SHEET}")
                write_average_sheet(workbook, conn)
                update_log_sheets(workbook, grouped_rows)
                workbook.save(args.workbook)
            finally:
                workbook.close()
        validate_workbook(args.workbook, grouped_rows)
        validate_average_values(args.workbook, conn)
    finally:
        conn.close()
    print(f"workbook={args.workbook.resolve()}")
    print(f"detail_rows={sum(len(rows) for rows in grouped_rows.values())}")
    print(f"detail_sheets={len(grouped_rows)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
