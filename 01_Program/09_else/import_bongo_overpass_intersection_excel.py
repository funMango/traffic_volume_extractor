#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Import May--June 2026 Bongo Overpass intersection Excel data into SQLite."""

from __future__ import annotations

import argparse
import sqlite3
import warnings
from collections.abc import Iterable
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parents[2]
DB_PATH = BASE_DIR / "00_Data" / "삼정동_레미콘_데이터.db"
INPUT_DIR = BASE_DIR / "00_Data" / "2605_06_삼정동_레미콘_데이터" / "봉오고가교사거리"
SHEET_NAME = "data_0"
TABLE_NAME = "vehicle_detection"
SUMMARY_TABLE_NAME = "import_summary"
SOURCE_FILES = (
    "260501_260510_봉오고가교사거리.xlsx",
    "260511_260522_봉오고가교사거리.xlsx",
    "260523_260531_봉오고가교사거리.xlsx",
    "260601_260608_봉오고가교사거리.xlsx",
    "260609_260616_봉오고가교사거리.xlsx",
    "260617_260624_봉오고가교사거리.xlsx",
    "260625_260630_봉오고가교사거리.xlsx",
)
EXPECTED_HEADER = (
    "수집일시",
    "등록일시",
    "장비ID",
    "위치명",
    "차선",
    "차선방향명",
    "차량번호",
    "차량소유주등록지",
)
INSERT_COLUMNS = (
    "collected_at",
    "registered_at",
    "device_id",
    "location_name",
    "lane",
    "lane_direction_name",
    "vehicle_number",
    "owner_registered_area",
    "source_file",
    "source_sheet",
)
EXPECTED_FILE_ROWS = {
    SOURCE_FILES[0]: 96843,
    SOURCE_FILES[1]: 97606,
    SOURCE_FILES[2]: 99483,
    SOURCE_FILES[3]: 88120,
    SOURCE_FILES[4]: 97691,
    SOURCE_FILES[5]: 99002,
    SOURCE_FILES[6]: 68606,
}
EXPECTED_LOCATION_ROWS = {"봉오고가교 사거리[남향]": 511847, "봉오고가교 사거리[서향]": 135504}
EXPECTED_TOTAL_ROWS = 3345300


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--db", type=Path, default=DB_PATH, help="Target SQLite DB path")
    parser.add_argument("--input-dir", type=Path, default=INPUT_DIR, help="Source Excel directory")
    parser.add_argument(
        "--verify-only", action="store_true", help="Validate imported rows without writing"
    )
    return parser.parse_args()


def cell_to_text(value: Any) -> str:
    return "" if value is None else str(value)


def validate_db_schema(conn: sqlite3.Connection) -> None:
    vehicle_columns = {row[1] for row in conn.execute(f"PRAGMA table_info({TABLE_NAME})")}
    if missing_columns := [column for column in INSERT_COLUMNS if column not in vehicle_columns]:
        raise RuntimeError("Missing vehicle_detection columns: " + ", ".join(missing_columns))
    summary_columns = {row[1] for row in conn.execute(f"PRAGMA table_info({SUMMARY_TABLE_NAME})")}
    required_columns = {"source_file", "source_sheet", "column_count", "imported_rows"}
    if missing_columns := sorted(required_columns - summary_columns):
        raise RuntimeError("Missing import_summary columns: " + ", ".join(missing_columns))


def source_paths(input_dir: Path) -> list[Path]:
    paths = [input_dir / name for name in SOURCE_FILES]
    if missing_paths := [str(path) for path in paths if not path.is_file()]:
        raise FileNotFoundError("Missing source files: " + ", ".join(missing_paths))
    return paths


def iter_source_rows(path: Path) -> Iterable[tuple[str, ...]]:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise RuntimeError(f"Missing sheet {SHEET_NAME!r}: {path.name}")
        worksheet = workbook[SHEET_NAME]
        worksheet.reset_dimensions()
        rows = worksheet.iter_rows(values_only=True)
        header = tuple(cell_to_text(value) for value in next(rows, ()))
        if header != EXPECTED_HEADER:
            raise RuntimeError(f"Header mismatch in {path.name}: {header!r}")
        for row in rows:
            if len(row) != len(EXPECTED_HEADER):
                raise RuntimeError(f"Column count mismatch in {path.name}: {len(row)}")
            collected_at, registered_at, device_id, lane, direction, location, number, owner = row
            location_name = cell_to_text(location)
            if location_name not in EXPECTED_LOCATION_ROWS:
                raise RuntimeError(f"Unexpected location in {path.name}: {location!r}")
            yield (
                cell_to_text(collected_at),
                cell_to_text(registered_at),
                cell_to_text(device_id),
                location_name,
                cell_to_text(lane),
                cell_to_text(direction),
                cell_to_text(number),
                cell_to_text(owner),
                path.name,
                SHEET_NAME,
            )
    finally:
        workbook.close()


def delete_existing_rows(conn: sqlite3.Connection) -> tuple[int, int]:
    placeholders = ", ".join("?" for _ in SOURCE_FILES)
    vehicle_rows = conn.execute(
        f"DELETE FROM {TABLE_NAME} WHERE source_file IN ({placeholders})", SOURCE_FILES
    ).rowcount
    summary_rows = conn.execute(
        f"DELETE FROM {SUMMARY_TABLE_NAME} WHERE source_file IN ({placeholders})", SOURCE_FILES
    ).rowcount
    return vehicle_rows, summary_rows


def import_file(conn: sqlite3.Connection, path: Path) -> int:
    sql = f"INSERT INTO {TABLE_NAME} ({', '.join(INSERT_COLUMNS)}) VALUES ({', '.join('?' for _ in INSERT_COLUMNS)})"
    batch: list[tuple[str, ...]] = []
    row_count = 0
    for row in iter_source_rows(path):
        batch.append(row)
        if len(batch) == 5000:
            conn.executemany(sql, batch)
            row_count += len(batch)
            batch.clear()
    if batch:
        conn.executemany(sql, batch)
        row_count += len(batch)
    if row_count != EXPECTED_FILE_ROWS[path.name]:
        raise RuntimeError(f"Row count mismatch in {path.name}: {row_count:,}")
    conn.execute(
        f"INSERT INTO {SUMMARY_TABLE_NAME} (source_file, source_sheet, column_count, imported_rows) VALUES (?, ?, ?, ?)",
        (path.name, SHEET_NAME, len(EXPECTED_HEADER), row_count),
    )
    return row_count


def validate_import(conn: sqlite3.Connection) -> None:
    placeholders = ", ".join("?" for _ in SOURCE_FILES)
    file_counts = dict(
        conn.execute(
            f"SELECT source_file, COUNT(*) FROM {TABLE_NAME} WHERE source_file IN ({placeholders}) GROUP BY source_file",
            SOURCE_FILES,
        )
    )
    location_counts = dict(
        conn.execute(
            f"SELECT location_name, COUNT(*) FROM {TABLE_NAME} WHERE source_file IN ({placeholders}) GROUP BY location_name",
            SOURCE_FILES,
        )
    )
    summary_counts = dict(
        conn.execute(
            f"SELECT source_file, imported_rows FROM {SUMMARY_TABLE_NAME} WHERE source_file IN ({placeholders})",
            SOURCE_FILES,
        )
    )
    total_rows = conn.execute(f"SELECT COUNT(*) FROM {TABLE_NAME}").fetchone()[0]
    if file_counts != EXPECTED_FILE_ROWS:
        raise RuntimeError(f"Imported file counts mismatch: {file_counts!r}")
    if location_counts != EXPECTED_LOCATION_ROWS:
        raise RuntimeError(f"Imported location counts mismatch: {location_counts!r}")
    if summary_counts != EXPECTED_FILE_ROWS:
        raise RuntimeError(f"Summary counts mismatch: {summary_counts!r}")
    if total_rows != EXPECTED_TOTAL_ROWS:
        raise RuntimeError(f"Total row count mismatch: {total_rows:,} != {EXPECTED_TOTAL_ROWS:,}")


def main() -> None:
    args = parse_args()
    if args.verify_only:
        conn = sqlite3.connect(f"file:{args.db.resolve().as_posix()}?mode=ro", uri=True)
        try:
            validate_db_schema(conn)
            validate_import(conn)
        finally:
            conn.close()
        print("verification=passed")
        return
    paths = source_paths(args.input_dir)
    conn = sqlite3.connect(args.db)
    try:
        validate_db_schema(conn)
        with conn:
            deleted_vehicle, deleted_summary = delete_existing_rows(conn)
            inserted_rows = sum(import_file(conn, path) for path in paths)
            validate_import(conn)
    finally:
        conn.close()
    print(f"deleted_vehicle_rows={deleted_vehicle}")
    print(f"deleted_import_summary_rows={deleted_summary}")
    print(f"inserted_rows={inserted_rows}")


if __name__ == "__main__":
    main()
