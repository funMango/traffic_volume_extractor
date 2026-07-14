#!/usr/bin/env python3
"""Generate a Leaflet map showing Samjeong-dong remicon camera locations."""

from __future__ import annotations

import argparse
import html
import json
import math
import re
from collections.abc import Iterable
from dataclasses import dataclass, replace
from pathlib import Path
from typing import Final

from openpyxl import load_workbook


PROJECT_ROOT: Final = Path(__file__).resolve().parents[2]
DEFAULT_INPUT_PATH: Final = PROJECT_ROOT / "00_Data" / "삼정동_위경도.xlsx"
DEFAULT_OUTPUT_PATH: Final = (
    PROJECT_ROOT / "02_Result" / "09_기타" / "삼정동_레미콘_카메라_위치지도.html"
)
REQUIRED_HEADERS: Final = ("도로", "위치명", "카메라 종류", "위도", "경도")
EXPECTED_CAMERA_COUNT: Final = 7
PROXIMITY_METERS: Final = 35.0
DISPLAY_OFFSET_METERS: Final = 20.0
# Input-row camera number -> map marker number: 1→7, 2→5, 3→6, 4→1, 5→2, 6→4, 7→3.
DISPLAY_NUMBERS: Final = (7, 5, 6, 1, 2, 4, 3)


@dataclass(frozen=True)
class CameraLocation:
    road: str
    name: str
    camera_type: str
    latitude: float
    longitude: float
    display_latitude: float | None = None
    display_longitude: float | None = None
    proximity_group: int | None = None

    @property
    def is_adjusted(self) -> bool:
        return self.display_latitude is not None and self.display_longitude is not None


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="삼정동 레미콘 카메라 위치 지도 HTML 생성")
    parser.add_argument(
        "--input", type=Path, default=DEFAULT_INPUT_PATH, help="좌표 입력 엑셀 경로"
    )
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT_PATH, help="생성할 HTML 경로")
    return parser.parse_args()


def normalized_text(value: object) -> str:
    return str(value).strip() if value is not None else ""


def load_locations(input_path: Path) -> list[CameraLocation]:
    if not input_path.is_file():
        raise FileNotFoundError(f"입력 엑셀을 찾을 수 없습니다: {input_path}")

    workbook = load_workbook(input_path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = list(worksheet.iter_rows(values_only=True))
    workbook.close()
    if not rows:
        raise RuntimeError("입력 엑셀이 비어 있습니다.")

    headers = tuple(normalized_text(value) for value in rows[0])
    if headers != REQUIRED_HEADERS:
        raise RuntimeError(
            f"엑셀 헤더가 일치하지 않습니다: expected={REQUIRED_HEADERS}, actual={headers}"
        )

    locations: list[CameraLocation] = []
    last_road = ""
    for row_number, row in enumerate(rows[1:], start=2):
        if not any(value is not None and normalized_text(value) for value in row):
            continue
        if len(row) != len(REQUIRED_HEADERS):
            raise RuntimeError(f"{row_number}행의 열 수가 올바르지 않습니다.")
        road, name, camera_type, latitude, longitude = row
        road_text = normalized_text(road) or last_road
        name_text = normalized_text(name)
        type_text = normalized_text(camera_type)
        if not all((road_text, name_text, type_text)):
            raise RuntimeError(f"{row_number}행에 도로, 위치명 또는 카메라 종류가 없습니다.")
        last_road = road_text
        try:
            latitude_value = float(latitude)
            longitude_value = float(longitude)
        except (TypeError, ValueError) as exc:
            raise RuntimeError(f"{row_number}행의 위도 또는 경도가 숫자가 아닙니다.") from exc
        if not 33.0 <= latitude_value <= 39.0 or not 124.0 <= longitude_value <= 132.0:
            raise RuntimeError(f"{row_number}행의 좌표가 대한민국 범위를 벗어났습니다.")
        locations.append(
            CameraLocation(road_text, name_text, type_text, latitude_value, longitude_value)
        )

    if len(locations) != EXPECTED_CAMERA_COUNT:
        raise RuntimeError(
            f"카메라 수가 일치하지 않습니다: expected={EXPECTED_CAMERA_COUNT}, actual={len(locations)}"
        )
    names = [location.name for location in locations]
    if len(set(names)) != len(names):
        raise RuntimeError("중복된 위치명이 있습니다.")
    return locations


def distance_meters(first: CameraLocation, second: CameraLocation) -> float:
    earth_radius = 6_371_000.0
    lat1, lat2 = math.radians(first.latitude), math.radians(second.latitude)
    delta_lat = lat2 - lat1
    delta_lon = math.radians(second.longitude - first.longitude)
    haversine = (
        math.sin(delta_lat / 2) ** 2
        + math.cos(lat1) * math.cos(lat2) * math.sin(delta_lon / 2) ** 2
    )
    return 2 * earth_radius * math.asin(math.sqrt(haversine))


def proximity_groups(locations: list[CameraLocation]) -> list[list[int]]:
    remaining = set(range(len(locations)))
    groups: list[list[int]] = []
    while remaining:
        seed = remaining.pop()
        component = {seed}
        frontier = [seed]
        while frontier:
            current = frontier.pop()
            adjacent = {
                candidate
                for candidate in remaining
                if distance_meters(locations[current], locations[candidate]) <= PROXIMITY_METERS
            }
            remaining -= adjacent
            component |= adjacent
            frontier.extend(adjacent)
        if len(component) > 1:
            groups.append(sorted(component))
    return groups


def offset_coordinate(
    latitude: float, longitude: float, bearing_radians: float
) -> tuple[float, float]:
    north_meters = DISPLAY_OFFSET_METERS * math.cos(bearing_radians)
    east_meters = DISPLAY_OFFSET_METERS * math.sin(bearing_radians)
    display_latitude = latitude + north_meters / 111_320.0
    display_longitude = longitude + east_meters / (111_320.0 * math.cos(math.radians(latitude)))
    return display_latitude, display_longitude


def adjust_display_locations(
    locations: list[CameraLocation],
) -> tuple[list[CameraLocation], list[list[int]]]:
    adjusted = locations.copy()
    groups = proximity_groups(locations)
    for group_number, indexes in enumerate(groups, start=1):
        center_latitude = sum(locations[index].latitude for index in indexes) / len(indexes)
        center_longitude = sum(locations[index].longitude for index in indexes) / len(indexes)
        for position, index in enumerate(indexes):
            angle = 2 * math.pi * position / len(indexes) - math.pi / 2
            display_latitude, display_longitude = offset_coordinate(
                center_latitude, center_longitude, angle
            )
            adjusted[index] = replace(
                locations[index],
                display_latitude=display_latitude,
                display_longitude=display_longitude,
                proximity_group=group_number,
            )
    return adjusted, groups


def map_record(location: CameraLocation, display_number: int) -> dict[str, object]:
    return {
        "displayNumber": display_number,
        "road": location.road,
        "name": location.name,
        "cameraType": location.camera_type,
        "latitude": location.latitude,
        "longitude": location.longitude,
        "displayLatitude": location.display_latitude or location.latitude,
        "displayLongitude": location.display_longitude or location.longitude,
        "isAdjusted": location.is_adjusted,
        "proximityGroup": location.proximity_group,
    }


def render_camera_index(locations: list[CameraLocation]) -> str:
    items: list[str] = []
    for location, display_number in zip(locations, DISPLAY_NUMBERS, strict=True):
        marker_class = "cctv" if location.camera_type == "방범CCTV" else "edge"
        items.append(
            "<li>"
            f'<span class="index-number {marker_class}">{display_number}</span>'
            "<div>"
            f"<strong>{html.escape(location.name)}</strong>"
            f"<span>{html.escape(location.road)} · {html.escape(location.camera_type)}</span>"
            "</div>"
            "</li>"
        )
    return "\n".join(items)


def render_html(locations: Iterable[CameraLocation]) -> str:
    location_list = list(locations)
    if len(location_list) != len(DISPLAY_NUMBERS):
        raise RuntimeError("지도 마커 번호 구성과 카메라 수가 일치하지 않습니다.")
    camera_data = json.dumps(
        [
            map_record(location, DISPLAY_NUMBERS[index])
            for index, location in enumerate(location_list)
        ],
        ensure_ascii=False,
    )
    camera_index = render_camera_index(location_list)
    return f"""<!doctype html>
<html lang="ko">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>삼정동 레미콘 카메라 위치 지도</title>
  <link rel="stylesheet" href="https://unpkg.com/leaflet@1.9.4/dist/leaflet.css">
  <style>
    body {{ margin: 0; background: #f1f5f9; color: #172033; font-family: "Malgun Gothic", "Noto Sans KR", sans-serif; }}
    main {{ max-width: 1200px; margin: 0 auto; padding: 28px 20px 36px; }}
    h1 {{ margin: 0 0 8px; font-size: 27px; }}
    .notice {{ margin: 0 0 18px; padding: 14px 16px; background: #fff7ed; border: 1px solid #fed7aa; border-radius: 10px; color: #7c2d12; line-height: 1.6; }}
    #map {{ height: min(72vh, 720px); min-height: 480px; border: 1px solid #cbd5e1; border-radius: 12px; box-shadow: 0 4px 16px #0f172a18; }}
    /* Prevent host-page image rules (for example max-width: 100%) from breaking Leaflet tile positions. */
    .leaflet-container img.leaflet-tile {{ max-width: none !important; max-height: none !important; }}
    .camera-marker {{ display: flex; align-items: center; justify-content: center; width: 32px; height: 32px; border: 3px solid #fff; border-radius: 50%; box-shadow: 0 2px 7px #0f172a80; color: #fff; font-size: 13px; font-weight: 700; }}
    .camera-marker.cctv {{ background: #2563eb; }}
    .camera-marker.edge {{ background: #dc2626; }}
    .legend {{ padding: 10px 12px; background: #fff; border: 1px solid #cbd5e1; border-radius: 8px; box-shadow: 0 2px 8px #0f172a1f; line-height: 1.8; }}
    .legend-title {{ font-weight: 700; margin-bottom: 3px; }}
    .legend-dot {{ display: inline-block; width: 12px; height: 12px; margin-right: 6px; border-radius: 50%; vertical-align: -1px; }}
    .legend-cctv {{ background: #2563eb; }} .legend-edge {{ background: #dc2626; }}
    .popup-title {{ margin: 0 0 8px; font-size: 15px; }} .popup-table {{ border-collapse: collapse; font-size: 13px; }} .popup-table th {{ text-align: left; padding: 3px 10px 3px 0; color: #475569; }} .popup-table td {{ padding: 3px 0; }}
    .camera-index {{ margin-top: 20px; padding: 18px; background: #fff; border: 1px solid #cbd5e1; border-radius: 12px; }}
    .camera-index h2 {{ margin: 0 0 14px; font-size: 18px; }} .camera-index ol {{ display: grid; grid-template-columns: repeat(2, minmax(0, 1fr)); gap: 10px 18px; margin: 0; padding: 0; list-style: none; }}
    .camera-index li {{ display: flex; align-items: center; gap: 9px; min-width: 0; }} .index-number {{ display: inline-flex; flex: 0 0 28px; align-items: center; justify-content: center; width: 28px; height: 28px; border-radius: 50%; color: #fff; font-size: 13px; font-weight: 700; }} .index-number.cctv {{ background: #2563eb; }} .index-number.edge {{ background: #dc2626; }}
    .camera-index strong, .camera-index span {{ display: block; }} .camera-index strong {{ overflow: hidden; text-overflow: ellipsis; white-space: nowrap; font-size: 14px; }} .camera-index li div > span {{ margin-top: 2px; color: #64748b; font-size: 12px; }}
    @media (max-width: 600px) {{ main {{ padding: 18px 10px 24px; }} h1 {{ font-size: 22px; }} #map {{ min-height: 420px; }} .camera-index ol {{ grid-template-columns: 1fr; }} }}
  </style>
</head>
<body>
  <main>
    <h1>삼정동 레미콘 카메라 위치 지도</h1>
    <p class="notice">가까운 카메라(35m 이내)는 실제 좌표를 보존한 채 가독성을 위해 표시 위치만 약 20m 벌려 표시했습니다. 점선은 실제 위치와 보정된 표시 위치를 연결합니다.</p>
    <div id="map" aria-label="삼정동 레미콘 카메라 위치 지도"></div>
    <section class="camera-index" id="camera-index" aria-label="카메라 번호 안내">
      <h2>카메라 번호 안내</h2>
      <ol>
        {camera_index}
      </ol>
    </section>
  </main>
  <script src="https://unpkg.com/leaflet@1.9.4/dist/leaflet.js" integrity="sha256-20nQCchB9co0qIjJZRGuk2/Z9VM+kNiyxNV1lvTlZBo=" crossorigin=""></script>
  <script>
    const cameras = {camera_data};
    const map = L.map('map', {{ scrollWheelZoom: true }});
    L.tileLayer('https://{{s}}.tile.openstreetmap.org/{{z}}/{{x}}/{{y}}.png', {{
      maxZoom: 19,
      attribution: '&copy; <a href="https://www.openstreetmap.org/copyright">OpenStreetMap</a> contributors'
    }}).addTo(map);

    const bounds = [];
    const typeClass = (cameraType) => cameraType === '방범CCTV' ? 'cctv' : 'edge';
    const popup = (camera) => {{
      const adjustment = camera.isAdjusted
        ? '<p style="margin:9px 0 0;color:#9a3412;">가독성을 위한 표시 위치 보정 적용 (점선은 실제 위치 연결)</p>'
        : '';
      return `<h2 class="popup-title">${{camera.name}}</h2><table class="popup-table"><tr><th>도로</th><td>${{camera.road}}</td></tr><tr><th>카메라 종류</th><td>${{camera.cameraType}}</td></tr><tr><th>실제 위도</th><td>${{camera.latitude.toFixed(6)}}</td></tr><tr><th>실제 경도</th><td>${{camera.longitude.toFixed(6)}}</td></tr></table>${{adjustment}}`;
    }};
    cameras.forEach((camera) => {{
      const marker = L.marker([camera.displayLatitude, camera.displayLongitude], {{
        icon: L.divIcon({{
          className: '',
          html: `<span class="camera-marker ${{typeClass(camera.cameraType)}}">${{camera.displayNumber}}</span>`,
          iconSize: [32, 32], iconAnchor: [16, 16]
        }}),
        title: camera.name
      }}).addTo(map).bindPopup(popup(camera));
      bounds.push([camera.latitude, camera.longitude], [camera.displayLatitude, camera.displayLongitude]);
      if (camera.isAdjusted) {{
        L.circleMarker([camera.latitude, camera.longitude], {{ radius: 4, color: '#475569', weight: 1, fillColor: '#ffffff', fillOpacity: 1 }}).addTo(map);
        L.polyline([[camera.latitude, camera.longitude], [camera.displayLatitude, camera.displayLongitude]], {{ color: '#64748b', weight: 1.5, dashArray: '4 5', opacity: 0.9 }}).addTo(map);
      }}
    }});
    map.fitBounds(bounds, {{ padding: [36, 36], maxZoom: 17 }});
    const legend = L.control({{ position: 'bottomright' }});
    legend.onAdd = () => {{
      const element = L.DomUtil.create('div', 'legend');
      element.innerHTML = '<div class="legend-title">카메라 종류</div><div><span class="legend-dot legend-cctv"></span>방범CCTV</div><div><span class="legend-dot legend-edge"></span>엣지카메라</div>';
      return element;
    }};
    legend.addTo(map);
  </script>
</body>
</html>"""


def validate_html(
    output_path: Path, locations: list[CameraLocation], groups: list[list[int]]
) -> None:
    content = output_path.read_text(encoding="utf-8")
    required_fragments = (
        "leaflet@1.9.4",
        "openstreetmap.org",
        "const cameras =",
        "가독성을 위한 표시 위치 보정",
        "L.polyline",
        "map.fitBounds",
        "카메라 종류",
        "카메라 번호 안내",
    )
    if any(fragment not in content for fragment in required_fragments):
        raise RuntimeError("생성 HTML에 필수 지도 구성 요소가 없습니다.")
    match = re.search(r"const cameras = (\[.*?\]);", content, flags=re.DOTALL)
    if match is None:
        raise RuntimeError("HTML 카메라 데이터가 없습니다.")
    camera_data = json.loads(match.group(1))
    if len(camera_data) != EXPECTED_CAMERA_COUNT or len(camera_data) != len(locations):
        raise RuntimeError("HTML 카메라 데이터 수가 올바르지 않습니다.")
    if sorted(camera["displayNumber"] for camera in camera_data) != list(
        range(1, EXPECTED_CAMERA_COUNT + 1)
    ):
        raise RuntimeError("HTML 지도 마커 번호가 1부터 7까지 한 번씩 포함되지 않습니다.")
    adjusted_count = sum(bool(camera["isAdjusted"]) for camera in camera_data)
    expected_adjusted_count = sum(len(group) for group in groups)
    if adjusted_count != expected_adjusted_count:
        raise RuntimeError("근접 카메라 표시 위치 보정 데이터가 올바르지 않습니다.")


def main() -> int:
    args = parse_args()
    locations = load_locations(args.input.resolve())
    display_locations, groups = adjust_display_locations(locations)
    output_path = args.output.resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(render_html(display_locations), encoding="utf-8")
    validate_html(output_path, display_locations, groups)
    print(f"output={output_path}")
    print(f"camera_count={len(locations)}")
    print(f"proximity_groups={len(groups)}")
    print(f"adjusted_camera_count={sum(len(group) for group in groups)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
