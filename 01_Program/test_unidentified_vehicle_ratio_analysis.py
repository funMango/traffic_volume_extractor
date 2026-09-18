from __future__ import annotations

import importlib.util
import sys
from datetime import date
from pathlib import Path

from openpyxl import load_workbook


MODULE_PATH = Path(__file__).parent / "00_main" / "차종_미확인_비율_분석.py"
SPEC = importlib.util.spec_from_file_location("unidentified_vehicle_ratio_analysis", MODULE_PATH)
assert SPEC is not None and SPEC.loader is not None
analysis = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = analysis
SPEC.loader.exec_module(analysis)


def test_period_label_uses_dates_only():
    assert analysis.Period(date(2026, 7, 1), date(2026, 7, 3)).label == "2026.07.01-2026.07.03"
    assert analysis.Period(date(2026, 9, 17), date(2026, 9, 17)).label == "2026.09.17"
    assert [period.label for period in analysis.parse_periods("260701~260703,260917")] == [
        "2026.07.01-2026.07.03",
        "2026.09.17",
    ]


def test_hours_label_compresses_sorted_contiguous_slots():
    assert analysis.hours_label([8, 7]) == "07:00-09:00"
    assert analysis.hours_label([10, 7]) == "07:00-08:00, 10:00-11:00"


def test_save_xlsx_separates_period_and_time_conditions(tmp_path):
    output = tmp_path / "analysis.xlsx"
    periods = [analysis.Period(date(2026, 7, 1), date(2026, 7, 3))]
    rows = [analysis.AnalysisRow(periods[0].label, "교차로", "북", "접근로", 2, 10)]

    analysis.save_xlsx(rows, periods, [analysis.Target(intersection="교차로")], [7, 8], output)

    workbook = load_workbook(output, data_only=False)
    try:
        sheet = workbook.active
        assert [(sheet.cell(row, 1).value, sheet.cell(row, 2).value) for row in range(4, 7)] == [
            ("기간", "2026.07.01-2026.07.03"),
            ("시간대", "07:00-09:00"),
            ("대상", "교차로"),
        ]
        assert [sheet.cell(10, column).value for column in range(1, 9)] == [
            "기간",
            "시간",
            "교차로",
            "방향",
            "접근로",
            "미확인 교통량",
            "전체 교통량",
            "미확인 교통량 비율(%)",
        ]
        assert sheet["A11"].value == "2026.07.01-2026.07.03"
        assert sheet["B11"].value == "07:00-09:00"
        assert sheet["H11"].value == "=IF(G11=0,0,F11/G11)"
        assert sheet["H11"].number_format == "0.00%"
        assert sheet.auto_filter.ref == "A10:H11"
        assert sheet.tables["UnidentifiedVehicleAnalysis"].ref == "A10:H11"
        assert sheet.freeze_panes == "A11"
    finally:
        workbook.close()


def test_parse_target_supports_multiple_normalized_directions():
    targets = analysis.parse_target("고강지하차도사거리-서(동향), 고강지하차도사거리-남 (북향)")

    assert [target.label for target in targets] == [
        "고강지하차도사거리-서(동향)",
        "고강지하차도사거리-남(북향)",
    ]


def test_build_day_sql_uses_one_or_filter_without_duplicate_aggregation():
    targets = analysis.parse_target("고강지하차도사거리, 고강지하차도사거리-서(동향)")

    sql, params = analysis.build_day_sql(targets, [7])

    assert (
        "((c.CRSRD_NM = :intersection_name0) OR (c.CRSRD_NM = :intersection_name1 AND a.ACSR_NM = :direction_name1))"
        in sql
    )
    assert params["intersection_name0"] == "고강지하차도사거리"
    assert params["direction_name1"] == "서(동향)"


def test_parse_target_rejects_all_mixed_with_explicit_target():
    try:
        analysis.parse_target("*, 고강지하차도사거리")
    except ValueError as exc:
        assert "함께 입력할 수 없습니다" in str(exc)
    else:
        raise AssertionError("전체와 명시 대상을 함께 입력하면 오류가 나야 합니다.")


def test_save_xlsx_shows_all_selected_targets(tmp_path):
    output = tmp_path / "analysis.xlsx"
    periods = [analysis.Period(date(2026, 7, 1), date(2026, 7, 1))]
    targets = analysis.parse_target("교차로-북(동향), 교차로-남(북향)")

    analysis.save_xlsx([], periods, targets, [7], output)

    workbook = load_workbook(output, data_only=False)
    try:
        assert workbook.active["B6"].value == "교차로-북(동향), 교차로-남(북향)"
    finally:
        workbook.close()
